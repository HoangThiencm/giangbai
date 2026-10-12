# -*- coding: utf-8 -*-
import os
import sys
sys.stdout.reconfigure(encoding='utf-8')

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
import pymupdf

def set_cell_margins(cell, top=6, bottom=6, left=10, right=10):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="1F2937", sz="4"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            </w:tblBorders>
        ''')
        tblPr[0].append(borders)

def add_single_slip(doc, slip_idx, is_last=False):
    # 1. Header
    p_h = doc.add_paragraph()
    p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_h.paragraph_format.line_spacing = 1.05
    p_h.paragraph_format.space_before = Pt(0)
    p_h.paragraph_format.space_after = Pt(1)
    
    r1 = p_h.add_run("TRƯỜNG THCS TRẦN PHÚ — TOÁN 9\n")
    r1.font.bold = True
    r1.font.size = Pt(9.0)
    r1.font.name = 'Times New Roman'
    
    r2 = p_h.add_run("PHIẾU TRẢ LỜI KIỂM TRA 15 PHÚT")
    r2.font.bold = True
    r2.font.size = Pt(10.0)
    r2.font.name = 'Times New Roman'
    r2.font.color.rgb = RGBColor(0x1e, 0x3a, 0x8a) # Blue

    # 2. Student Info & Code
    p_info = doc.add_paragraph()
    p_info.paragraph_format.line_spacing = 1.05
    p_info.paragraph_format.space_before = Pt(0)
    p_info.paragraph_format.space_after = Pt(2)
    
    r_name = p_info.add_run("Họ tên: ........................................................................ Lớp: 9.....\n")
    r_name.font.size = Pt(9.0)
    r_name.font.name = 'Times New Roman'
    
    r_code = p_info.add_run("MÃ ĐỀ: [      ]       ")
    r_code.font.bold = True
    r_code.font.size = Pt(9.5)
    r_code.font.name = 'Times New Roman'
    r_code.font.color.rgb = RGBColor(0xb9, 0x1c, 0x1c) # Red
    
    r_pts = p_info.add_run("ĐIỂM: [        / 10 ]")
    r_pts.font.bold = True
    r_pts.font.size = Pt(9.5)
    r_pts.font.name = 'Times New Roman'
    r_pts.font.color.rgb = RGBColor(0x1e, 0x3a, 0x8a) # Blue

    # 3. Answer Table (3 rows x 10 cols, total width ~ 7.4 inches)
    t = doc.add_table(rows=3, cols=10)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_table_borders(t, color="374151", sz="5")

    # Column widths: 
    # 4 cols for P1 (each 0.75 in -> 3.0 in)
    # 4 cols for P2 (each 0.60 in -> 2.4 in)
    # 2 cols for P3 (each 1.00 in -> 2.0 in)
    # Total = 7.4 inches!
    widths = [Inches(0.75)]*4 + [Inches(0.60)]*4 + [Inches(1.00)]*2
    for row in t.rows:
        for c_i, w in enumerate(widths):
            row.cells[c_i].width = w
            set_cell_margins(row.cells[c_i], top=4, bottom=4, left=4, right=4)

    # Row 0: Merge headers
    t.cell(0, 0).merge(t.cell(0, 3))
    p0 = t.cell(0, 0).paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(0)
    r = p0.add_run("PHẦN 1: TRẮC NGHIỆM KHÁCH QUAN")
    r.font.bold = True
    r.font.size = Pt(8.0)
    r.font.name = 'Times New Roman'
    set_cell_background(t.cell(0, 0), "F3F4F6")

    t.cell(0, 4).merge(t.cell(0, 7))
    p1 = t.cell(0, 4).paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(0)
    r = p1.add_run("PHẦN 2: TRẮC NGHIỆM ĐÚNG SAI")
    r.font.bold = True
    r.font.size = Pt(8.0)
    r.font.name = 'Times New Roman'
    set_cell_background(t.cell(0, 4), "F3F4F6")

    t.cell(0, 8).merge(t.cell(0, 9))
    p2 = t.cell(0, 8).paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    r = p2.add_run("PHẦN 3: TRẢ LỜI NGẮN")
    r.font.bold = True
    r.font.size = Pt(8.0)
    r.font.name = 'Times New Roman'
    set_cell_background(t.cell(0, 8), "F3F4F6")

    # Row 1: Sub-headers
    sub_titles = ["Câu 1", "Câu 2", "Câu 3", "Câu 4", "a.", "b.", "c.", "d.", "Câu 1", "Câu 2"]
    for idx, txt in enumerate(sub_titles):
        c = t.cell(1, idx)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(txt)
        r.font.bold = True
        r.font.size = Pt(8.0)
        r.font.name = 'Times New Roman'

    # Row 2: Empty student answer box
    for idx in range(10):
        c = t.cell(2, idx)
        set_cell_margins(c, top=12, bottom=12, left=4, right=4)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(" ")
        r.font.size = Pt(9.0)

    # Cut line between slips (if not last slip)
    if not is_last:
        p_cut = doc.add_paragraph()
        p_cut.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cut.paragraph_format.space_before = Pt(3)
        p_cut.paragraph_format.space_after = Pt(3)
        p_cut.paragraph_format.line_spacing = 1.0
        r_cut = p_cut.add_run("✂ - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - -")
        r_cut.font.size = Pt(7.0)
        r_cut.font.color.rgb = RGBColor(0x9c, 0xa3, 0xaf) # Gray

doc = docx.Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.PORTRAIT
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = Inches(0.3)
sec.bottom_margin = Inches(0.3)
sec.left_margin = Inches(0.4)
sec.right_margin = Inches(0.4)

# Add exactly 6 slips
for i in range(6):
    add_single_slip(doc, slip_idx=i+1, is_last=(i==5))

out_docx = r"TROLYTHIEN\2_TAO_BAI_TAP\Ket_qua\Phieu_Bai_Lam_6_Phieu_A4_Chuan.docx"
doc.save(out_docx)
print("Saved Phieu_Bai_Lam_6_Phieu_A4_Chuan.docx successfully!")
