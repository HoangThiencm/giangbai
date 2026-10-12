# -*- coding: utf-8 -*-
"""Bóc khung ma trận và bản đặc tả từ file Word .docx."""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

DEFAULT_DAU_VAO = Path(__file__).resolve().parents[1] / "2_TAO_BAI_TAP" / "Dau_vao"

MUC_DO_THU_TU = ("Nhận biết", "Thông hiểu", "Vận dụng", "Vận dụng cao")

_LEVEL_PATTERNS = (
    (re.compile(r"vận\s*dụng\s*cao|van\s*dung\s*cao", re.IGNORECASE), "Vận dụng cao"),
    (re.compile(r"thông\s*hiểu|thong\s*hieu", re.IGNORECASE), "Thông hiểu"),
    (re.compile(r"nhận\s*biết|nhan\s*biet", re.IGNORECASE), "Nhận biết"),
    (re.compile(r"vận\s*dụng|van\s*dung", re.IGNORECASE), "Vận dụng"),
    (re.compile(r"(?<![\w])biết(?![\w])|(?<![\w])biet(?![\w])", re.IGNORECASE), "Nhận biết"),
    (re.compile(r"(?<![\w])hiểu(?![\w])|(?<![\w])hieu(?![\w])", re.IGNORECASE), "Thông hiểu"),
)

_DANG_CAU = (
    (re.compile(r"đúng\s*[-–]?\s*sai|dung\s*[-–]?\s*sai", re.IGNORECASE), "Đúng - Sai"),
    (re.compile(r"trả\s*lời\s*ngắn|tra\s*loi\s*ngan", re.IGNORECASE), "Trả lời ngắn"),
    (re.compile(r"nhiều\s*lựa\s*chọn|nhieu\s*lua\s*chon|TNKQ|\bTN(?=\d|\b)", re.IGNORECASE), "Trắc nghiệm"),
    (re.compile(r"tự\s*luận|tu\s*luan|\bTL(?=\d|\b)", re.IGNORECASE), "Tự luận"),
)

_MA_CAU = re.compile(r"\b(TN|TL)\s*([0-9][0-9A-Da-d;,]*)", re.IGNORECASE)
_DIEM = re.compile(r"(\d+(?:[.,]\d+)?)\s*đ", re.IGNORECASE)
_SO_CAU_WORD = re.compile(r"(\d+)\s*câu", re.IGNORECASE)


class DuongDanKhongHopLe(ValueError):
    """Đường dẫn thư mục hoặc tệp ma trận không dùng được."""


def _fold(text: str) -> str:
    raw = unicodedata.normalize("NFD", text or "")
    raw = "".join(ch for ch in raw if unicodedata.category(ch) != "Mn")
    raw = raw.replace("đ", "d").replace("Đ", "D")
    return re.sub(r"\s+", " ", raw).strip().lower()


def _gop_o(row) -> list[str]:
    cells: list[str] = []
    for cell in row.cells:
        text = " ".join(cell.text.split())
        if not cells or cells[-1] != text:
            cells.append(text)
    return cells


def _bang_thanh_hang(table) -> list[list[str]]:
    return [_gop_o(row) for row in table.rows]


def chuan_hoa_duong_dan(raw: str | Path) -> Path:
    text = str(raw).strip().strip('"')
    if not text:
        raise DuongDanKhongHopLe("Đường dẫn trống.")
    path = Path(text).expanduser()
    if not path.is_absolute():
        path = Path.cwd() / path
    try:
        return path.resolve(strict=False)
    except OSError as exc:
        raise DuongDanKhongHopLe(f"Đường dẫn không hợp lệ: {text}") from exc


def _la_docx(path: Path) -> bool:
    return path.is_file() and path.suffix.lower() == ".docx" and not path.name.startswith("~$")


def _liet_ke_docx(folder: Path) -> list[Path]:
    found = [path for path in folder.rglob("*") if _la_docx(path)]
    return sorted(found, key=lambda item: str(item).lower())


def quet_file_dau_vao(input_path: str | Path | None = None) -> list[Path]:
    """Tìm file .docx ma trận. None nghĩa là quét đệ quy thư mục Dau_vao mặc định."""
    if input_path is None or str(input_path).strip() == "":
        root = DEFAULT_DAU_VAO
        if not root.is_dir():
            raise DuongDanKhongHopLe(f"Thư mục mặc định không tồn tại: {root}")
        return _liet_ke_docx(root)

    resolved = chuan_hoa_duong_dan(input_path)
    if not resolved.exists():
        raise DuongDanKhongHopLe(
            "Đường dẫn không tồn tại: "
            f"{resolved}. Kiểm tra khoảng trắng, tiếng Việt hoặc đường dẫn tương đối."
        )
    if resolved.is_file():
        if not _la_docx(resolved):
            raise DuongDanKhongHopLe(f"Không phải tệp ma trận .docx: {resolved}")
        return [resolved]
    if resolved.is_dir():
        return _liet_ke_docx(resolved)
    raise DuongDanKhongHopLe(f"Đường dẫn không hợp lệ: {resolved}")


def _co_yeu_cau_trong_bang(rows: list[list[str]]) -> int:
    hits = 0
    for row in rows:
        if any(_la_o_yeu_cau(cell) for cell in row):
            hits += 1
    return hits


def _nhan_loai_bang(rows: list[list[str]]) -> str:
    blob = _fold(" ".join(" ".join(row) for row in rows[:4]))
    if "yeu cau can dat" in blob or _co_yeu_cau_trong_bang(rows) >= 2:
        return "ban_dac_ta"
    if blob.startswith("tong so cau") or blob.startswith("tong so diem") or blob.startswith("ti le"):
        return "tong_hop"
    if any(token in blob for token in ("chu de", "chuong", "noi dung kien thuc", "ma tran", "muc do")):
        return "khung_ma_tran"
    return "khac"


def _la_dong_tong_hop(row: list[str]) -> bool:
    for cell in row:
        if not cell.strip():
            continue
        label = _fold(cell)
        return label.startswith("tong") or label.startswith("ti le")
    return False


def _tach_muc_do(text: str) -> tuple[str, str]:
    found: list[tuple[int, int, str, int]] = []
    for pattern, label in _LEVEL_PATTERNS:
        match = pattern.search(text or "")
        if match:
            found.append((match.start(), -(match.end() - match.start()), label, match.end()))
    if not found:
        return "", " ".join((text or "").split())
    found.sort()
    _start, _length, label, end = found[0]
    rest = (text or "")[end:].lstrip(" \t:–—-.")
    return label, " ".join(rest.split())


def _la_o_yeu_cau(text: str) -> bool:
    _label, rest = _tach_muc_do(text)
    return bool(_label) and len(rest) >= 12


def _dem_ma(blob: str) -> int:
    count = 0
    for part in re.split(r"[;,]", blob):
        token = part.strip()
        if re.fullmatch(r"\d+[A-Da-d]?", token) or re.fullmatch(r"[A-Da-d]", token):
            count += 1
    return count


def _dem_cau(text: str) -> int:
    word_counts = [int(item) for item in _SO_CAU_WORD.findall(text or "")]
    if word_counts:
        return sum(word_counts)
    total = 0
    for match in _MA_CAU.finditer(text or ""):
        total += _dem_ma(match.group(2))
    if total:
        return total
    if re.search(r"\b(TN|TL)\b", text or ""):
        return 1
    return 0


def _cong_diem(text: str) -> str:
    values = [item.replace(".", ",") for item in _DIEM.findall(text or "")]
    if not values:
        return ""
    total = sum(float(item.replace(",", ".")) for item in values)
    if abs(total - round(total)) < 1e-9:
        return str(int(round(total)))
    rendered = f"{total:.2f}".rstrip("0").rstrip(".")
    return rendered.replace(".", ",")


def _dang_cau(text: str) -> list[str]:
    found: list[str] = []
    for pattern, label in _DANG_CAU:
        if pattern.search(text or "") and label not in found:
            found.append(label)
    return found


def _so_nguyen_thuan(row: list[str]) -> list[int]:
    return [int(cell) for cell in row if re.fullmatch(r"\d+", cell.strip())]


def _tong_cau_tu_hang(row: list[str]) -> int | None:
    label = ""
    for cell in row:
        if cell.strip():
            label = _fold(cell)
            break
    blob = _fold(" ".join(row))
    if not label.startswith("tong"):
        return None
    if "diem" in label and "cau" not in label and "cau" not in blob:
        return None
    if label.startswith("ti le"):
        return None
    if "cau" not in blob and label not in {"tong", "tong:"} and not label.startswith("tong:"):
        return None

    pure = _so_nguyen_thuan(row)
    if len(pure) >= 2:
        for index, number in enumerate(pure):
            others = pure[:index] + pure[index + 1 :]
            if others and sum(others) == number:
                return number
    if pure and "cau" in blob:
        return max(pure)

    picked: list[int] = []
    for cell in row:
        match = re.search(r"(\d+)\s+\d+(?:[.,]\d+)?\s*đ", cell, flags=re.IGNORECASE)
        if match:
            picked.append(int(match.group(1)))
    if picked:
        return max(picked)
    return None


def _tong_diem_tu_hang(row: list[str]) -> float | None:
    label = ""
    for cell in row:
        if cell.strip():
            label = _fold(cell)
            break
    blob = _fold(" ".join(row))
    if not label.startswith("tong"):
        return None
    if "diem" not in blob and not _DIEM.search(" ".join(row)):
        return None
    numbers: list[float] = []
    for cell in row:
        if re.fullmatch(r"\d+(?:[.,]\d+)?", cell.strip()):
            numbers.append(float(cell.strip().replace(",", ".")))
        numbers.extend(float(item.replace(",", ".")) for item in _DIEM.findall(cell))
    if not numbers:
        return None
    if any(abs(number - 10) < 1e-9 for number in numbers):
        return 10.0
    if "diem" in label and "cau" not in label:
        return max(numbers)
    return None


def _boc_dong_dac_ta(row: list[str], nho: dict[str, str]) -> dict | None:
    if _la_dong_tong_hop(row):
        return None
    index = next((pos for pos, cell in enumerate(row) if _la_o_yeu_cau(cell)), -1)
    if index < 0:
        return None
    muc_do, yeu_cau = _tach_muc_do(row[index])
    content = [cell for cell in row[:index] if cell.strip() and not re.fullmatch(r"\d+", cell.strip())]
    chu_de = nho.get("chu_de", "")
    don_vi = nho.get("don_vi", "")
    if len(content) >= 2:
        chu_de, don_vi = content[0], content[1]
    elif len(content) == 1:
        cell = content[0]
        folded = _fold(cell)
        if re.match(r"^\d+\.", cell) or ("chuong" not in folded and chu_de):
            don_vi = cell
        else:
            chu_de = cell
            don_vi = cell
    if chu_de:
        nho["chu_de"] = chu_de
    if don_vi:
        nho["don_vi"] = don_vi
    tail = " ".join(cell for cell in row[index + 1 :] if cell.strip())
    return {
        "chu_de": chu_de,
        "don_vi_kien_thuc": don_vi,
        "muc_do": muc_do,
        "yeu_cau_can_dat": yeu_cau,
        "dang_cau": _dang_cau(tail),
        "so_cau": _dem_cau(tail),
        "so_diem": _cong_diem(tail),
        "mo_ta_cau": tail,
    }


def _meta_van_ban(doc: Document) -> tuple[list[str], str, str]:
    paragraphs = [" ".join(para.text.split()) for para in doc.paragraphs if para.text.strip()]
    blob = _fold(" ".join(paragraphs))
    mon = "Toán" if "toan" in blob else ""
    lop_match = re.search(r"lop\s*(\d+)", blob)
    lop = lop_match.group(1) if lop_match else ""
    return paragraphs[:8], mon, lop


def phan_tich_ma_tran_dac_ta(docx_path: str | Path) -> dict:
    """Đọc một file .docx và trả về khung ma trận, bản đặc tả và các yêu cầu cần đạt."""
    path = chuan_hoa_duong_dan(docx_path)
    if not path.is_file():
        raise DuongDanKhongHopLe(f"Không thấy tệp: {path}")
    if path.suffix.lower() != ".docx" or path.name.startswith("~$"):
        raise DuongDanKhongHopLe(f"Không phải tệp ma trận .docx: {path}")
    try:
        document = Document(str(path))
    except (PackageNotFoundError, OSError, ValueError) as exc:
        raise DuongDanKhongHopLe(f"Không đọc được tệp Word: {path}") from exc

    tieu_de, mon, lop = _meta_van_ban(document)
    khung: list[dict] = []
    dac_ta: list[dict] = []
    tong_hop: list[dict] = []
    cau_hoi: list[dict] = []
    tong_cau_candidates: list[int] = []
    tong_diem: float | None = None

    for table in document.tables:
        rows = _bang_thanh_hang(table)
        if not rows:
            continue
        loai = _nhan_loai_bang(rows)
        block = {"loai": loai, "hang": rows}
        if loai == "ban_dac_ta":
            dac_ta.append(block)
            nho = {"chu_de": "", "don_vi": ""}
            for row in rows:
                item = _boc_dong_dac_ta(row, nho)
                if item:
                    cau_hoi.append(item)
        elif loai == "khung_ma_tran":
            khung.append(block)
        elif loai == "tong_hop":
            tong_hop.append(block)
        else:
            khung.append(block)

        for row in rows:
            total = _tong_cau_tu_hang(row)
            if total:
                tong_cau_candidates.append(total)
            diem = _tong_diem_tu_hang(row)
            if diem is not None:
                tong_diem = diem

    if not tong_cau_candidates:
        tong_so_cau = sum(item["so_cau"] for item in cau_hoi)
    else:
        tong_so_cau = Counter(tong_cau_candidates).most_common(1)[0][0]

    muc_do = [label for label in MUC_DO_THU_TU if any(item["muc_do"] == label for item in cau_hoi)]
    yeu_cau = [item["yeu_cau_can_dat"] for item in cau_hoi if item["yeu_cau_can_dat"]]
    return {
        "tep": str(path),
        "ten_tep": path.name,
        "tieu_de": tieu_de,
        "mon": mon,
        "lop": lop,
        "khung_ma_tran": khung,
        "ban_dac_ta": dac_ta,
        "tong_hop": tong_hop,
        "cau_hoi": cau_hoi,
        "tong_so_cau": tong_so_cau,
        "tong_so_diem": tong_diem,
        "muc_do": muc_do,
        "yeu_cau_can_dat": yeu_cau,
    }


def phan_tich_dau_vao(input_path: str | Path | None = None) -> dict:
    files = quet_file_dau_vao(input_path)
    return {
        "so_tep": len(files),
        "tep": [phan_tich_ma_tran_dac_ta(path) for path in files],
    }


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Bóc ma trận và bản đặc tả từ file Word.")
    parser.add_argument("--input-dir", dest="input_dir", default=None, help="Thư mục chứa file .docx")
    parser.add_argument("--file", dest="file", default=None, help="Một tệp .docx")
    args = parser.parse_args(argv)
    if args.file and args.input_dir:
        parser.error("Chỉ dùng một trong --file hoặc --input-dir.")
    target = args.file if args.file else args.input_dir
    try:
        payload = phan_tich_dau_vao(target)
    except DuongDanKhongHopLe as exc:
        print(str(exc), file=sys.stderr)
        return 2
    json.dump(payload, sys.stdout, ensure_ascii=False, indent=2)
    sys.stdout.write("\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
