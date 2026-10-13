# -*- coding: utf-8 -*-
"""
Helper chuẩn hóa tạo tài liệu Đề cương ôn tập Word (.docx)
Năm học 2026 - 2027 (Đã bỏ cấp Phòng Giáo Dục)
Trường THCS Trần Phú - Tổ Toán - Tin
Tích hợp 100% Microsoft Word Equation (OMML m:oMath)
"""

import os, re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

import latex2mathml.commands
from latex2mathml.converter import convert
import lxml.etree as ET

# Khởi tạo XSLT Engine chuyển đổi MathML -> OMML
XSLT_PATHS = [
    r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files\Microsoft Office\Office16\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\Office16\MML2OMML.XSL",
]

_transform = None
for path in XSLT_PATHS:
    if os.path.exists(path):
        try:
            _xslt_doc = ET.parse(path)
            _transform = ET.XSLT(_xslt_doc)
            break
        except Exception:
            pass

def latex_to_omml(latex_code):
    """Chuyển chuỗi LaTeX sang XML Element OMML (m:oMath) của Microsoft Word."""
    if _transform is None:
        return None
    try:
        code = latex_code.strip()
        # Chuẩn hóa an toàn ký hiệu tiếng Việt ƯCLN và BCNN trong môi trường toán học
        code = code.replace(r"\mathrm{UCLN}", r"\text{ƯCLN}").replace(r"\mathrm{BCNN}", r"\text{BCNN}")
        if "ƯCLN" in code and r"\text{ƯCLN}" not in code:
            code = code.replace("ƯCLN", r"\text{ƯCLN}")
        if "BCNN" in code and r"\text{BCNN}" not in code:
            code = code.replace("BCNN", r"\text{BCNN}")
        
        # 1. Tránh lỗi cặp ngoặc nhọn rỗng \{ \} trong equation bị sinh thẻ m:e rỗng:
        code = re.sub(r'\\\{\s*\\\}', r'\{ \\text{ } \}', code)
        
        mathml = convert(code)
        tree = ET.fromstring(mathml)
        omml = _transform(tree)
        
        ns = {'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}

        # KHẮC PHỤC TRIỆT ĐỂ LỖI DẤU TRỪ SAU DẤU MỞ NGOẶC (MML2OMML separator bug):
        # Tránh ô vuông rỗng (<m:e/>) trong các biểu thức: (-2; 1), (-1/2)^2, (-x)...
        for d in omml.xpath('//m:d', namespaces=ns):
            sep = d.xpath('./m:dPr/m:sepChr', namespaces=ns)
            e_list = d.xpath('./m:e', namespaces=ns)
            if sep and sep[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/math}val') in ['\u2212', '-']:
                if len(e_list) >= 2 and len(e_list[0]) == 0 and not (e_list[0].text or '').strip():
                    sep[0].set('{http://schemas.openxmlformats.org/officeDocument/2006/math}val', '')
                    minus_r = ET.Element('{http://schemas.openxmlformats.org/officeDocument/2006/math}r')
                    minus_t = ET.SubElement(minus_r, '{http://schemas.openxmlformats.org/officeDocument/2006/math}t')
                    minus_t.text = '\u2212'
                    e_list[1].insert(0, minus_r)
                    d.remove(e_list[0])

        # Triệt tiêu các thẻ m:e rỗng mồ côi nếu còn sót trong delimiter
        for empty_e in omml.xpath('//m:e[not(node())]', namespaces=ns):
            p_d = empty_e.getparent()
            if p_d is not None and p_d.tag.endswith('d'):
                sibs = p_d.xpath('./m:e', namespaces=ns)
                if len(sibs) > 1:
                    p_d.remove(empty_e)

        # SỬA LỖI GẠCH TRÊN ĐẦU (\overline):
        # MML2OMML mặc định chuyển overline thành m:acc (dấu mũ accent) làm Word hiển thị thành ký tự dị dạng / chắp vá.
        # Chuyển m:acc thành m:bar với pos=top (thanh gạch ngang trên đầu chính thức của Word Office Math).
        for acc in omml.xpath('//m:acc', namespaces=ns):
            chr_elm = acc.xpath('./m:accPr/m:chr', namespaces=ns)
            if chr_elm:
                val = chr_elm[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/math}val')
                if val in ['―', '¯', '-', '\u2015', '\u00af']:
                    e_elms = acc.xpath('./m:e', namespaces=ns)
                    if e_elms:
                        bar = ET.Element('{http://schemas.openxmlformats.org/officeDocument/2006/math}bar')
                        barPr = ET.SubElement(bar, '{http://schemas.openxmlformats.org/officeDocument/2006/math}barPr')
                        pos = ET.SubElement(barPr, '{http://schemas.openxmlformats.org/officeDocument/2006/math}pos')
                        pos.set('{http://schemas.openxmlformats.org/officeDocument/2006/math}val', 'top')
                        bar.append(e_elms[0])
                        acc.getparent().replace(acc, bar)
                        
        omml_str = ET.tostring(omml, encoding='unicode')
        return parse_xml(omml_str)
    except Exception as e:
        print(f"LỖI LATEX_TO_OMML trên mã: '{latex_code}': {e}")
        return None

def add_text_with_math(paragraph, text, default_size_pt=13, default_font="Times New Roman", bold=False, italic=False, color_rgb=(0,0,0)):
    """
    Thêm đoạn văn bản hỗn hợp (gồm text thường và công thức toán $...$).
    Công thức toán kẹp giữa $...$ sẽ được hiển thị dưới dạng Microsoft Word Equation (m:oMath).
    Tự động rà soát bọc $...$ cho các công thức LaTeX bị lọt ra ngoài text thường.
    """
    from de_cuong_docx_auditor import auto_wrap_unwrapped_latex
    text = auto_wrap_unwrapped_latex(text)
    
    pattern = r'\$([^\$]+)\$'
    pos = 0
    for match in re.finditer(pattern, text):
        start, end = match.span()
        if start > pos:
            run = paragraph.add_run(text[pos:start])
            run.font.name = default_font
            run.font.size = Pt(default_size_pt)
            run.bold = bold
            run.italic = italic
            run.font.color.rgb = RGBColor(*color_rgb)
        
        latex_str = match.group(1).strip()
        omml_elm = latex_to_omml(latex_str)
        if omml_elm is not None:
            paragraph._p.append(omml_elm)
        else:
            # TUYỆT ĐỐI KHÔNG NUỐT LỖI LÀM GỘP SỐ / DÍNH CHỮ
            print(f"CRITICAL ERROR: Failed to convert LaTeX formula: '{latex_str}'")
            run_fb = paragraph.add_run(f"${latex_str}$")
            run_fb.font.name = default_font
            run_fb.font.size = Pt(default_size_pt)
            run_fb.italic = True
            run_fb.font.color.rgb = RGBColor(255, 0, 0)
            
        pos = end
        
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.font.name = default_font
        run.font.size = Pt(default_size_pt)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = RGBColor(*color_rgb)

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_table_borders(table, color="000000", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_elm = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders_elm)

def set_table_no_borders(table):
    tblPr = table._tbl.tblPr
    borders_elm = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders_elm)

def set_box_borders(table, color="1E3A8A", sz="8"):
    tblPr = table._tbl.tblPr
    borders_elm = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders_elm)

def apply_table_protection(table):
    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if i == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def format_run(run, font_name="Times New Roman", size_pt=13, bold=False, italic=False, color_rgb=(0,0,0)):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)

def create_base_document():
    doc = docx.Document()
    # Định lề trang chuẩn Nghị định 30 (Top: 20mm, Bottom: 20mm, Left: 30mm, Right: 15mm)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.79)     # ~20 mm
        section.bottom_margin = Inches(0.79)  # ~20 mm
        section.left_margin = Inches(1.18)    # ~30 mm
        section.right_margin = Inches(0.59)   # ~15 mm
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)
        
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(13)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_before = Pt(2)
    normal_style.paragraph_format.space_after = Pt(2)
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return doc

def add_header_quoc_hieu(doc, grade_str, school_year="2026 - 2027"):
    """Header chuẩn: Bỏ Phòng GD&ĐT, chỉ hiển thị TRƯỜNG THCS TRẦN PHÚ, Năm học 2026 - 2027."""
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_no_borders(table)
    
    cell_l, cell_r = table.rows[0].cells
    cell_l.width = Inches(3.0)
    cell_r.width = Inches(3.8)
    
    # Left (Đã bỏ Phòng GD&ĐT theo yêu cầu chỉ đạo mới)
    p_l = cell_l.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l.paragraph_format.line_spacing = 1.15
    p_l.paragraph_format.space_before = Pt(0)
    p_l.paragraph_format.space_after = Pt(0)
    
    r1 = p_l.add_run("TRƯỜNG THCS TRẦN PHÚ\n")
    format_run(r1, size_pt=12, bold=True)
    r2 = p_l.add_run("TỔ: TOÁN - TIN\n")
    format_run(r2, size_pt=12, bold=True)
    r3 = p_l.add_run("──────────")
    format_run(r3, size_pt=10, bold=True)
    
    # Right
    p_r = cell_r.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r.paragraph_format.line_spacing = 1.15
    p_r.paragraph_format.space_before = Pt(0)
    p_r.paragraph_format.space_after = Pt(0)
    
    r4 = p_r.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
    format_run(r4, size_pt=11.5, bold=True)
    r5 = p_r.add_run("Độc lập - Tự do - Hạnh phúc\n")
    format_run(r5, size_pt=12, bold=True)
    r6 = p_r.add_run("──────────────────")
    format_run(r6, size_pt=10, bold=True)
    
    # Title
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(6)
    p_sp.paragraph_format.space_after = Pt(2)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.line_spacing = 1.2
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    
    r_t = p_title.add_run("ĐỀ CƯƠNG ÔN TẬP GIỮA HỌC KỲ I\n")
    format_run(r_t, size_pt=15, bold=True, color_rgb=(15, 23, 42))
    
    r_m = p_title.add_run(f"MÔN: TOÁN {grade_str} — NĂM HỌC {school_year}\n")
    format_run(r_m, size_pt=13.5, bold=True, color_rgb=(30, 58, 138))
    
    r_sub = p_title.add_run("(Rèn luyện kỹ năng giải toán & Bộ 10 đề tự luận 100% thời gian 60 phút)")
    format_run(r_sub, size_pt=11.5, italic=True, color_rgb=(71, 85, 105))
    
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_before = Pt(0)
    p_line.paragraph_format.space_after = Pt(8)
    r_div = p_line.add_run("═" * 45)
    format_run(r_div, size_pt=10, bold=True, color_rgb=(148, 163, 184))

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    add_text_with_math(p, text, default_size_pt=14, bold=True, color_rgb=(30, 58, 138))

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    add_text_with_math(p, text, default_size_pt=13, bold=True, color_rgb=(180, 83, 9))

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    add_text_with_math(p, text, default_size_pt=13, bold=True, color_rgb=(15, 23, 42))

def add_body_paragraph(doc, text, bold_prefix="", indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        format_run(r_pre, size_pt=13, bold=True)
    add_text_with_math(p, text, default_size_pt=13)

def add_exercise(doc, num_str, content, sub_items=None):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True if sub_items else False
    
    # Bỏ các gợi ý phương pháp trong ngoặc đơn (ví dụ: (Phân phối có dấu trừ), (Lũy thừa...)),
    # chỉ giữ lại thang điểm nếu có chữ 'điểm' (ví dụ: '1 (2,5 điểm)') hoặc số bài thuần túy
    clean_num = re.sub(r"\s*\((?![^)]*điểm)[^)]*\)", "", str(num_str)).strip()
    
    if clean_num:
        r_num = p.add_run(f"Bài {clean_num}. ")
        format_run(r_num, size_pt=13, bold=True, color_rgb=(15, 23, 42))
    
    add_text_with_math(p, content, default_size_pt=13)
    
    if sub_items:
        for sub in sub_items:
            p_sub = doc.add_paragraph()
            p_sub.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_sub.paragraph_format.line_spacing = 1.15
            p_sub.paragraph_format.space_before = Pt(1)
            p_sub.paragraph_format.space_after = Pt(2)
            p_sub.paragraph_format.left_indent = Inches(0.4)
            add_text_with_math(p_sub, sub, default_size_pt=12.5)

def add_test_header(doc, test_num, grade_str, time_min=60):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    
    r_t = p.add_run(f"ĐỀ ÔN TẬP TỰ LUẬN SỐ {test_num:02d} — MÔN TOÁN {grade_str}\n")
    format_run(r_t, size_pt=13.5, bold=True, color_rgb=(180, 83, 9))
    
    r_sub = p.add_run(f"(Thời gian làm bài: {time_min} phút — Hình thức: 100% Tự luận — Thang điểm: 10 điểm)")
    format_run(r_sub, size_pt=11.5, italic=True, color_rgb=(71, 85, 105))

def add_answer_header(doc, test_num):
    """Tiêu đề cho phần hướng dẫn giải/đáp số của từng đề thi (tránh lỗi Bài Đáp số Đề X)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"Đáp số & Hướng dẫn Đề số {test_num:02d}")
    format_run(r, size_pt=13, bold=True, color_rgb=(30, 58, 138))

def add_answer_item(doc, item_text):
    """Xuất từng dòng đáp án bài tập."""
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.4)
    add_text_with_math(p, item_text, default_size_pt=12.5)
