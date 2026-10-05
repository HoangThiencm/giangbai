import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def safe_save_doc(doc, target_path):
    try:
        doc.save(target_path)
        print("Saved successfully:", target_path)
        return target_path
    except PermissionError:
        base, ext = os.path.splitext(target_path)
        fallback_path = f"{base}_CapNhat{ext}"
        try:
            doc.save(fallback_path)
            print(f"File locked by Word. Saved to fallback: {fallback_path}")
            return fallback_path
        except Exception as e:
            print(f"Error saving docx: {e}")
            return None

def add_system_comment_box(doc, comment_text, is_reject=False):
    """Tạo ô đóng khung tiện copy-paste nhận xét trực tiếp lên hệ thống duyệt giáo án"""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(7.0)
    
    border_color = "DC2626" if is_reject else "2563EB"
    bg_color = "FEF2F2" if is_reject else "F0FDF4"
    text_color = RGBColor(185, 28, 28) if is_reject else RGBColor(22, 101, 52)
    
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    
    # Border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    r_label = p.add_run("👉 [MẪU NHẬN XÉT HỆ THỐNG DUYỆT]: ")
    r_label.bold = True
    r_label.font.size = Pt(11)
    r_label.font.color.rgb = text_color
    
    r_content = p.add_run(f'"{comment_text}"')
    r_content.italic = True
    r_content.font.size = Pt(11)
    r_content.font.color.rgb = RGBColor(30, 41, 59)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)

def create_full_minutes():
    doc = docx.Document()
    
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.59)
        s.bottom_margin = Inches(0.59)
        s.left_margin = Inches(0.79)
        s.right_margin = Inches(0.59)
        s.different_first_page_header_footer = False
    
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(13)
    normal_style.paragraph_format.space_before = Pt(0)
    normal_style.paragraph_format.space_after = Pt(3)
    normal_style.paragraph_format.line_spacing = 1.0

    # Header section
    tbl_hdr = doc.add_table(rows=1, cols=2)
    tbl_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left, c_right = tbl_hdr.rows[0].cells
    c_left.width = Inches(3.2)
    c_right.width = Inches(3.8)

    p_l = c_left.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_l1 = p_l.add_run("TRƯỜNG THCS TRẦN PHÚ\n")
    r_l1.font.bold = True
    r_l1.font.size = Pt(12)
    r_l2 = p_l.add_run("TỔ TOÁN – TIN\n")
    r_l2.font.bold = True
    r_l2.font.size = Pt(12)
    r_l3 = p_l.add_run("Số: ... /BB-TT")
    r_l3.font.italic = True
    r_l3.font.size = Pt(11)

    p_r = c_right.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_r1 = p_r.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
    r_r1.font.bold = True
    r_r1.font.size = Pt(12)
    r_r2 = p_r.add_run("Độc lập - Tự do - Hạnh phúc\n")
    r_r2.font.bold = True
    r_r2.font.size = Pt(12)
    r_r3 = p_r.add_run("Xuân Đông, ngày 30 tháng 9 năm 2026")
    r_r3.font.italic = True
    r_r3.font.size = Pt(11)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(6)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("BIÊN BẢN KIỂM TRA HỒ SƠ GIÁO ÁN TỔ CHUYÊN MÔN\n")
    r_t.font.bold = True
    r_t.font.size = Pt(14)
    r_sub = p_title.add_run("ĐỢT THÁNG 9/2026 — NĂM HỌC 2026 - 2027")
    r_sub.font.bold = True
    r_sub.font.size = Pt(13)

    # Section I
    p_sec1 = doc.add_paragraph()
    p_sec1.add_run("I. THÀNH PHẦN KIỂM TRA\n").bold = True
    p_sec1.add_run("1. Tổ trưởng chuyên môn: ").bold = True
    p_sec1.add_run("Trưởng ban kiểm tra.\n")
    p_sec1.add_run("2. Tổ phó chuyên môn: ").bold = True
    p_sec1.add_run("Phó ban kiểm tra.\n")
    p_sec1.add_run("3. Các thành viên: ").bold = True
    p_sec1.add_run("Toàn thể giáo viên tổ Toán – Tin.")

    # Section II
    p_sec2 = doc.add_paragraph()
    p_sec2.add_run("II. CĂN CỨ VÀ TIÊU CHÍ ĐÁNH GIÁ CHUYÊN MÔN\n").bold = True
    p_sec2.add_run("1. Căn cứ Phụ lục 4 Kế hoạch bài dạy (Công văn 5512/BGDĐT-GDTrH) và Quy định chuyên môn trường THCS Trần Phú:\n")
    p_sec2.add_run("   • Hình thức thể thức: Font Times New Roman 13pt (in đứng); Căn lề Trên 1.5cm, Dưới 1.5cm, Trái 2.0cm, Phải 1.5cm; Giãn dòng 0pt / 3pt / Single.\n")
    p_sec2.add_run("   • Header bắt buộc: Trường THCS Trần Phú | Giáo viên: ... ; Footer bắt buộc: Môn/Phân môn | Trang | Năm học.\n")
    p_sec2.add_run("2. Tiến độ thực hiện: Đúng theo phân môn và tuần dạy đã duyệt tại Phụ lục 1 và Phụ lục 3.\n")
    p_sec2.add_run("3. Tích hợp NLS / AI: Khớp 1-1 cả Mã và Câu mô tả với Phụ lục 3; ")
    r_req = p_sec2.add_run("trong bài dạy bắt buộc phải in đậm, nghiêng. Nếu không khớp đề nghị trả hồ sơ.\n")
    r_req.bold = True
    r_req.italic = True
    p_sec2.add_run("4. Toán học: Đảm bảo tính chuẩn xác tuyệt đối về ký hiệu dấu mũ góc, dấu toán học (+/-), biến đổi đại số.\n")

    # Section III: Bảng tổng hợp nhanh toàn tổ
    p_sec3 = doc.add_paragraph()
    p_sec3.add_run("III. BẢNG TỔNG HỢP TIẾN ĐỘ VÀ XẾP LOẠI TOÀN TỔ").bold = True

    t_sum = doc.add_table(rows=3, cols=6)
    t_sum.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_sum, color="1F2937", sz="4")

    headers = [
        "STT", "Họ và tên GV", "Môn / Khối lớp", "Tổng số tiết đã nộp",
        "Tiến độ PPCT Tháng 9", "Kết luận chung"
    ]
    hdr_cells = t_sum.rows[0].cells
    col_widths = [Inches(0.5), Inches(1.8), Inches(1.4), Inches(1.2), Inches(1.1), Inches(1.0)]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(10)
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "E2E8F0")
        set_cell_margins(hdr_cells[i], top=70, bottom=70, left=70, right=70)

    # Row 1: Thầy Danh
    r1_cells = t_sum.rows[1].cells
    r1_data = [
        "1",
        "HỒ ĐĂNG DANH",
        "Toán 6, Toán 9",
        "33 tiết (4 phân môn)",
        "Thiếu 02 tiết (Tuần 4)",
        "TRẢ HỒ SƠ\n(Bổ sung bài)"
    ]
    for i, d in enumerate(r1_data):
        r1_cells[i].text = d
        p = r1_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in [0, 4, 5] else WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.size = Pt(9.5)
            if i == 5:
                run.font.bold = True
                run.font.color.rgb = RGBColor(180, 83, 9)
            elif i == 1:
                run.font.bold = True
            elif i == 4:
                run.font.bold = True
                run.font.color.rgb = RGBColor(185, 28, 28)
        r1_cells[i].width = col_widths[i]
        set_cell_margins(r1_cells[i], top=60, bottom=60, left=60, right=60)

    # Row 2: Thầy Hải
    r2_cells = t_sum.rows[2].cells
    r2_data = [
        "2",
        "TRẦN LONG HẢI",
        "Toán 8",
        "20 tiết (2 phân môn)",
        "Đủ 100% Tháng 9 & vượt Tuần 5",
        "TRẢ HỒ SƠ\n(Sửa sai sót toán)"
    ]
    for i, d in enumerate(r2_data):
        r2_cells[i].text = d
        p = r2_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in [0, 4, 5] else WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.size = Pt(9.5)
            if i == 5:
                run.font.bold = True
                run.font.color.rgb = RGBColor(185, 28, 28)
            elif i == 1:
                run.font.bold = True
        r2_cells[i].width = col_widths[i]
        set_cell_margins(r2_cells[i], top=60, bottom=60, left=60, right=60)

    # Section IV: KẾT LUẬN CHI TIẾT TỪNG GIÁO ÁN
    p_sec4 = doc.add_paragraph()
    p_sec4.paragraph_format.space_before = Pt(12)
    r_s4 = p_sec4.add_run("IV. KẾT LUẬN VÀ Ý KIẾN NHẬN XÉT TỪNG GIÁO ÁN\n(Sử dụng để nhập trực tiếp lên hệ thống quản lý duyệt giáo án điện tử)\n")
    r_s4.bold = True
    r_s4.font.size = Pt(13)

    # ------------------ 1. THẦY HỒ ĐĂNG DANH ------------------
    p_danh = doc.add_paragraph()
    p_danh.paragraph_format.space_before = Pt(6)
    r_d_name = p_danh.add_run("1. GIÁO VIÊN: HỒ ĐĂNG DANH (MÔN TOÁN 6, TOÁN 9)")
    r_d_name.bold = True
    r_d_name.font.size = Pt(12)

    # Danh - Phân môn 1: Đại số 9
    p_d1 = doc.add_paragraph()
    p_d1.add_run("1.1. Giáo án Phân môn Đại số 9 (14 tiết — Bài 1, 2, LTC, Bài 3, Ôn tập C1, Bài 4):\n").bold = True
    p_d1.add_run("• Tiến độ: Đạt và vượt tiến độ tháng 9 (soạn đến tuần 5).\n")
    p_d1.add_run("• Cấu trúc: Đầy đủ 4 hoạt động theo CV 5512, phân hóa nhiệm vụ rõ ràng.\n")
    p_d1.add_run("• Tích hợp NLS/AI: Có tích hợp NLS (5.3) và AI (9.B2) tại Bài 1 đúng địa chỉ, nhưng chưa định dạng in đậm nghiêng.\n")
    r_c1 = p_d1.add_run("• Kết luận duyệt: ")
    r_c1.bold = True
    r_c1_res = p_d1.add_run("DUYỆT CÓ ĐIỀU KIỆN (Yêu cầu in đậm nghiêng NLS/AI)\n")
    r_c1_res.bold = True
    r_c1_res.font.color.rgb = RGBColor(180, 83, 9)
    add_system_comment_box(
        doc,
        "Giáo án Phân môn Đại số 9 soạn đủ 14 tiết theo đúng PPCT, đảm bảo tiến độ tháng 9 và vượt tuần 5; cấu trúc đủ 4 hoạt động theo CV 5512. Tích hợp NLS (5.3) và AI (9.B2) tại Bài 1 đúng địa chỉ. Lưu ý: Cần in đậm, nghiêng câu mô tả năng lực số và AI theo đúng quy định chuyên môn trước khi đưa vào giảng dạy.",
        is_reject=False
    )

    # Danh - Phân môn 2: Hình học 9
    p_d2 = doc.add_paragraph()
    p_d2.add_run("1.2. Giáo án Phân môn Hình học 9 (03 tiết — Bài 11: Tỉ số lượng giác):\n").bold = True
    p_d2.add_run("• Tiến độ: ")
    r_d2_err = p_d2.add_run("CHƯA ĐẠT. Tháng 9 có 4 tuần, giáo án mới nộp 03 tiết (Bài 11), thiếu Tiết 4 (Bài 12: Bảng tỉ số lượng giác - Tuần 4).\n")
    r_d2_err.bold = True
    r_d2_err.font.color.rgb = RGBColor(185, 28, 28)
    p_d2.add_run("• Chuyên môn: Các tiết Bài 11 soạn chi tiết, định nghĩa và ví dụ chuẩn xác.\n")
    r_c2 = p_d2.add_run("• Kết luận duyệt: ")
    r_c2.bold = True
    r_c2_res = p_d2.add_run("KHÔNG DUYỆT / TRẢ HỒ SƠ (Yêu cầu bổ sung Tiết 4 Bài 12)\n")
    r_c2_res.bold = True
    r_c2_res.font.color.rgb = RGBColor(185, 28, 28)
    add_system_comment_box(
        doc,
        "Hồ sơ chưa đủ số tiết quy định của tháng 9 (4 tuần). Hiện tại mới có 03 tiết (Bài 11), thiếu Tiết 4 (Bài 12: Bảng lượng giác) theo đúng PPCT Phụ lục 3. Trả hồ sơ, đề nghị thầy soạn bổ sung Tiết 4 nộp lại tổ chuyên môn thẩm định.",
        is_reject=True
    )

    # Danh - Phân môn 3: Số học 6
    p_d3 = doc.add_paragraph()
    p_d3.add_run("1.3. Giáo án Phân môn Số học 6 (11 tiết — Bài 1 đến Bài 7: Thứ tự thực hiện phép tính):\n").bold = True
    p_d3.add_run("• Tiến độ: ")
    r_d3_err = p_d3.add_run("CHƯA ĐẠT. Đã nộp 11 tiết (đến hết Bài 7), thiếu Tiết 12 (Luyện tập chung Tuần 4) theo PPCT Phụ lục 3.\n")
    r_d3_err.bold = True
    r_d3_err.font.color.rgb = RGBColor(185, 28, 28)
    p_d3.add_run("• Tích hợp NLS: Có mã NLS 5.3 (Bài 6) và NLS 3.1 (Bài 7) nhưng chưa in đậm nghiêng.\n")
    r_c3 = p_d3.add_run("• Kết luận duyệt: ")
    r_c3.bold = True
    r_c3_res = p_d3.add_run("KHÔNG DUYỆT / TRẢ HỒ SƠ (Yêu cầu bổ sung Tiết 12 và định dạng NLS)\n")
    r_c3_res.bold = True
    r_c3_res.font.color.rgb = RGBColor(185, 28, 28)
    add_system_comment_box(
        doc,
        "Hồ sơ chưa đảm bảo tiến độ tháng 9. Đã soạn 11 tiết (đến hết Bài 7), còn thiếu Tiết 12 (Luyện tập chung tuần 4) theo PPCT Phụ lục 3. Phần mô tả năng lực số tại Bài 6, Bài 7 chưa in đậm nghiêng. Trả hồ sơ, đề nghị bổ sung Tiết 12 và định dạng in đậm nghiêng NLS.",
        is_reject=True
    )

    # Danh - Phân môn 4: Hình học 6
    p_d4 = doc.add_paragraph()
    p_d4.add_run("1.4. Giáo án Phân môn Hình học 6 (05 tiết — Bài 18, Bài 19: Tam giác đều, Lục giác đều, Hình vuông):\n").bold = True
    p_d4.add_run("• Tiến độ: Đạt và vượt tiến độ tháng 9 (soạn đến tuần 5).\n")
    p_d4.add_run("• Cấu trúc & Chuyên môn: Soạn tốt, hình vẽ trực quan, câu hỏi gợi mở phù hợp học sinh lớp 6.\n")
    r_c4 = p_d4.add_run("• Kết luận duyệt: ")
    r_c4.bold = True
    r_c4_res = p_d4.add_run("ĐẠT / DUYỆT\n")
    r_c4_res.bold = True
    r_c4_res.font.color.rgb = RGBColor(22, 101, 52)
    add_system_comment_box(
        doc,
        "Giáo án soạn đạt yêu cầu, đủ 05 tiết (đảm bảo tiến độ tháng 9 và vượt tuần 5). Cấu trúc 4 hoạt động rõ ràng, phân chia hoạt động GV - HS mạch lạc, hình vẽ và hệ thống câu hỏi phù hợp đối tượng học sinh lớp 6. Duyệt.",
        is_reject=False
    )

    # ------------------ 2. THẦY TRẦN LONG HẢI ------------------
    p_hai = doc.add_paragraph()
    p_hai.paragraph_format.space_before = Pt(8)
    r_h_name = p_hai.add_run("2. GIÁO VIÊN: TRẦN LONG HẢI (MÔN TOÁN 8 — 20 TIẾT, 89 TRANG)")
    r_h_name.bold = True
    r_h_name.font.size = Pt(12)

    p_h_sec_a = doc.add_paragraph()
    p_h_sec_a.add_run("A. KẾT LUẬN TỪNG BÀI DẠY — PHÂN MÔN HÌNH HỌC 8 (49 TRANG):\n").bold = True

    # Hải - Hình 8 - Bài 10
    p_h_b10 = doc.add_paragraph()
    p_h_b10.add_run("2.1. Bài 10: Tứ giác (Tiết 1, 2 — Trang 1 đến 8):\n").bold = True
    p_h_b10.add_run("• Cấu trúc: Đầy đủ 4 hoạt động. Có tích hợp NLS 1.1 và NLS 5.3 nhưng chưa in đậm nghiêng.\n")
    r_h_err = p_h_b10.add_run("• TỒN TẠI NGHIÊM TRỌNG TẠI TRANG 5:\n")
    r_h_err.bold = True
    r_h_err.font.color.rgb = RGBColor(185, 28, 28)
    p_h_b10.add_run("   (1) Mất hoàn toàn dấu mũ góc: Viết trần chữ cái in hoa (A + B + C + D = 360°).\n")
    p_h_b10.add_run("   (2) Sai bản chất định lý tổng các góc: Viết thành H + E + F - G = 360° (nhầm dấu trừ '- G').\n")
    p_h_b10.add_run("   (3) Sai quy tắc chuyển vế và thứ tự phép tính: D = 360° - A + B + C = 50° và F = 360° - H - E + G = 125°.\n")
    r_ch1 = p_h_b10.add_run("• Kết luận duyệt: ")
    r_ch1.bold = True
    r_ch1_res = p_h_b10.add_run("KHÔNG DUYỆT / TRẢ HỒ SƠ (Yêu cầu đính chính gấp Trang 5)\n")
    r_ch1_res.bold = True
    r_ch1_res.font.color.rgb = RGBColor(185, 28, 28)
    add_system_comment_box(
        doc,
        "TRẢ HỒ SƠ do có sai sót nghiêm trọng về kiến thức và ký hiệu toán học tại Trang 5: (1) Mất dấu mũ góc ở toàn bộ các công thức (viết trần A, B, C, D); (2) Sai bản chất định lý tổng các góc: viết H + E + F - G = 360° (nhầm dấu trừ); (3) Sai quy tắc chuyển vế và thứ tự phép tính: viết D = 360° - A + B + C = 50° và F = 360° - H - E + G = 125°. Đề nghị đính chính chuẩn xác ký hiệu góc và biểu thức toán học, in đậm nghiêng NLS 1.1, 5.3 trước khi duyệt.",
        is_reject=True
    )

    # Hải - Hình 8 - Bài 11
    p_h_b11 = doc.add_paragraph()
    p_h_b11.add_run("2.2. Bài 11: Hình thang cân (Tiết 3, 4 — Trang 9 đến 18):\n").bold = True
    p_h_b11.add_run("• Cấu trúc: Chuẩn CV 5512. Ký hiệu góc đúng chuẩn: góc DAB + góc A1 = 180°, góc C = góc D.\n")
    r_ch2 = p_h_b11.add_run("• Kết luận duyệt: ")
    r_ch2.bold = True
    r_ch2_res = p_h_b11.add_run("ĐẠT / DUYỆT\n")
    r_ch2_res.bold = True
    r_ch2_res.font.color.rgb = RGBColor(22, 101, 52)
    add_system_comment_box(
        doc,
        "Soạn tốt, đủ 4 hoạt động theo CV 5512. Ký hiệu góc và các bước chứng minh tính chất hình thang cân chuẩn xác. Tiến độ đúng PPCT. Duyệt.",
        is_reject=False
    )

    # Hải - Hình 8 - LTC Tiết 5
    p_h_ltc5 = doc.add_paragraph()
    p_h_ltc5.add_run("2.3. Luyện tập chung (Tiết 5 — Trang 19 đến 24):\n").bold = True
    p_h_ltc5.add_run("• Cấu trúc & Chuyên môn: Hệ thống bài tập tính góc và chứng minh hình thang cân chuẩn xác.\n")
    r_ch3 = p_h_ltc5.add_run("• Kết luận duyệt: ")
    r_ch3.bold = True
    r_ch3_res = p_h_ltc5.add_run("ĐẠT / DUYỆT\n")
    r_ch3_res.bold = True
    r_ch3_res.font.color.rgb = RGBColor(22, 101, 52)
    add_system_comment_box(
        doc,
        "Hệ thống bài tập củng cố tính chất góc và cạnh của tứ giác, hình thang cân bám sát SGK. Ký hiệu toán học chuẩn. Duyệt.",
        is_reject=False
    )

    # Hải - Hình 8 - Bài 12
    p_h_b12 = doc.add_paragraph()
    p_h_b12.add_run("2.4. Bài 12: Hình bình hành (Tiết 6, 7 — Trang 25 đến 36):\n").bold = True
    p_h_b12.add_run("• Cấu trúc: Soạn rất công phu (12 trang), hoạt động nhóm và hình vẽ trực quan rõ nét. Ký hiệu góc chuẩn.\n")
    r_ch4 = p_h_b12.add_run("• Kết luận duyệt: ")
    r_ch4.bold = True
    r_ch4_res = p_h_b12.add_run("ĐẠT / DUYỆT\n")
    r_ch4_res.bold = True
    r_ch4_res.font.color.rgb = RGBColor(22, 101, 52)
    add_system_comment_box(
        doc,
        "Kế hoạch bài dạy chi tiết (12 trang), phân hóa đối tượng học sinh tốt, ứng dụng đồ dùng dạy học trực quan. Ký hiệu góc chuẩn xác. Duyệt.",
        is_reject=False
    )

    # Hải - Hình 8 - LTC Tiết 8
    p_h_ltc8 = doc.add_paragraph()
    p_h_ltc8.add_run("2.5. Luyện tập chung (Tiết 8 — Trang 37 đến 40):\n").bold = True
    p_h_ltc8.add_run("• Cấu trúc & Chuyên môn: Rèn kỹ năng chứng minh tứ giác là hình bình hành logic, bài tập chặt chẽ.\n")
    r_ch5 = p_h_ltc8.add_run("• Kết luận duyệt: ")
    r_ch5.bold = True
    r_ch5_res = p_h_ltc8.add_run("ĐẠT / DUYỆT\n")
    r_ch5_res.bold = True
    r_ch5_res.font.color.rgb = RGBColor(22, 101, 52)
    add_system_comment_box(
        doc,
        "Bài tập rèn luyện kỹ năng chứng minh hình bình hành chuẩn mực, các bước suy luận logic. Duyệt.",
        is_reject=False
    )

    # Hải - Hình 8 - Bài 13
    p_h_b13 = doc.add_paragraph()
    p_h_b13.add_run("2.6. Bài 13: Hình chữ nhật (Tiết 9, 10 — Trang 41 đến 49):\n").bold = True
    p_h_b13.add_run("• Tiến độ & Chuyên môn: Soạn vượt tiến độ tuần 5, định nghĩa và tính chất chuẩn xác.\n")
    r_ch6 = p_h_b13.add_run("• Kết luận duyệt: ")
    r_ch6.bold = True
    r_ch6_res = p_h_b13.add_run("ĐẠT / DUYỆT\n")
    r_ch6_res.bold = True
    r_ch6_res.font.color.rgb = RGBColor(22, 101, 52)
    add_system_comment_box(
        doc,
        "Soạn vượt tiến độ tuần 5, nội dung định nghĩa, tính chất và dấu hiệu nhận biết hình chữ nhật rất chuẩn xác, khoa học. Duyệt.",
        is_reject=False
    )

    # Đánh giá chung gói Hình 8
    add_system_comment_box(
        doc,
        "[NHẬN XÉT NẾU DUYỆT THEO GÓI/TỆP HÌNH HỌC 8]: TRẢ LẠI HỒ SƠ. Yêu cầu thầy Trần Long Hải đính chính dứt điểm toàn bộ lỗi kiến thức định lý và ký hiệu góc tại Trang 5 Bài 10, in đậm nghiêng NLS trước khi phê duyệt toàn tệp.",
        is_reject=True
    )

    p_h_sec_b = doc.add_paragraph()
    p_h_sec_b.paragraph_format.space_before = Pt(8)
    p_h_sec_b.add_run("B. KẾT LUẬN TỪNG BÀI DẠY — PHÂN MÔN ĐẠI SỐ 8 (40 TRANG):\n").bold = True

    # Hải - Đại 8 - Bài 1
    p_h_a1 = doc.add_paragraph()
    p_h_a1.add_run("2.7. Bài 1: Đơn thức (Tiết 1, 2 — Trang 1 đến 8):\n").bold = True
    p_h_a1.add_run("• Cấu trúc: Đầy đủ 4 hoạt động, khái niệm đơn thức, bậc, hệ số diễn giải chuẩn.\n")
    p_h_a1.add_run("• Kết luận duyệt: ").bold = True
    r_cha1 = p_h_a1.add_run("ĐẠT / DUYỆT\n")
    r_cha1.bold = True
    r_cha1.font.color.rgb = RGBColor(22, 101, 52)
    add_system_comment_box(
        doc,
        "Khái niệm đơn thức, đơn thức thu gọn, bậc và hệ số diễn giải rõ ràng, ví dụ trực quan, bám sát SGK. Duyệt.",
        is_reject=False
    )

    # Hải - Đại 8 - Bài 2
    p_h_a2 = doc.add_paragraph()
    p_h_a2.add_run("2.8. Bài 2: Đa thức (Tiết 3, 4 — Trang 9 đến 16):\n").bold = True
    p_h_a2.add_run("• Cấu trúc: Đầy đủ 4 hoạt động, quy tắc thu gọn đa thức nhiều biến chuẩn xác.\n")
    p_h_a2.add_run("• Kết luận duyệt: ").bold = True
    r_cha2 = p_h_a2.add_run("ĐẠT / DUYỆT\n")
    r_cha2.bold = True
    r_cha2.font.color.rgb = RGBColor(22, 101, 52)
    add_system_comment_box(
        doc,
        "Nội dung nhận biết và thu gọn đa thức nhiều biến chuẩn xác, bài tập phân hóa tốt, hoạt động nhóm tích cực. Duyệt.",
        is_reject=False
    )

    # Hải - Đại 8 - Bài 3
    p_h_a3 = doc.add_paragraph()
    p_h_a3.add_run("2.9. Bài 3: Phép cộng và phép trừ đa thức (Tiết 5 — Trang 16 đến 21):\n").bold = True
    p_h_a3.add_run("• Cấu trúc & Chuyên môn: Rèn kỹ năng quy tắc dấu ngoặc khi trừ đa thức chặt chẽ. Khớp mã NLS 5.3.TC2a nhưng chưa in đậm nghiêng.\n")
    p_h_a3.add_run("• Kết luận duyệt: ").bold = True
    r_cha3 = p_h_a3.add_run("DUYỆT CÓ LƯU Ý (In đậm nghiêng NLS)\n")
    r_cha3.bold = True
    r_cha3.font.color.rgb = RGBColor(180, 83, 9)
    add_system_comment_box(
        doc,
        "Phương pháp rèn kỹ năng quy tắc dấu ngoặc khi trừ đa thức chặt chẽ. Có tích hợp NLS (5.3.TC2a) khớp Phụ lục 3. Lưu ý: Cần in đậm, nghiêng câu mô tả năng lực số theo đúng quy định chuyên môn.",
        is_reject=False
    )

    # Hải - Đại 8 - LTC Tiết 6
    p_h_a4 = doc.add_paragraph()
    p_h_a4.add_run("2.10. Luyện tập chung (Tiết 6 — Trang 22 đến 26):\n").bold = True
    p_h_a4.add_run("• Cấu trúc & Chuyên môn: Hệ thống bài tập 1.18 đến 1.23 SGK giải chuẩn xác, rèn tính cẩn thận cho học sinh.\n")
    p_h_a4.add_run("• Kết luận duyệt: ").bold = True
    r_cha4 = p_h_a4.add_run("ĐẠT / DUYỆT\n")
    r_cha4.bold = True
    r_cha4.font.color.rgb = RGBColor(22, 101, 52)
    add_system_comment_box(
        doc,
        "Hệ thống bài tập 1.18 đến 1.23 SGK sắp xếp khoa học, củng cố vững chắc kỹ năng cộng trừ đa thức. Duyệt.",
        is_reject=False
    )

    # Hải - Đại 8 - Bài 4
    p_h_a5 = doc.add_paragraph()
    p_h_a5.add_run("2.11. Bài 4: Phép nhân đa thức (Tiết 7, 8 — Trang 27 đến 33):\n").bold = True
    p_h_a5.add_run("• Cấu trúc & Chuyên môn: Quy tắc nhân đơn thức với đa thức, đa thức với đa thức mạch lạc, bài tập đa dạng.\n")
    p_h_a5.add_run("• Kết luận duyệt: ").bold = True
    r_cha5 = p_h_a5.add_run("ĐẠT / DUYỆT\n")
    r_cha5.bold = True
    r_cha5.font.color.rgb = RGBColor(22, 101, 52)
    add_system_comment_box(
        doc,
        "Hình thành quy tắc nhân đơn thức với đa thức, đa thức với đa thức mạch lạc, bài tập vận dụng phong phú, đúng PPCT. Duyệt.",
        is_reject=False
    )

    # Hải - Đại 8 - Bài 5
    p_h_a6 = doc.add_paragraph()
    p_h_a6.add_run("2.12. Bài 5: Phép chia đa thức cho đơn thức (Tiết 9, 10 — Trang 34 đến 40):\n").bold = True
    p_h_a6.add_run("• Tiến độ & Chuyên môn: Soạn vượt tiến độ tuần 5, điều kiện chia hết và quy tắc chia chuẩn xác.\n")
    p_h_a6.add_run("• Kết luận duyệt: ").bold = True
    r_cha6 = p_h_a6.add_run("ĐẠT / DUYỆT\n")
    r_cha6.bold = True
    r_cha6.font.color.rgb = RGBColor(22, 101, 52)
    add_system_comment_box(
        doc,
        "Soạn vượt tiến độ tuần 5, kiến thức điều kiện chia hết và quy tắc chia đơn thức cho đơn thức, đa thức cho đơn thức chuẩn xác. Duyệt.",
        is_reject=False
    )

    # Đánh giá chung gói Đại số 8
    add_system_comment_box(
        doc,
        "[NHẬN XÉT NẾU DUYỆT THEO GÓI/TỆP ĐẠI SỐ 8]: ĐẠT YÊU CẦU. Kế hoạch bài dạy soạn đầy đủ 10 tiết, đúng PPCT và vượt tiến độ. Nhắc nhở giáo viên in đậm, nghiêng câu mô tả năng lực số tại Bài 3 theo đúng quy định.",
        is_reject=False
    )

    # ------------------ GIÁO VIÊN TIẾP THEO ------------------
    p_next = doc.add_paragraph()
    p_next.paragraph_format.space_before = Pt(8)
    p_next.add_run("3. GIÁO VIÊN TIẾP THEO: ").bold = True
    p_next.add_run("[Đang chờ nạp hồ sơ kiểm tra cuốn chiếu tiếp theo...]\n").italic = True

    # Section V: Chữ ký phê duyệt
    p_sec5 = doc.add_paragraph()
    p_sec5.paragraph_format.space_before = Pt(14)
    tbl_sign = doc.add_table(rows=1, cols=2)
    tbl_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_left, s_right = tbl_sign.rows[0].cells
    s_left.width = Inches(3.5)
    s_right.width = Inches(3.5)

    p_sl = s_left.paragraphs[0]
    p_sl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sl.add_run("TỔ TRƯỞNG CHUYÊN MÔN\n").bold = True
    p_sl.add_run("(Ký và ghi rõ họ tên)\n\n\n\n\n")
    p_sl.add_run("....................................................")

    p_sr = s_right.paragraphs[0]
    p_sr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sr.add_run("PHÊ DUYỆT CỦA BAN GIÁM HIỆU\n").bold = True
    p_sr.add_run("HIỆU TRƯỞNG / PHÓ HIỆU TRƯỞNG\n\n\n\n\n")
    p_sr.add_run("....................................................")

    out_path = os.path.abspath("TROLYTHIEN/3_DUYET_GIAO_AN/Ket_qua/Bien_Ban_Kiem_Tra_Ho_So_To_Toan_Thang_9.docx")
    safe_save_doc(doc, out_path)

if __name__ == "__main__":
    os.makedirs("TROLYTHIEN/3_DUYET_GIAO_AN/Ket_qua", exist_ok=True)
    create_full_minutes()
