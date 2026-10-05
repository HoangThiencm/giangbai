import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def safe_save_doc(doc, target_path):
    try:
        doc.save(target_path)
        print("Saved successfully:", target_path)
        return target_path
    except PermissionError:
        base, ext = os.path.splitext(target_path)
        fallback_path = f"{base}_CapNhat{ext}"
        try:
            doc.save(fallback_path)
            print(f"File locked by Word. Saved to fallback: {fallback_path}")
            return fallback_path
        except Exception as e:
            print(f"Error saving docx: {e}")
            return None

def update_department_minutes():
    doc = docx.Document()
    
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.59)
        s.bottom_margin = Inches(0.59)
        s.left_margin = Inches(0.79)
        s.right_margin = Inches(0.59)
        s.different_first_page_header_footer = False
    
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(13)
    normal_style.paragraph_format.space_before = Pt(0)
    normal_style.paragraph_format.space_after = Pt(3)
    normal_style.paragraph_format.line_spacing = 1.0

    tbl_hdr = doc.add_table(rows=1, cols=2)
    tbl_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left, c_right = tbl_hdr.rows[0].cells
    c_left.width = Inches(3.2)
    c_right.width = Inches(3.8)

    p_l = c_left.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_l1 = p_l.add_run("TRƯỜNG THCS TRẦN PHÚ\n")
    r_l1.font.bold = True
    r_l1.font.size = Pt(12)
    r_l2 = p_l.add_run("TỔ TOÁN – TIN\n")
    r_l2.font.bold = True
    r_l2.font.size = Pt(12)
    r_l3 = p_l.add_run("Số: ... /BB-TT")
    r_l3.font.italic = True
    r_l3.font.size = Pt(11)

    p_r = c_right.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_r1 = p_r.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
    r_r1.font.bold = True
    r_r1.font.size = Pt(12)
    r_r2 = p_r.add_run("Độc lập - Tự do - Hạnh phúc\n")
    r_r2.font.bold = True
    r_r2.font.size = Pt(12)
    r_r3 = p_r.add_run("Xuân Đông, ngày 30 tháng 9 năm 2026")
    r_r3.font.italic = True
    r_r3.font.size = Pt(11)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(6)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("BIÊN BẢN KIỂM TRA HỒ SƠ GIÁO ÁN TỔ CHUYÊN MÔN\n")
    r_t.font.bold = True
    r_t.font.size = Pt(14)
    r_sub = p_title.add_run("ĐỢT THÁNG 9/2026 — NĂM HỌC 2026 - 2027")
    r_sub.font.bold = True
    r_sub.font.size = Pt(13)

    p_sec1 = doc.add_paragraph()
    p_sec1.add_run("I. THÀNH PHẦN KIỂM TRA\n").bold = True
    p_sec1.add_run("1. Tổ trưởng chuyên môn: ").bold = True
    p_sec1.add_run("Thầy/Cô [Tổ trưởng] - Trưởng ban\n")
    p_sec1.add_run("2. Tổ phó chuyên môn: ").bold = True
    p_sec1.add_run("Thầy/Cô [Tổ phó] - Thành viên\n")
    p_sec1.add_run("3. Các thành viên trong tổ: ").bold = True
    p_sec1.add_run("Toàn thể giáo viên tổ Toán – Tin.")

    p_sec2 = doc.add_paragraph()
    p_sec2.add_run("II. CĂN CỨ VÀ NỘI DUNG ĐÁNH GIÁ CHUYÊN MÔN\n").bold = True
    p_sec2.add_run("1. Căn cứ Phụ lục 4 Kế hoạch bài dạy (Công văn 5512/BGDĐT-GDTrH) và Quy định chuyên môn trường THCS Trần Phú:\n")
    p_sec2.add_run("   • Hình thức: Font Times New Roman 13pt; Lề trên/dưới/phải 1.5cm, lề trái 2.0cm; Giãn dòng 0pt/3pt/Single.\n")
    p_sec2.add_run("   • Bắt buộc có Header (Trường THCS Trần Phú | Giáo viên: ...) và Footer (Môn/Phân môn | Trang | Năm học).\n")
    p_sec2.add_run("2. Tiến độ thực hiện: Đúng theo phân môn và tuần dạy đã duyệt tại Phụ lục 1 và Phụ lục 3 (Kế hoạch giáo dục của giáo viên).\n")
    p_sec2.add_run("3. Nội dung tích hợp: Khớp 1-1 cả Mã và Câu mô tả với Phụ lục 1 & 3; ")
    r_req = p_sec2.add_run("bắt buộc phải in đậm, nghiêng. Không khớp đề nghị trả hồ sơ.\n")
    r_req.bold = True
    r_req.italic = True
    p_sec2.add_run("4. Toán học: Bắt buộc chuẩn xác tuyệt đối về ký hiệu dấu mũ góc, dấu toán học (+/-), biến đổi đại số.\n")

    p_sec3 = doc.add_paragraph()
    p_sec3.add_run("III. BẢNG TỔNG HỢP KẾT QUẢ KIỂM TRA TỪNG GIÁO VIÊN").bold = True

    table = doc.add_table(rows=3, cols=8)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="1F2937", sz="4")

    headers = [
        "STT", "Họ và tên GV", "Môn / Phân môn", "Số bài / tiết đã duyệt",
        "Tiến độ PPCT (PL3)", "Thể thức & CV 5512", "Tồn tại chuyên môn / Ký hiệu", "Xếp loại / Xử lý"
    ]
    hdr_cells = table.rows[0].cells
    col_widths = [Inches(0.4), Inches(1.2), Inches(1.0), Inches(0.9), Inches(1.0), Inches(0.9), Inches(1.2), Inches(0.8)]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(9.5)
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "E2E8F0")
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=80, right=80)

    # Row 1: Thầy Hồ Đăng Danh
    r1_cells = table.rows[1].cells
    r1_data = [
        "1",
        "HỒ ĐĂNG DANH",
        "• Đại số 9\n• Hình học 9\n• Số học 6\n• Hình học 6",
        "• Đại 9: 14 tiết\n• Hình 9: 3 tiết\n• Số 6: 11 tiết\n• Hình 6: 5 tiết",
        "Thiếu 01t Hình 9 (Bài 12) & 01t Số 6 (LTC tuần 4)",
        "Đạt CV 5512;\nHeader/Footer đúng mẫu",
        "• Chưa in đậm, nghiêng mô tả NLS/AI",
        "TRẢ HỒ SƠ\n(Bổ sung bài)"
    ]
    for i, d in enumerate(r1_data):
        r1_cells[i].text = d
        p = r1_cells[i].paragraphs[0]
        if i in [0, 7]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.size = Pt(9)
            if i == 7:
                run.font.bold = True
                run.font.color.rgb = RGBColor(180, 83, 9)
            elif i == 1:
                run.font.bold = True
            elif i == 4 and "Thiếu" in d:
                run.font.bold = True
                run.font.color.rgb = RGBColor(185, 28, 28)
        r1_cells[i].width = col_widths[i]
        set_cell_margins(r1_cells[i], top=70, bottom=70, left=70, right=70)

    # Row 2: Thầy Trần Long Hải
    r2_cells = table.rows[2].cells
    r2_data = [
        "2",
        "TRẦN LONG HẢI",
        "• Đại số 8\n• Hình học 8",
        "• Đại 8: 10 tiết\n• Hình 8: 10 tiết\n(20 tiết - 89 trang)",
        "Đủ 100% Tháng 9 (16t) & vượt Tuần 5",
        "Đạt CV 5512;\nHeader/Footer đúng mẫu",
        "• Trang 5 Hình 8: Mất dấu mũ góc (A, B, C..)\n• Sai định lý: H+E+F-G=360°\n• Sai dấu chuyển vế & thiếu ngoặc\n• Chưa in đậm, nghiêng NLS",
        "TRẢ HỒ SƠ\n(Sửa sai sót toán học)"
    ]
    for i, d in enumerate(r2_data):
        r2_cells[i].text = d
        p = r2_cells[i].paragraphs[0]
        if i in [0, 7]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.size = Pt(8.5)
            if i == 7:
                run.font.bold = True
                run.font.color.rgb = RGBColor(185, 28, 28)
            elif i == 1:
                run.font.bold = True
            elif i == 6 and "Sai" in d:
                run.font.bold = True
                run.font.color.rgb = RGBColor(185, 28, 28)
        r2_cells[i].width = col_widths[i]
        set_cell_margins(r2_cells[i], top=70, bottom=70, left=70, right=70)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    p_sec4 = doc.add_paragraph()
    p_sec4.add_run("IV. NHẬN XÉT CHI TIẾT VÀ KIẾN NGHỊ TỪNG THÀNH VIÊN\n").bold = True
    
    # GV 1
    p_danh = doc.add_paragraph()
    p_danh.add_run("1. Thầy Hồ Đăng Danh (Giáo viên Toán 6, Toán 9):\n").bold = True
    p_danh.add_run("   • Ưu điểm: Hồ sơ giáo án nộp đầy đủ các phân môn. Thể thức và Header/Footer đúng mẫu.\n")
    p_danh.add_run("   • Tồn tại: ").bold = True
    p_danh.add_run("Thiếu Tiết 4 (Bài 12 Hình học 9) và Tiết 12 (Số học 6) của Tuần 4 tháng 9. Các câu mô tả tích hợp NLS/AI chưa in đậm, nghiêng.\n")
    p_danh.add_run("   • Xử lý: ").bold = True
    p_danh.add_run("Trả hồ sơ đề nghị bổ sung đủ 2 tiết bài học và định dạng in đậm nghiêng đúng quy định.\n")

    # GV 2
    p_hai = doc.add_paragraph()
    p_hai.add_run("2. Thầy Trần Long Hải (Giáo viên Toán 8):\n").bold = True
    p_hai.add_run("   • Ưu điểm: ").bold = True
    p_hai.add_run("Hồ sơ chuẩn bị số lượng lớn (20 tiết - 89 trang), tiến độ đủ tháng 9 và vượt tuần 5. Thể thức văn bản và Header/Footer đúng mẫu trường THCS Trần Phú.\n")
    
    p_hai.add_run("   • SAI SÓT CHUYÊN MÔN TOÁN HỌC NGHIÊM TRỌNG (CẦN ĐÍNH CHÍNH GẤP): ").bold = True
    r_err = p_hai.add_run("Tại Trang 5 - Bài 10: Tứ giác (Hình học 8):\n")
    r_err.font.color.rgb = RGBColor(185, 28, 28)
    
    p_hai.add_run("     1) Vi phạm ký hiệu góc: ")
    r_e1 = p_hai.add_run("Toàn bộ các góc viết trần bằng chữ cái in hoa (A, B, C, D và H, E, F, G), không hề có dấu mũ góc (phải viết là góc A, góc B... hoặc ký hiệu mũ).\n")
    r_e1.font.color.rgb = RGBColor(185, 28, 28)
    
    p_hai.add_run("     2) Sai bản chất định lý: ")
    r_e2 = p_hai.add_run("Tại Luyện tập 2, công thức định lý tổng các góc viết thành: H + E + F - G = 360° (sai dấu trừ '- G' rất tai hại).\n")
    r_e2.font.color.rgb = RGBColor(185, 28, 28)
    
    p_hai.add_run("     3) Sai quy tắc chuyển vế và thứ tự phép tính: ")
    r_e3 = p_hai.add_run("Dòng biến đổi viết D = 360° - A + B + C và F = 360° - H - E + G (thiếu ngoặc đơn, sai dấu hoàn toàn). Dòng thay số: = 360° - 110° + 120° + 80° = 50° và = 360° - 55° + 90° + 90° = 125° là sai quy tắc toán học.\n")
    r_e3.font.color.rgb = RGBColor(185, 28, 28)
    
    p_hai.add_run("     4) Năng lực số: ")
    p_hai.add_run("Nội dung mô tả tích hợp NLS chưa được in đậm, nghiêng.\n")

    p_hai.add_run("   • XỬ LÝ CHUYÊN MÔN: ").bold = True
    r_act = p_hai.add_run("ĐỀ NGHỊ TRẢ HỒ SƠ. Yêu cầu thầy Trần Long Hải đính chính và sửa lại toàn bộ công thức, ký hiệu toán học tại Trang 5 và định dạng in đậm nghiêng các phần tích hợp NLS trước khi phê duyệt.")
    r_act.bold = True
    r_act.font.color.rgb = RGBColor(185, 28, 28)

    # Note for next teachers
    p_next = doc.add_paragraph()
    p_next.add_run("\n3. Giáo viên tiếp theo: ").bold = True
    p_next.add_run("[Đang chờ nạp hồ sơ kiểm tra cuốn chiếu...]\n").italic = True

    # Section V: Ký duyệt
    p_sec5 = doc.add_paragraph()
    p_sec5.paragraph_format.space_before = Pt(12)
    tbl_sign = doc.add_table(rows=1, cols=2)
    tbl_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_left, s_right = tbl_sign.rows[0].cells
    s_left.width = Inches(3.5)
    s_right.width = Inches(3.5)

    p_sl = s_left.paragraphs[0]
    p_sl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sl.add_run("TỔ TRƯỞNG CHUYÊN MÔN\n").bold = True
    p_sl.add_run("(Ký và ghi rõ họ tên)\n\n\n\n\n")
    p_sl.add_run("....................................................")

    p_sr = s_right.paragraphs[0]
    p_sr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sr.add_run("PHÊ DUYỆT CỦA BAN GIÁM HIỆU\n").bold = True
    p_sr.add_run("HIỆU TRƯỞNG / PHÓ HIỆU TRƯỞNG\n\n\n\n\n")
    p_sr.add_run("....................................................")

    out_path = os.path.abspath("TROLYTHIEN/3_DUYET_GIAO_AN/Ket_qua/Bien_Ban_Kiem_Tra_Ho_So_To_Toan_Thang_9.docx")
    safe_save_doc(doc, out_path)

def create_hai_individual_feedback():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.59)
        s.bottom_margin = Inches(0.59)
        s.left_margin = Inches(0.79)
        s.right_margin = Inches(0.59)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(13)
    normal_style.paragraph_format.space_before = Pt(0)
    normal_style.paragraph_format.space_after = Pt(3)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_school = p_title.add_run("TRƯỜNG THCS TRẦN PHÚ — TỔ TOÁN - TIN\n")
    r_school.font.size = Pt(11)
    r_school.font.bold = True
    r_t = p_title.add_run("PHIẾU NHẬN XÉT ĐÁNH GIÁ HỒ SƠ GIÁO ÁN\n")
    r_t.font.bold = True
    r_t.font.size = Pt(14)
    r_sub = p_title.add_run("Tháng 9/2026 — Năm học: 2026 - 2027")
    r_sub.font.italic = True
    r_sub.font.size = Pt(12)

    p_info = doc.add_paragraph()
    p_info.add_run("• Họ và tên giáo viên: ").bold = True
    p_info.add_run("TRẦN LONG HẢI\n").bold = True
    p_info.add_run("• Phân công giảng dạy: ").bold = True
    p_info.add_run("Toán 8 (Đại số 8 và Hình học 8)\n")
    p_info.add_run("• Số lượng tệp hồ sơ kiểm tra: ").bold = True
    p_info.add_run("02 tệp phân môn (Tổng cộng: 20 tiết — 89 trang giáo án)")

    p_sec = doc.add_paragraph()
    p_sec.add_run("I. ĐỐI CHIẾU TIẾN ĐỘ VÀ NỘI DUNG TỪNG PHÂN MÔN (THEO PHỤ LỤC 3)\n").bold = True

    items = [
        ("1. Phân môn Đại số 8 (10 tiết - 40 trang):", 
         "Đã soạn từ Tiết 1 đến Tiết 10 (Bài 1 -> Bài 5). Đủ tiến độ tháng 9 và vượt tuần 5. Đã tích hợp NLS 5.3.TC2a tại Bài 3."),
        ("2. Phân môn Hình học 8 (10 tiết - 49 trang) — [CÓ SAI SÓT TOÁN HỌC NGHIÊM TRỌNG]:", 
         "Đã soạn từ Tiết 1 đến Tiết 10 (Bài 10 -> Bài 13). Đủ số lượng tiết nhưng PHÁT HIỆN SAI SÓT KIẾN THỨC VÀ KÝ HIỆU TOÁN HỌC NGHIÊM TRỌNG TẠI TRANG 5 (Bài 10: Tứ giác).")
    ]

    for title, desc in items:
        p_item = doc.add_paragraph()
        r_tit = p_item.add_run(f"{title}\n")
        r_tit.bold = True
        if "SAI SÓT" in title:
            r_tit.font.color.rgb = RGBColor(185, 28, 28)
        p_item.add_run(f"   {desc}")

    p_eval = doc.add_paragraph()
    p_eval.paragraph_format.space_before = Pt(6)
    p_eval.add_run("II. CHI TIẾT SAI SÓT CHUYÊN MÔN CẦN ĐÍNH CHÍNH (TRANG 5 - HÌNH HỌC 8)\n").bold = True
    
    p_e1 = doc.add_paragraph()
    p_e1.add_run("1. Lỗi ký hiệu góc: ").bold = True
    p_e1.add_run("Viết góc hoàn toàn không có dấu mũ: A + B + C + D = 360° và H + E + F - G = 360°. Viết chữ trần A, B, C gây nhầm lẫn giữa góc và đỉnh/đoạn thẳng.\n")
    
    p_e2 = doc.add_paragraph()
    p_e2.add_run("2. Lỗi công thức định lý: ").bold = True
    p_e2.add_run("Viết H + E + F - G = 360° (nhầm dấu '+' thành dấu trừ '- G').\n")

    p_e3 = doc.add_paragraph()
    p_e3.add_run("3. Lỗi biến đổi đại số và thứ tự phép tính: ").bold = True
    p_e3.add_run("• Viết: D = 360° - A + B + C và F = 360° - H - E + G (thiếu ngoặc đơn, sai dấu hoàn toàn).\n")
    p_e3.add_run("• Dòng thay số: = 360° - 110° + 120° + 80° = 50° và = 360° - 55° + 90° + 90° = 125° (tính từ trái sang phải sẽ ra 450° và 485°, không thể ra 50° và 125°).\n")

    p_e4 = doc.add_paragraph()
    p_e4.add_run("4. Lỗi định dạng NLS: ").bold = True
    p_e4.add_run("Các nội dung tích hợp NLS chưa được in đậm, nghiêng theo Phụ lục 4.\n")

    p_res = doc.add_paragraph()
    p_res.add_run("III. KẾT LUẬN & XỬ LÝ CHUYÊN MÔN:\n").bold = True
    r_res = p_res.add_run("ĐỀ NGHỊ TRẢ HỒ SƠ ĐỂ CHỈNH SỬA, HOÀN THIỆN LẠI TRƯỚC KHI DUYỆT.")
    r_res.bold = True
    r_res.font.color.rgb = RGBColor(185, 28, 28)

    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.space_before = Pt(14)
    tbl_s = doc.add_table(rows=1, cols=2)
    tbl_s.alignment = WD_TABLE_ALIGNMENT.CENTER
    sl, sr = tbl_s.rows[0].cells
    sl.width = Inches(3.5)
    sr.width = Inches(3.5)

    psl = sl.paragraphs[0]
    psl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    psl.add_run("GIÁO VIÊN ĐƯỢC KIỂM TRA\n").bold = True
    psl.add_run("(Ký nhận lại hồ sơ để chỉnh sửa)\n\n\n\n\n")
    psl.add_run("Trần Long Hải")

    psr = sr.paragraphs[0]
    psr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    psr.add_run("NGƯỜI KIỂM TRA (TỔ TRƯỞNG)\n").bold = True
    psr.add_run("(Ký và ghi rõ họ tên)\n\n\n\n\n")
    psr.add_run("....................................................")

    out_path = os.path.abspath("TROLYTHIEN/3_DUYET_GIAO_AN/Ket_qua/Phieu_Nhan_Xet_Ho_So_Tran_Long_Hai_Thang_9.docx")
    safe_save_doc(doc, out_path)

if __name__ == "__main__":
    os.makedirs("TROLYTHIEN/3_DUYET_GIAO_AN/Ket_qua", exist_ok=True)
    update_department_minutes()
    create_hai_individual_feedback()
