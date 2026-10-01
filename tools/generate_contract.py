import os
import docx
from docx import Document
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_contract_docx(output_path):
    doc = Document()

    # 1. Page Margins A4 chuẩn NĐ 30/2020: Top 20mm, Bottom 20mm, Left 30mm, Right 15mm
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(30)
        section.right_margin = Mm(15)

    # 2. Base Normal Style: Times New Roman 13pt, line spacing 1.2, space 2pt/3pt
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(13)
    normal_font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_before = Pt(2)
    normal_style.paragraph_format.space_after = Pt(3)

    def set_font_run(run, name='Times New Roman', size_pt=13, bold=False, italic=False, color_rgb=(0,0,0)):
        run.font.name = name
        run.font.size = Pt(size_pt)
        run.bold = bold
        run.italic = italic
        if color_rgb:
            run.font.color.rgb = RGBColor(*color_rgb)
        rPr = run._r.get_or_add_rPr()
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{name}" w:hAnsi="{name}" w:cs="{name}" w:eastAsia="{name}"/>')
        rPr.append(rFonts)

    def remove_table_borders(table):
        tblPr = table._tbl.tblPr
        tblBorders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="none"/>\n'
            f'  <w:left w:val="none"/>\n'
            f'  <w:bottom w:val="none"/>\n'
            f'  <w:right w:val="none"/>\n'
            f'  <w:insideH w:val="none"/>\n'
            f'  <w:insideV w:val="none"/>\n'
            f'</w:tblBorders>'
        )
        tblPr.append(tblBorders)

    # --- TIÊU NGỮ & QUỐC HIỆU ---
    # Dùng bảng 2 cột hoặc 1 cột căn giữa. Với hợp đồng kinh tế/dân sự, Quốc hiệu Tiêu ngữ căn giữa trang hoặc góc phải.
    header_table = doc.add_table(rows=1, cols=2)
    remove_table_borders(header_table)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths = [Mm(70), Mm(95)]
    for row in header_table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    cell_left = header_table.cell(0, 0)
    p_left1 = cell_left.paragraphs[0]
    p_left1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_left1.paragraph_format.space_before = Pt(0)
    p_left1.paragraph_format.space_after = Pt(2)
    p_left1.paragraph_format.line_spacing = 1.15
    r = p_left1.add_run("NHẬT NAM BARBER STUDIO")
    set_font_run(r, size_pt=12, bold=True)

    p_left2 = cell_left.add_paragraph()
    p_left2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_left2.paragraph_format.space_before = Pt(0)
    p_left2.paragraph_format.space_after = Pt(0)
    p_left2.paragraph_format.line_spacing = 1.15
    r = p_left2.add_run("Số: ... /HĐĐT-NNBS")
    set_font_run(r, size_pt=11.5, italic=True)

    cell_right = header_table.cell(0, 1)
    p_right1 = cell_right.paragraphs[0]
    p_right1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_right1.paragraph_format.space_before = Pt(0)
    p_right1.paragraph_format.space_after = Pt(2)
    p_right1.paragraph_format.line_spacing = 1.15
    r = p_right1.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    set_font_run(r, size_pt=12, bold=True)

    p_right2 = cell_right.add_paragraph()
    p_right2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_right2.paragraph_format.space_before = Pt(0)
    p_right2.paragraph_format.space_after = Pt(2)
    p_right2.paragraph_format.line_spacing = 1.15
    r = p_right2.add_run("Độc lập - Tự do - Hạnh phúc")
    set_font_run(r, size_pt=13, bold=True)

    p_right3 = cell_right.add_paragraph()
    p_right3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_right3.paragraph_format.space_before = Pt(0)
    p_right3.paragraph_format.space_after = Pt(0)
    r = p_right3.add_run("───────────")
    set_font_run(r, size_pt=10, bold=True)

    # Khoảng cách sau header
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(6)
    p_space.paragraph_format.space_after = Pt(6)

    # --- TIÊU ĐỀ HỢP ĐỒNG ---
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("HỢP ĐỒNG ĐÀO TẠO NGHỀ")
    set_font_run(r_title, size_pt=15, bold=True)

    p_sub_title = doc.add_paragraph()
    p_sub_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub_title.paragraph_format.space_before = Pt(0)
    p_sub_title.paragraph_format.space_after = Pt(10)
    r_sub = p_sub_title.add_run("(V/v: Đào tạo nghề cắt tóc nam – Barber chuyên nghiệp)")
    set_font_run(r_sub, size_pt=13, italic=True)

    # --- CĂN CỨ PHÁP LÝ ---
    cancu_list = [
        "Căn cứ Bộ luật Dân sự số 91/2015/QH13 ngày 24 tháng 11 năm 2015 của Quốc hội;",
        "Căn cứ Bộ luật Lao động số 45/2019/QH14 ngày 20 tháng 11 năm 2019 của Quốc hội;",
        "Căn cứ Luật Giáo dục nghề nghiệp số 74/2014/QH13 ngày 27 tháng 11 năm 2014 của Quốc hội;",
        "Căn cứ vào nhu cầu học nghề thực tế của học viên và năng lực tiếp nhận đào tạo của Nhật Nam Barber Studio,"
    ]
    for idx, cc in enumerate(cancu_list):
        p_cc = doc.add_paragraph()
        p_cc.paragraph_format.first_line_indent = Inches(0.5)
        p_cc.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_cc.paragraph_format.space_before = Pt(2)
        p_cc.paragraph_format.space_after = Pt(2)
        p_cc.paragraph_format.line_spacing = 1.2
        r_cc = p_cc.add_run(cc)
        set_font_run(r_cc, size_pt=13, italic=True)

    # --- THỜI GIAN ĐỊA ĐIỂM ---
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.first_line_indent = Inches(0.5)
    p_intro.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_intro.paragraph_format.space_before = Pt(4)
    p_intro.paragraph_format.space_after = Pt(4)
    p_intro.paragraph_format.line_spacing = 1.2
    r_intro = p_intro.add_run("Hôm nay, ngày ... tháng ... năm 20..., tại cơ sở Nhật Nam Barber Studio, chúng tôi gồm có:")
    set_font_run(r_intro, size_pt=13)

    # --- BÊN A ---
    p_a_title = doc.add_paragraph()
    p_a_title.paragraph_format.first_line_indent = Inches(0.5)
    p_a_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_a_title.paragraph_format.space_before = Pt(4)
    p_a_title.paragraph_format.space_after = Pt(2)
    r_a_head = p_a_title.add_run("BÊN A: BÊN ĐÀO TẠO NGHỀ (NHẬT NAM BARBER STUDIO)")
    set_font_run(r_a_head, size_pt=13, bold=True)

    info_a = [
        ("Tên cơ sở đào tạo", "NHẬT NAM BARBER STUDIO"),
        ("Người đại diện", "[...]"),
        ("Chức vụ", "Chủ nhiệm / Đại diện cơ sở"),
        ("Địa chỉ cơ sở", "[...]"),
        ("Số điện thoại", "[...]"),
        ("Mã số thuế / CCCD", "[...]"),
        ("Tài khoản ngân hàng", "[...] tại Ngân hàng: [...]")
    ]
    for label, val in info_a:
        p_info = doc.add_paragraph()
        p_info.paragraph_format.first_line_indent = Inches(0.5)
        p_info.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_info.paragraph_format.space_before = Pt(2)
        p_info.paragraph_format.space_after = Pt(2)
        p_info.paragraph_format.line_spacing = 1.2
        r1 = p_info.add_run(f"- {label}: ")
        set_font_run(r1, size_pt=13)
        r2 = p_info.add_run(val)
        set_font_run(r2, size_pt=13, bold=(label in ["Tên cơ sở đào tạo"]))

    # --- BÊN B ---
    p_b_title = doc.add_paragraph()
    p_b_title.paragraph_format.first_line_indent = Inches(0.5)
    p_b_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_b_title.paragraph_format.space_before = Pt(4)
    p_b_title.paragraph_format.space_after = Pt(2)
    r_b_head = p_b_title.add_run("BÊN B: HỌC VIÊN ĐÀO TẠO NGHỀ")
    set_font_run(r_b_head, size_pt=13, bold=True)

    info_b = [
        ("Họ và tên học viên", "[...]"),
        ("Ngày, tháng, năm sinh", ".../.../.......               Giới tính: [...]"),
        ("Số CCCD/Định danh", "[...]"),
        ("Ngày cấp", ".../.../.......                 Nơi cấp: [...]"),
        ("Nơi đăng ký thường trú", "[...]"),
        ("Nơi ở hiện nay", "[...]"),
        ("Số điện thoại liên hệ", "[...]"),
        ("Người liên hệ khẩn cấp", "[...] - SĐT: [...] (Quan hệ: [...])")
    ]
    for label, val in info_b:
        p_info = doc.add_paragraph()
        p_info.paragraph_format.first_line_indent = Inches(0.5)
        p_info.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_info.paragraph_format.space_before = Pt(2)
        p_info.paragraph_format.space_after = Pt(2)
        p_info.paragraph_format.line_spacing = 1.2
        r1 = p_info.add_run(f"- {label}: ")
        set_font_run(r1, size_pt=13)
        r2 = p_info.add_run(val)
        set_font_run(r2, size_pt=13)

    # Lời dẫn
    p_lead = doc.add_paragraph()
    p_lead.paragraph_format.first_line_indent = Inches(0.5)
    p_lead.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_lead.paragraph_format.space_before = Pt(4)
    p_lead.paragraph_format.space_after = Pt(6)
    p_lead.paragraph_format.line_spacing = 1.2
    r_lead = p_lead.add_run("Sau khi trao đổi, bàn bạc trên tinh thần hoàn toàn tự nguyện, bình đẳng và cùng có lợi, hai bên thống nhất ký kết Hợp đồng đào tạo nghề với các điều khoản cụ thể sau đây:")
    set_font_run(r_lead, size_pt=13)

    # Hàm thêm điều khoản
    def add_article_title(title_text):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0.5)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(title_text)
        set_font_run(r, size_pt=13, bold=True)
        return p

    def add_body_p(text, bold_prefix="", indent_inch=0.5, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(indent_inch)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.2
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            set_font_run(r_pre, size_pt=13, bold=True)
        r_txt = p.add_run(text)
        set_font_run(r_txt, size_pt=13, italic=italic)
        return p

    # --- ĐIỀU 1 ---
    add_article_title("ĐIỀU 1. MỤC TIÊU VÀ NỘI DUNG ĐÀO TẠO NGHỀ")
    add_body_p("Ngành nghề đào tạo: Cắt tóc nam – Tạo mẫu tóc Barber chuyên nghiệp.", bold_prefix="1. ")
    add_body_p("Mục tiêu đào tạo: Giúp học viên nắm vững kiến thức lý thuyết cơ bản về tóc và da đầu; làm chủ kỹ thuật thao tác dụng cụ nghề (kéo, tông đơ, dao cạo); thành thạo các kỹ thuật cắt, cạo, sấy, tạo kiểu tóc nam từ cổ điển đến hiện đại; kỹ thuật uốn, nhuộm, tẩy hóa chất; xây dựng tác phong nghề nghiệp chuẩn mực, văn hóa ứng xử chuyên nghiệp và khả năng tự tin làm việc độc lập hoặc khởi nghiệp.", bold_prefix="2. ")
    add_body_p("Nội dung chương trình đào tạo bao gồm:", bold_prefix="3. ")
    add_body_p("Kiến thức tổng quan: Cấu trúc tóc, form đầu, nhận diện chất tóc; quy chuẩn an toàn lao động, vệ sinh dụng cụ và tiệt trùng y tế trong ngành tóc.", bold_prefix="a) ")
    add_body_p("Kỹ thuật dụng cụ: Tư thế đứng chuẩn công thái học, kỹ năng sử dụng thành thạo kéo cắt, kéo tỉa, tông đơ (Clipper, Trimmer, Shaver), lược và dao cạo.", bold_prefix="b) ")
    add_body_p("Kỹ thuật cắt và hoàn thiện mẫu tóc: Phân chia khối tóc, các góc độ cắt cơ bản và nâng cao; kỹ thuật Fade chuẩn (Skin Fade, Low/Mid/High Fade, Taper Fade); các mẫu tóc thịnh hành (Pompadour, Side Part, Undercut, Mohican, Mullet, Layer, Two Block...); kỹ thuật tỉa texture tạo liên kết form tóc.", bold_prefix="c) ")
    add_body_p("Kỹ thuật hóa chất: Kỹ thuật uốn lạnh, uốn nóng, uốn texture, uốn con sâu; quy trình nhuộm màu thời trang, nâng tone, tẩy tóc an toàn và phục hồi tóc chuyên sâu.", bold_prefix="d) ")
    add_body_p("Kỹ năng bổ trợ: Cạo mặt, tỉa râu, gội đầu thư giãn, massage đầu cơ bản; kỹ năng tư vấn phong cách theo khuôn mặt; phương pháp xây dựng thương hiệu cá nhân và quản lý vận hành studio.", bold_prefix="đ) ")

    # --- ĐIỀU 2 ---
    add_article_title("ĐIỀU 2. THỜI GIAN VÀ ĐỊA ĐIỂM ĐÀO TẠO")
    add_body_p("Thời gian đào tạo: Dự kiến [...] tháng (hoặc đào tạo theo tiến độ kèm cặp cho đến khi học viên vững tay nghề và đạt chuẩn đánh giá đầu ra của Studio), bắt đầu từ ngày .../.../20... đến ngày .../.../20...", bold_prefix="1. ")
    add_body_p("Lịch học: Học viên tham gia học tập từ Thứ [...] đến Thứ [...] hàng tuần, khung giờ từ [...] giờ đến [...] giờ.", bold_prefix="2. ")
    add_body_p("Địa điểm đào tạo: Tại cơ sở Nhật Nam Barber Studio, địa chỉ: [...]", bold_prefix="3. ")

    # --- ĐIỀU 3 ---
    add_article_title("ĐIỀU 3. HỌC PHÍ VÀ PHƯƠNG THỨC THANH TOÁN")
    add_body_p("Tổng học phí trọn gói khóa đào tạo: [...] VNĐ (Bằng chữ: [...]).", bold_prefix="1. ")
    add_body_p("Phương thức thanh toán: Học viên thanh toán bằng tiền mặt hoặc chuyển khoản vào tài khoản ngân hàng của Bên A theo tiến độ sau:", bold_prefix="2. ")
    add_body_p("Đợt 1: Thanh toán số tiền [...] VNĐ ngay khi ký kết Hợp đồng này.", bold_prefix="a) ")
    add_body_p("Đợt 2: Thanh toán số tiền [...] VNĐ sau [...] ngày kể từ ngày bắt đầu khóa học.", bold_prefix="b) ")
    add_body_p("Học phí nêu trên đã bao gồm toàn bộ chi phí giáo trình, chi phí đào tạo trực tiếp của người hướng dẫn, trang thiết bị học tập tại tiệm và nguồn mẫu thực hành dưới sự hướng dẫn của Bên A.", bold_prefix="3. ")
    add_body_p("Các khoản chi phí phát sinh (nếu có) phải được hai bên thống nhất trước khi thực hiện.", bold_prefix="4. ")

    # --- ĐIỀU 4 (NỘI DUNG CHÍNH XÁC TỪ USER) ---
    add_article_title("ĐIỀU 4. QUYỀN VÀ TRÁCH NHIỆM CỦA BÊN A")
    add_body_p("Bên A có trách nhiệm:", bold_prefix="1. ")
    add_body_p("Hướng dẫn học viên theo nội dung đào tạo đã thỏa thuận.", bold_prefix="a) ")
    add_body_p("Tạo điều kiện để học viên được thực hành trong phạm vi phù hợp.", bold_prefix="b) ")
    add_body_p("Hướng dẫn sử dụng máy móc, dụng cụ và kỹ thuật nghề.", bold_prefix="c) ")
    add_body_p("Nhắc nhở, góp ý và đánh giá quá trình học tập của học viên.", bold_prefix="d) ")
    add_body_p("Đảm bảo môi trường học tập nghiêm túc, chuyên nghiệp.", bold_prefix="đ) ")
    add_body_p("Bên A có quyền:", bold_prefix="2. ")
    add_body_p("Yêu cầu học viên tuân thủ nội quy của tiệm.", bold_prefix="a) ")
    add_body_p("Tạm dừng hoặc chấm dứt đào tạo nếu học viên vi phạm nghiêm trọng nội quy, gây thiệt hại tài sản hoặc ảnh hưởng đến khách hàng và hoạt động của tiệm.", bold_prefix="b) ")

    # --- ĐIỀU 5 (NỘI DUNG CHÍNH XÁC TỪ USER) ---
    add_article_title("ĐIỀU 5. QUYỀN VÀ TRÁCH NHIỆM CỦA BÊN B")
    add_body_p("Học viên có trách nhiệm:", bold_prefix="1. ")
    add_body_p("Đi học đúng giờ, đầy đủ và nghiêm túc.", bold_prefix="a) ")
    add_body_p("Tuân thủ sự hướng dẫn của người đào tạo.", bold_prefix="b) ")
    add_body_p("Giữ gìn máy móc, dụng cụ và tài sản của tiệm.", bold_prefix="c) ")
    add_body_p("Giữ vệ sinh khu vực làm việc.", bold_prefix="d) ")
    add_body_p("Có thái độ lịch sự, tôn trọng khách hàng và nhân viên.", bold_prefix="đ) ")
    add_body_p("Không tự ý sử dụng hoặc mang tài sản của tiệm ra ngoài.", bold_prefix="e) ")
    add_body_p("Không tự ý thực hiện các kỹ thuật trên khách khi chưa được người hướng dẫn cho phép.", bold_prefix="g) ")
    add_body_p("Chủ động luyện tập để nâng cao tay nghề.", bold_prefix="h) ")
    add_body_p("Học viên có quyền: Được học tập, thực hành theo đúng nội dung thỏa thuận; được người hướng dẫn tận tình chỉ dạy, giải đáp các thắc mắc chuyên môn và được hỗ trợ đánh giá, công nhận tay nghề khi hoàn thành khóa học.", bold_prefix="2. ")

    # --- ĐIỀU 6 (NỘI DUNG CHÍNH XÁC TỪ USER) ---
    add_article_title("ĐIỀU 6. NỘI QUY HỌC VIÊN")
    add_body_p("Học viên không được:", bold_prefix="1. ")
    add_body_p("Gây mất trật tự hoặc ảnh hưởng đến hoạt động của tiệm.", bold_prefix="a) ")
    add_body_p("Có hành vi thiếu tôn trọng khách hàng, nhân viên hoặc người hướng dẫn.", bold_prefix="b) ")
    add_body_p("Tự ý nghỉ học nhiều ngày mà không báo trước.", bold_prefix="c) ")
    add_body_p("Sử dụng rượu bia hoặc chất kích thích trong thời gian học.", bold_prefix="d) ")
    add_body_p("Làm hư hỏng tài sản do cố ý hoặc sử dụng sai quy định.", bold_prefix="đ) ")
    add_body_p("Tự ý nhận khách hoặc thực hiện dịch vụ bên ngoài với danh nghĩa của tiệm khi chưa được phép.", bold_prefix="e) ")
    add_body_p("Trường hợp vi phạm, Bên A có quyền nhắc nhở, cảnh cáo hoặc chấm dứt hợp đồng tùy theo mức độ vi phạm.", bold_prefix="2. ")

    # --- ĐIỀU 7 (NỘI DUNG CHÍNH XÁC TỪ USER) ---
    add_article_title("ĐIỀU 7. BẢO LƯU VÀ NGHỈ HỌC")
    add_body_p("Trường hợp học viên có lý do chính đáng cần nghỉ học hoặc bảo lưu, học viên phải thông báo cho Bên A trước.", bold_prefix="1. ")
    add_body_p("Việc bảo lưu, thời gian bảo lưu và các khoản học phí liên quan sẽ được hai bên thống nhất cụ thể bằng văn bản hoặc thỏa thuận riêng.", bold_prefix="2. ")

    # --- ĐIỀU 8 (NỘI DUNG CHÍNH XÁC TỪ USER) ---
    add_article_title("ĐIỀU 8. CHẤM DỨT HỢP ĐỒNG")
    add_body_p("Hợp đồng có thể chấm dứt trong các trường hợp:", bold_prefix="1. ")
    add_body_p("Hai bên thống nhất chấm dứt.", bold_prefix="a) ")
    add_body_p("Học viên hoàn thành chương trình đào tạo.", bold_prefix="b) ")
    add_body_p("Học viên tự nguyện nghỉ học.", bold_prefix="c) ")
    add_body_p("Học viên vi phạm nghiêm trọng nội quy của Nhật Nam Barber Studio.", bold_prefix="d) ")
    add_body_p("Một trong hai bên không thực hiện đúng các thỏa thuận trong hợp đồng.", bold_prefix="đ) ")
    add_body_p("Các vấn đề về học phí, tài sản và trách nhiệm liên quan khi chấm dứt hợp đồng sẽ được hai bên đối chiếu và giải quyết trên tinh thần thỏa thuận.", bold_prefix="2. ")

    # --- ĐIỀU 9 (NỘI DUNG CHÍNH XÁC TỪ USER) ---
    add_article_title("ĐIỀU 9. CAM KẾT CỦA HAI BÊN")
    add_body_p("Hai bên cam kết những thông tin cung cấp trong hợp đồng là đúng sự thật và tự nguyện thực hiện các nội dung đã thỏa thuận.", bold_prefix="1. ")
    add_body_p("Trong quá trình thực hiện, nếu phát sinh vấn đề chưa được quy định trong hợp đồng, hai bên sẽ trao đổi và thống nhất giải quyết.", bold_prefix="2. ")
    add_body_p("Hợp đồng được lập thành 02 bản, mỗi bên giữ 01 bản, có giá trị như nhau.", bold_prefix="3. ")
    add_body_p("Hợp đồng có hiệu lực kể từ ngày ký.", bold_prefix="4. ")

    # --- PHẦN KÝ TÊN HAI BÊN ---
    p_sign_space = doc.add_paragraph()
    p_sign_space.paragraph_format.space_before = Pt(8)
    p_sign_space.paragraph_format.space_after = Pt(4)

    sign_table = doc.add_table(rows=1, cols=2)
    remove_table_borders(sign_table)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in sign_table.rows:
        row.cells[0].width = Mm(80)
        row.cells[1].width = Mm(85)

    # Chữ ký Bên A
    c_a = sign_table.cell(0, 0)
    p1 = c_a.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.line_spacing = 1.15
    r = p1.add_run("ĐẠI DIỆN BÊN A\nNHẬT NAM BARBER STUDIO")
    set_font_run(r, size_pt=13, bold=True)

    p2 = c_a.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(50)
    p2.paragraph_format.line_spacing = 1.15
    r = p2.add_run("(Ký và ghi rõ họ tên)")
    set_font_run(r, size_pt=12, italic=True)

    p3 = c_a.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(0)
    p3.paragraph_format.space_after = Pt(0)
    r = p3.add_run("............................................")
    set_font_run(r, size_pt=12)

    # Chữ ký Bên B
    c_b = sign_table.cell(0, 1)
    p1_b = c_b.paragraphs[0]
    p1_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1_b.paragraph_format.space_before = Pt(0)
    p1_b.paragraph_format.space_after = Pt(2)
    p1_b.paragraph_format.line_spacing = 1.15
    r = p1_b.add_run("HỌC VIÊN – BÊN B")
    set_font_run(r, size_pt=13, bold=True)

    p2_b = c_b.add_paragraph()
    p2_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2_b.paragraph_format.space_before = Pt(0)
    p2_b.paragraph_format.space_after = Pt(50)
    p2_b.paragraph_format.line_spacing = 1.15
    r = p2_b.add_run("(Ký và ghi rõ họ tên)")
    set_font_run(r, size_pt=12, italic=True)

    p3_b = c_b.add_paragraph()
    p3_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3_b.paragraph_format.space_before = Pt(0)
    p3_b.paragraph_format.space_after = Pt(0)
    r = p3_b.add_run("............................................")
    set_font_run(r, size_pt=12)

    # Đảm bảo thư mục tồn tại và lưu file
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Successfully generated DOCX at: {output_path}")

if __name__ == '__main__':
    out_file = r"C:\Users\HoangThien\Documents\GitHub\giangbai\TROLYTHIEN\8_TAO_BAO_CAO\Ket_qua\Hop_Dong_Dao_Tao_Nghe_Nhat_Nam_Barber.docx"
    create_contract_docx(out_file)
