import os
import docx
from docx import Document
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def standardize_document(output_path):
    doc = Document()

    # 1. Page Margins chuẩn NĐ 30/2020: Top 20mm, Bottom 20mm, Left 30mm, Right 15mm
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(30)
        section.right_margin = Mm(15)

    # 2. Base Normal Style
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(13)
    normal_font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_before = Pt(2)
    normal_style.paragraph_format.space_after = Pt(3)

    def set_font(run, name='Times New Roman', size_pt=13, bold=False, italic=False, color_rgb=(0,0,0)):
        run.font.name = name
        run.font.size = Pt(size_pt)
        run.bold = bold
        run.italic = italic
        if color_rgb:
            run.font.color.rgb = RGBColor(*color_rgb)
        rPr = run._r.get_or_add_rPr()
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{name}" w:hAnsi="{name}" w:cs="{name}" w:eastAsia="{name}"/>')
        rPr.append(rFonts)

    def remove_table_borders(table):
        tblPr = table._tbl.tblPr
        tblBorders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="none"/>\n'
            f'  <w:left w:val="none"/>\n'
            f'  <w:bottom w:val="none"/>\n'
            f'  <w:right w:val="none"/>\n'
            f'  <w:insideH w:val="none"/>\n'
            f'  <w:insideV w:val="none"/>\n'
            f'</w:tblBorders>'
        )
        tblPr.append(tblBorders)

    def set_table_borders(table, color="000000", sz="4"):
        tblPr = table._tbl.tblPr
        tblBorders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'</w:tblBorders>'
        )
        tblPr.append(tblBorders)

    def set_cant_split(row):
        trPr = row._tr.get_or_add_trPr()
        cantSplit = parse_xml(f'<w:cantSplit {nsdecls("w")}/>')
        trPr.append(cantSplit)

    def set_tbl_header(row):
        trPr = row._tr.get_or_add_trPr()
        tblHeader = parse_xml(f'<w:tblHeader {nsdecls("w")}/>')
        trPr.append(tblHeader)

    def set_no_wrap(paragraph):
        pPr = paragraph._p.get_or_add_pPr()
        pPr.append(parse_xml(f'<w:noWrap {nsdecls("w")}/>'))

    def tighten_run(run, twips):
        rPr = run._r.get_or_add_rPr()
        rPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:val="{twips}"/>'))

    def add_vector_line(container, length_mm, space_before=0, space_after=2):
        paragraph = container.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(space_before)
        paragraph.paragraph_format.space_after = Pt(space_after)
        paragraph.paragraph_format.line_spacing = 1.0
        length_pt = round(length_mm * 72 / 25.4, 2)
        run_xml = (
            f'<w:r {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml" '
            f'xmlns:o="urn:schemas-microsoft-com:office:office">'
            f'<w:pict><v:line style="position:relative" from="0pt,0pt" to="{length_pt}pt,0pt" '
            f'strokecolor="#000000" strokeweight="0.75pt"/></w:pict></w:r>'
        )
        paragraph._p.append(parse_xml(run_xml))
        return paragraph

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>\n'
            f'  <w:top w:w="{top}" w:type="dxa"/>\n'
            f'  <w:bottom w:w="{bottom}" w:type="dxa"/>\n'
            f'  <w:left w:w="{left}" w:type="dxa"/>\n'
            f'  <w:right w:w="{right}" w:type="dxa"/>\n'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    # ==========================
    # 1. HEADER (BẢNG 2 CỘT)
    # ==========================
    header_table = doc.add_table(rows=3, cols=2)
    remove_table_borders(header_table)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    col_widths = [Mm(65), Mm(100)]
    tbl = header_table._tbl
    tblPr = tbl.tblPr
    tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="9354" w:type="dxa"/>'))
    tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
    grid = parse_xml(
        f'<w:tblGrid {nsdecls("w")}>'
        f'<w:gridCol w:w="3685"/>'
        f'<w:gridCol w:w="5669"/>'
        f'</w:tblGrid>'
    )
    tbl.insert(1, grid)
    for row in header_table.rows:
        for idx, width in enumerate(col_widths):
            cell = row.cells[idx]
            cell.width = width
            set_cell_margins(cell, top=0, bottom=0, left=0, right=0)

    # Hàng 0, Cột 0: UBND XÃ XUÂN ĐÔNG — một dòng, cột 65mm
    p00 = header_table.cell(0, 0).paragraphs[0]
    p00.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p00.paragraph_format.space_before = Pt(0)
    p00.paragraph_format.space_after = Pt(2)
    p00.paragraph_format.line_spacing = 1.0
    set_no_wrap(p00)
    r = p00.add_run("ỦY BAN NHÂN DÂN XÃ XUÂN ĐÔNG")
    set_font(r, size_pt=12, bold=False)
    tighten_run(r, -18)

    # Hàng 1, Cột 0: TRƯỜNG THCS TRẦN PHÚ
    p10 = header_table.cell(1, 0).paragraphs[0]
    p10.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p10.paragraph_format.space_before = Pt(0)
    p10.paragraph_format.space_after = Pt(1)
    p10.paragraph_format.line_spacing = 1.15
    r = p10.add_run("TRƯỜNG THCS TRẦN PHÚ")
    set_font(r, size_pt=12, bold=True)
    set_no_wrap(p10)
    add_vector_line(header_table.cell(1, 0), 28, space_before=0, space_after=2)

    # Hàng 2, Cột 0: Số ký hiệu
    p20 = header_table.cell(2, 0).paragraphs[0]
    p20.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p20.paragraph_format.space_before = Pt(0)
    p20.paragraph_format.space_after = Pt(0)
    p20.paragraph_format.line_spacing = 1.15
    r = p20.add_run("Số: … /KH-THCSTP")
    set_font(r, size_pt=12)

    # Hàng 0, Cột 1: CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM
    p01 = header_table.cell(0, 1).paragraphs[0]
    p01.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p01.paragraph_format.space_before = Pt(0)
    p01.paragraph_format.space_after = Pt(2)
    p01.paragraph_format.line_spacing = 1.0
    set_no_wrap(p01)
    r = p01.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    set_font(r, size_pt=12, bold=True)
    tighten_run(r, -8)

    # Hàng 1, Cột 1: Độc lập - Tự do - Hạnh phúc
    p11 = header_table.cell(1, 1).paragraphs[0]
    p11.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p11.paragraph_format.space_before = Pt(0)
    p11.paragraph_format.space_after = Pt(1)
    p11.paragraph_format.line_spacing = 1.15
    r = p11.add_run("Độc lập - Tự do - Hạnh phúc")
    set_font(r, size_pt=13, bold=True)
    add_vector_line(header_table.cell(1, 1), 38, space_before=0, space_after=2)

    # Hàng 2, Cột 1: Địa danh, ngày tháng
    p21 = header_table.cell(2, 1).paragraphs[0]
    p21.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p21.paragraph_format.space_before = Pt(0)
    p21.paragraph_format.space_after = Pt(0)
    p21.paragraph_format.line_spacing = 1.15
    r = p21.add_run("Xuân Đông, ngày 01 tháng 10 năm 2026")
    set_font(r, size_pt=13, italic=True)

    # Khoảng cách sau header
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(6)
    p_sp.paragraph_format.space_after = Pt(6)

    # ==========================
    # 2. TIÊU ĐỀ KẾ HOẠCH
    # ==========================
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(3)
    r = p_title.add_run("KẾ HOẠCH")
    set_font(r, size_pt=14.5, bold=True)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(2)
    r = p_sub.add_run("Tổ chức dạy học trực tuyến hỗ trợ dạy học trực tiếp\nvà kiểm tra, đánh giá trực tuyến năm học 2026 - 2027")
    set_font(r, size_pt=13.5, bold=True)

    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(0)
    p_sub.paragraph_format.space_after = Pt(0)
    add_vector_line(doc, 45, space_before=2, space_after=6)

    # ==========================
    # 3. CĂN CỨ PHÁP LÝ
    # ==========================
    cancu_items = [
        "Căn cứ Thông tư số 09/2021/TT-BGDĐT ngày 30 tháng 3 năm 2021 của Bộ trưởng Bộ Giáo dục và Đào tạo quy định về quản lý và tổ chức dạy học trực tuyến trong cơ sở giáo dục phổ thông và cơ sở giáo dục thường xuyên;",
        "Căn cứ Thông tư số 22/2021/TT-BGDĐT ngày 20 tháng 7 năm 2021 của Bộ trưởng Bộ Giáo dục và Đào tạo quy định về đánh giá học sinh trung học cơ sở và học sinh trung học phổ thông;",
        "Căn cứ Quyết định số 4725/QĐ-BGDĐT ngày 30 tháng 12 năm 2022 của Bộ trưởng Bộ Giáo dục và Đào tạo ban hành Bộ chỉ số đánh giá mức độ chuyển đổi số của cơ sở giáo dục phổ thông và giáo dục thường xuyên;",
        "Căn cứ Thông tư số 15/2026/TT-BGDĐT ngày 24 tháng 3 năm 2026 của Bộ trưởng Bộ Giáo dục và Đào tạo ban hành Điều lệ trường tiểu học, trường trung học cơ sở, trường trung học phổ thông và trường phổ thông có nhiều cấp học;",
        "Căn cứ Kế hoạch số 5264/KH-SGDĐT ngày 22 tháng 9 năm 2026 của Sở Giáo dục và Đào tạo tỉnh Đồng Nai về việc thực hiện nhiệm vụ ứng dụng công nghệ thông tin, chuyển đổi số và thống kê giáo dục năm học 2026 - 2027;",
        "Căn cứ Kế hoạch số 101/KH-UBND ngày 21 tháng 4 năm 2026 của Ủy ban nhân dân xã Xuân Đông về chuyển đổi số trong các cơ quan nhà nước xã Xuân Đông năm 2026;",
        "Căn cứ tình hình thực tế và điều kiện cơ sở vật chất của Trường THCS Trần Phú,"
    ]
    for cc in cancu_items:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0.5)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(cc)
        set_font(r, size_pt=13, italic=True)

    # Lời dẫn
    p_lead = doc.add_paragraph()
    p_lead.paragraph_format.first_line_indent = Inches(0.5)
    p_lead.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_lead.paragraph_format.space_before = Pt(4)
    p_lead.paragraph_format.space_after = Pt(4)
    p_lead.paragraph_format.line_spacing = 1.2
    r = p_lead.add_run("Trường THCS Trần Phú xây dựng Kế hoạch tổ chức dạy học trực tuyến hỗ trợ dạy học trực tiếp và kiểm tra, đánh giá trực tuyến năm học 2026 - 2027 cụ thể như sau:")
    set_font(r, size_pt=13)

    # Helpers
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0.5)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(text)
        set_font(r, size_pt=13, bold=True)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0.5)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(text)
        set_font(r, size_pt=13, bold=True)
        return p

    def add_bullet(text):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0.5)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(f"- {text}")
        set_font(r, size_pt=13)
        return p

    def add_step(step_name, items):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0.5)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(step_name)
        set_font(r, size_pt=13, bold=True)
        for it in items:
            add_bullet(it)

    # ==========================
    # I. MỤC ĐÍCH, YÊU CẦU
    # ==========================
    add_h1("I. MỤC ĐÍCH, YÊU CẦU")
    add_h2("1. Mục đích")
    add_bullet("Tổ chức dạy học trực tuyến linh hoạt, hiệu quả, phù hợp điều kiện thực tế; hỗ trợ dạy học trực tiếp, bảo đảm thực hiện chương trình giáo dục phổ thông và kế hoạch giáo dục của nhà trường.")
    add_bullet("Khai thác eNetViet của Quảng Ích là kênh thông tin chính thức giữa nhà trường, giáo viên, học sinh và cha mẹ học sinh; sử dụng các chức năng được nhà trường cấp quyền để hỗ trợ quản lý, giao tiếp, theo dõi việc học.")
    add_bullet("Phát triển học liệu số, nâng cao năng lực số của cán bộ quản lý, giáo viên, nhân viên và học sinh; từng bước hoàn thiện minh chứng chuyển đổi số theo Quyết định số 4725/QĐ-BGDĐT.")
    add_bullet("Bảo đảm an toàn thông tin, bảo vệ dữ liệu cá nhân, tính trung thực trong học tập và kiểm tra, đánh giá.")

    add_h2("2. Yêu cầu")
    add_bullet("Dạy học trực tuyến được tổ chức theo kế hoạch giáo dục của nhà trường, bảo đảm yêu cầu cần đạt của chương trình; không gây quá tải cho học sinh và giáo viên.")
    add_bullet("eNetViet được sử dụng thống nhất là kênh thông báo, phối hợp với cha mẹ học sinh và các chức năng chuyên môn do nhà trường cấu hình. Không sử dụng tài khoản cá nhân để thay thế kênh thông tin chính thức.")
    add_bullet("Việc kiểm tra, đánh giá trực tuyến chỉ thực hiện khi bảo đảm điều kiện kỹ thuật, công bằng, xác thực người học và phù hợp quy định hiện hành; ưu tiên đánh giá thường xuyên.")
    add_bullet("Có phương án hỗ trợ học sinh khó khăn về thiết bị, đường truyền; không để điều kiện công nghệ trở thành rào cản đối với việc học.")

    # ==========================
    # II. ĐẶC ĐIỂM TÌNH HÌNH
    # ==========================
    add_h1("II. ĐẶC ĐIỂM TÌNH HÌNH")
    add_h2("1. Thuận lợi")
    add_bullet("Nhà trường triển khai eNetViet của Quảng Ích, tạo kênh kết nối chính thức với phụ huynh, giáo viên và học sinh.")
    add_bullet("Đội ngũ cán bộ, giáo viên có khả năng sử dụng các thiết bị số cơ bản, sẵn sàng bồi dưỡng và chia sẻ kinh nghiệm ứng dụng công nghệ trong dạy học.")
    add_bullet("Nhà trường có hệ thống Internet, phòng máy và thiết bị trình chiếu phục vụ dạy học, kiểm tra, đánh giá và phát triển học liệu số.")

    add_h2("2. Khó khăn")
    add_bullet("Một số học sinh còn hạn chế về thiết bị hoặc kết nối Internet tại nhà; kỹ năng tự học trên môi trường số chưa đồng đều.")
    add_bullet("Giáo viên cần thêm thời gian để thiết kế học liệu số, tổ chức hoạt động học tập trực tuyến và kiểm tra, đánh giá bảo đảm thực chất.")
    add_bullet("Nguy cơ lộ lọt dữ liệu cá nhân, sử dụng công nghệ không đúng mục đích và gian lận trong kiểm tra trực tuyến cần được chủ động phòng ngừa.")

    # ==========================
    # III. NỘI DUNG VÀ NHIỆM VỤ TRỌNG TÂM
    # ==========================
    add_h1("III. NỘI DUNG VÀ NHIỆM VỤ TRỌNG TÂM")
    add_h2("1. Quản lý và khai thác eNetViet")
    add_bullet("Rà soát, cập nhật chính xác tài khoản, thông tin liên hệ của cán bộ, giáo viên, nhân viên, học sinh và cha mẹ học sinh theo phân quyền.")
    add_bullet("Sử dụng eNetViet để thông báo kế hoạch giáo dục, lịch học, lịch kiểm tra, kết quả học tập theo thẩm quyền; trao đổi kịp thời với cha mẹ học sinh.")
    add_bullet("Hướng dẫn người dùng bảo mật tài khoản, đổi mật khẩu khi cần thiết; không chia sẻ thông tin đăng nhập, không đăng tải dữ liệu cá nhân nhạy cảm ngoài phạm vi được phép.")

    add_h2("2. Tổ chức dạy học trực tuyến hỗ trợ dạy học trực tiếp")
    add_bullet("Giáo viên chủ động sử dụng học liệu số, giao nhiệm vụ tự học, hướng dẫn học sinh chuẩn bị bài và củng cố sau giờ học thông qua các nền tảng được nhà trường cho phép.")
    add_bullet("Khi tổ chức dạy học trực tuyến, giáo viên xây dựng kế hoạch bài dạy, xác định mục tiêu, nội dung, thời lượng, học liệu, phương thức tương tác và cách kiểm tra mức độ hoàn thành của học sinh.")
    add_bullet("Các tổ chuyên môn xây dựng, thẩm định và cập nhật kho học liệu số; bảo đảm học liệu đúng chương trình, phù hợp lứa tuổi, tôn trọng bản quyền và có nguồn gốc rõ ràng.")
    add_bullet("Trong trường hợp học sinh không thể tham gia trực tuyến vì điều kiện khách quan, giáo viên chủ nhiệm và giáo viên bộ môn phối hợp thực hiện phương án hỗ trợ phù hợp.")

    add_h2("3. Kiểm tra, đánh giá trực tuyến")
    add_bullet("Thực hiện kiểm tra, đánh giá trực tuyến chủ yếu đối với đánh giá thường xuyên, khi giáo viên bảo đảm được điều kiện kỹ thuật, tính công bằng và khả năng xác thực kết quả học tập.")
    add_bullet("Đề kiểm tra trực tuyến phải bám sát yêu cầu cần đạt; có ma trận, bản đặc tả hoặc hướng dẫn chấm theo quy định; được tổ chuyên môn thống nhất trước khi sử dụng đối với các bài kiểm tra chung.")
    add_bullet("Kết quả do hệ thống tự chấm là dữ liệu tham khảo; giáo viên chịu trách nhiệm kiểm tra, xác nhận kết quả trước khi sử dụng trong đánh giá học sinh.")
    add_bullet("Không quy định thực hiện toàn bộ bài kiểm tra thường xuyên bằng hình thức trực tuyến; giáo viên lựa chọn hình thức phù hợp đặc điểm môn học và điều kiện của học sinh.")

    add_h2("4. Bảo đảm an toàn thông tin và hỗ trợ người học")
    add_bullet("Không chia sẻ công khai danh sách lớp, số điện thoại, hình ảnh, thông tin sức khỏe, kết quả học tập chi tiết hoặc dữ liệu cá nhân của học sinh trên các nền tảng không được phép.")
    add_bullet("Tổ Công nghệ thông tin tiếp nhận, hỗ trợ xử lý sự cố tài khoản, truy cập, thiết bị và đường truyền; kịp thời báo cáo Ban Giám hiệu các sự cố có nguy cơ ảnh hưởng dữ liệu hoặc hoạt động dạy học.")
    add_bullet("Rà soát học sinh thiếu thiết bị, kết nối; hướng dẫn sử dụng phòng máy, cho mượn hoặc bố trí phương án học tập thay thế trong điều kiện cho phép.")

    # ==========================
    # IV. LỘ TRÌNH THỰC HIỆN
    # ==========================
    add_h1("IV. LỘ TRÌNH THỰC HIỆN")
    p_tbl_intro = doc.add_paragraph()
    p_tbl_intro.paragraph_format.first_line_indent = Inches(0.5)
    p_tbl_intro.paragraph_format.space_before = Pt(2)
    p_tbl_intro.paragraph_format.space_after = Pt(4)
    r = p_tbl_intro.add_run("Lộ trình triển khai thực hiện nhiệm vụ trong năm học 2026 - 2027 cụ thể như sau:")
    set_font(r, size_pt=13)

    table_data = [
        ["Giai đoạn", "Thời gian", "Nội dung trọng tâm", "Sản phẩm / Kết quả"],
        ["1", "09 - 10/2026", "Rà soát tài khoản eNetViet; phân quyền; cập nhật dữ liệu; tập huấn sử dụng, bảo mật tài khoản và quy trình hỗ trợ kỹ thuật.", "Danh sách tài khoản; biên bản tập huấn; hướng dẫn sử dụng."],
        ["2", "11 - 12/2026", "Tổ chức dạy học trực tuyến hỗ trợ dạy trực tiếp; xây dựng, thẩm định học liệu số; thực hiện kiểm tra thường xuyên phù hợp.", "Kho học liệu số; minh chứng hoạt động dạy học; báo cáo học kỳ I."],
        ["3", "01 - 03/2027", "Bồi dưỡng chuyên môn, chia sẻ kinh nghiệm; củng cố việc giao nhiệm vụ, phản hồi học tập; hỗ trợ học sinh khó khăn về thiết bị, kết nối.", "Chuyên đề/tập huấn; danh sách hỗ trợ; báo cáo tiến độ."],
        ["4", "04 - 05/2027", "Rà soát chất lượng học liệu, kết quả triển khai; đánh giá việc bảo đảm an toàn thông tin, tính thực chất của kiểm tra, đánh giá.", "Báo cáo tổng kết; danh mục học liệu; đề xuất cải tiến."],
        ["5", "06 - 08/2027", "Lưu trữ hồ sơ, làm sạch dữ liệu cần thiết; chuẩn bị tài khoản và kế hoạch cho năm học tiếp theo.", "Hồ sơ lưu trữ; kế hoạch dự kiến năm học 2027 - 2028."]
    ]

    t_plan = doc.add_table(rows=len(table_data), cols=4)
    set_table_borders(t_plan, color="000000", sz="4")
    t_plan.alignment = WD_TABLE_ALIGNMENT.CENTER

    widths = [Mm(20), Mm(28), Mm(67), Mm(50)]
    for r_idx, row in enumerate(t_plan.rows):
        set_cant_split(row)
        if r_idx == 0:
            set_tbl_header(row)
        for c_idx, cell in enumerate(row.cells):
            cell.width = widths[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            cell_text = table_data[r_idx][c_idx]
            r = p.add_run(cell_text)
            if r_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_font(r, size_pt=11.5, bold=True)
            else:
                if c_idx in [0, 1]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                set_font(r, size_pt=11.5, bold=False)

    p_sp3 = doc.add_paragraph()
    p_sp3.paragraph_format.space_before = Pt(4)
    p_sp3.paragraph_format.space_after = Pt(4)

    # ==========================
    # V. PHÂN CÔNG NHIỆM VỤ
    # ==========================
    add_h1("V. PHÂN CÔNG NHIỆM VỤ")
    add_h2("1. Hiệu trưởng")
    add_bullet("Chỉ đạo chung, phê duyệt kế hoạch và quyết định các biện pháp tổ chức thực hiện.")
    add_bullet("Chỉ đạo kiểm tra, đánh giá kết quả thực hiện; giải quyết các vấn đề vượt thẩm quyền của các bộ phận.")

    add_h2("2. Phó Hiệu trưởng phụ trách chuyên môn")
    add_bullet("Theo dõi việc tổ chức dạy học trực tuyến, kiểm tra, đánh giá trực tuyến của các tổ chuyên môn và giáo viên.")
    add_bullet("Chỉ đạo thẩm định học liệu số, ngân hàng câu hỏi và việc thực hiện kế hoạch giáo dục.")

    add_h2("3. Tổ Công nghệ thông tin")
    add_bullet("Quản trị kỹ thuật theo phân công; hỗ trợ tài khoản, phân quyền và hướng dẫn sử dụng eNetViet.")
    add_bullet("Theo dõi sự cố kỹ thuật, tham mưu giải pháp bảo đảm an toàn thông tin, hạ tầng mạng và thiết bị.")
    add_bullet("Tổng hợp số liệu triển khai phục vụ công tác quản lý, báo cáo.")

    add_h2("4. Tổ chuyên môn")
    add_bullet("Xây dựng kế hoạch triển khai của tổ; chỉ đạo giáo viên thiết kế học liệu số, nhiệm vụ học tập phù hợp.")
    add_bullet("Thẩm định nội dung học liệu, ngân hàng câu hỏi; tổ chức sinh hoạt chuyên môn, rút kinh nghiệm.")

    add_h2("5. Giáo viên bộ môn")
    add_bullet("Tổ chức dạy học, giao nhiệm vụ, phản hồi và đánh giá học sinh theo kế hoạch; bảo đảm chất lượng, tính khoa học và an toàn thông tin.")
    add_bullet("Theo dõi mức độ tham gia của học sinh; phối hợp giáo viên chủ nhiệm hỗ trợ học sinh gặp khó khăn.")
    add_bullet("Không sử dụng kết quả tự chấm hoặc dữ liệu hệ thống một cách máy móc; chịu trách nhiệm xác nhận kết quả đánh giá.")

    add_h2("6. Giáo viên chủ nhiệm và Tổng phụ trách Đội")
    add_bullet("Hướng dẫn, tuyên truyền học sinh và cha mẹ học sinh sử dụng eNetViet đúng mục đích, an toàn, văn minh.")
    add_bullet("Nắm bắt khó khăn của học sinh; phối hợp với nhà trường hỗ trợ và nhắc nhở việc thực hiện nhiệm vụ học tập.")

    add_h2("7. Học sinh và cha mẹ học sinh")
    add_bullet("Sử dụng tài khoản đúng mục đích, bảo mật thông tin đăng nhập; thực hiện nghiêm túc nhiệm vụ học tập và kiểm tra, đánh giá.")
    add_bullet("Kịp thời phản ánh với giáo viên chủ nhiệm hoặc nhà trường khi có sự cố về tài khoản, thiết bị, đường truyền hoặc dấu hiệu mất an toàn thông tin.")

    # ==========================
    # VI. QUY TRÌNH TỔ CHỨC KIỂM TRA, ĐÁNH GIÁ TRỰC TUYẾN
    # ==========================
    add_h1("VI. QUY TRÌNH TỔ CHỨC KIỂM TRA, ĐÁNH GIÁ TRỰC TUYẾN")
    add_step("Bước 1. Chuẩn bị", [
        "Giáo viên xác định mục tiêu, hình thức, thời gian, yêu cầu kỹ thuật; chuẩn bị đề, đáp án, hướng dẫn chấm và phương án xử lý sự cố.",
        "Đối với bài kiểm tra chung, tổ chuyên môn tổ chức thẩm định theo quy định."
    ])
    add_step("Bước 2. Thông báo và hướng dẫn", [
        "Thông báo cho học sinh tối thiểu 02 ngày trước thời điểm thực hiện; hướng dẫn kiểm tra thiết bị, đường truyền và tài khoản.",
        "Thông báo rõ quy định về trung thực học tập, thời gian làm bài và cách xử lý khi xảy ra sự cố kỹ thuật."
    ])
    add_step("Bước 3. Tổ chức thực hiện", [
        "Tổ chức theo thời gian đã thông báo; áp dụng biện pháp phù hợp để hạn chế gian lận, bảo đảm quyền lợi học sinh khi gặp sự cố khách quan.",
        "Giáo viên theo dõi quá trình thực hiện trong phạm vi chức năng của nền tảng được nhà trường cho phép."
    ])
    add_step("Bước 4. Xử lý kết quả", [
        "Giáo viên kiểm tra, xác nhận kết quả; chấm phần tự luận, nhận xét và phản hồi cho học sinh theo quy định.",
        "Nếu có dấu hiệu bất thường, giáo viên trao đổi với học sinh, báo cáo tổ chuyên môn và Ban Giám hiệu để xem xét, không tự động kết luận vi phạm chỉ dựa trên dữ liệu kỹ thuật."
    ])

    # ==========================
    # VII. KIỂM TRA, GIÁM SÁT VÀ CHẾ ĐỘ BÁO CÁO
    # ==========================
    add_h1("VII. KIỂM TRA, GIÁM SÁT VÀ CHẾ ĐỘ BÁO CÁO")
    add_h2("1. Kiểm tra, giám sát")
    add_bullet("Ban Giám hiệu kiểm tra định kỳ hoặc đột xuất việc thực hiện kế hoạch; tập trung vào chất lượng học liệu, việc giao nhiệm vụ, phản hồi học sinh, an toàn dữ liệu và tính thực chất của kiểm tra, đánh giá.")
    add_bullet("Tổ chuyên môn kiểm tra nội bộ hoạt động của giáo viên theo nhiệm vụ được phân công; kịp thời hỗ trợ, chấn chỉnh các hạn chế.")

    add_h2("2. Chế độ báo cáo")
    add_bullet("Tổ chuyên môn báo cáo định kỳ cuối mỗi học kỳ và báo cáo đột xuất khi có yêu cầu hoặc sự cố đáng kể.")
    add_bullet("Tổ Công nghệ thông tin tổng hợp số liệu liên quan đến tài khoản, hỗ trợ kỹ thuật, hạ tầng và sự cố an toàn thông tin để báo cáo Ban Giám hiệu.")
    add_bullet("Cuối năm học, nhà trường tổ chức đánh giá, rút kinh nghiệm, cập nhật kế hoạch và đề xuất giải pháp cho năm học tiếp theo.")

    # ==========================
    # VIII. TỔ CHỨC THỰC HIỆN
    # ==========================
    add_h1("VIII. TỔ CHỨC THỰC HIỆN")
    add_bullet("Các bộ phận, cá nhân được phân công nghiêm túc triển khai thực hiện Kế hoạch này.")
    add_bullet("Trong quá trình thực hiện, nếu có khó khăn, vướng mắc, các bộ phận kịp thời báo cáo Ban Giám hiệu để xem xét, chỉ đạo giải quyết.")
    add_bullet("Kế hoạch này được phổ biến đến toàn thể cán bộ, giáo viên, nhân viên, học sinh và cha mẹ học sinh để phối hợp thực hiện.")

    p_sp4 = doc.add_paragraph()
    p_sp4.paragraph_format.space_before = Pt(8)
    p_sp4.paragraph_format.space_after = Pt(4)

    # ==========================
    # NƠI NHẬN & CHỮ KÝ (BẢNG 2 CỘT CHUẨN NĐ 30)
    # ==========================
    footer_table = doc.add_table(rows=1, cols=2)
    remove_table_borders(footer_table)
    footer_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    footer_table.rows[0].cells[0].width = Mm(65)
    footer_table.rows[0].cells[1].width = Mm(100)

    # Cột Trái: Nơi nhận
    cell_nn = footer_table.cell(0, 0)
    p_nn_head = cell_nn.paragraphs[0]
    p_nn_head.paragraph_format.space_before = Pt(0)
    p_nn_head.paragraph_format.space_after = Pt(2)
    p_nn_head.paragraph_format.line_spacing = 1.15
    r = p_nn_head.add_run("Nơi nhận:")
    set_font(r, size_pt=12, bold=True, italic=True)

    nn_items = [
        "Ban Giám hiệu;",
        "Các tổ chuyên môn;",
        "Cán bộ, giáo viên, nhân viên;",
        "Lưu: VT."
    ]
    for it in nn_items:
        p_it = cell_nn.add_paragraph()
        p_it.paragraph_format.space_before = Pt(0)
        p_it.paragraph_format.space_after = Pt(1)
        p_it.paragraph_format.line_spacing = 1.1
        r = p_it.add_run(f"- {it}")
        set_font(r, size_pt=11, bold=False)

    # Cột Phải: Chữ ký Hiệu trưởng
    cell_sign = footer_table.cell(0, 1)
    p_sign_role = cell_sign.paragraphs[0]
    p_sign_role.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sign_role.paragraph_format.space_before = Pt(0)
    p_sign_role.paragraph_format.space_after = Pt(2)
    p_sign_role.paragraph_format.line_spacing = 1.15
    r = p_sign_role.add_run("HIỆU TRƯỞNG")
    set_font(r, size_pt=13, bold=True)

    p_sign_note = cell_sign.add_paragraph()
    p_sign_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sign_note.paragraph_format.space_before = Pt(0)
    p_sign_note.paragraph_format.space_after = Pt(45) # Khoảng trống ký tên đóng dấu
    p_sign_note.paragraph_format.line_spacing = 1.15
    r = p_sign_note.add_run("(Ký, đóng dấu và ghi rõ họ tên)")
    set_font(r, size_pt=11, italic=True)

    p_sign_name = cell_sign.add_paragraph()
    p_sign_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sign_name.paragraph_format.space_before = Pt(0)
    p_sign_name.paragraph_format.space_after = Pt(0)
    p_sign_name.paragraph_format.line_spacing = 1.15
    r = p_sign_name.add_run("Bùi Ngọc Nam")
    set_font(r, size_pt=13, bold=True)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print("SUCCESS: Standardized docx generated at:", output_path)

if __name__ == '__main__':
    out_file = r"C:\Users\HoangThien\Documents\GitHub\giangbai\TROLYTHIEN\12_CHUAN_HOA_VAN_BAN\Ket_qua\KH_TO_CHUC_DAY_HOC_TRUC_TUYEN_2026_2027_TRANPHU_Chuan_Hoa.docx"
    standardize_document(out_file)
