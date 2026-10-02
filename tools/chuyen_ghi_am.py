import os
import sys
import datetime
from docx import Document
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_font(run, name='Times New Roman', size_pt=13, bold=False, italic=False, color_rgb=(0,0,0)):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = RGBColor(*color_rgb)
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{name}" w:hAnsi="{name}" w:cs="{name}" w:eastAsia="{name}"/>')
    rPr.append(rFonts)

def export_transcript_to_docx_and_txt(source_file_name, transcript_text, output_dir):
    """
    Xuất nội dung bóc tách âm thanh sang tệp Word .docx chuẩn in ấn và .txt.
    """
    os.makedirs(output_dir, exist_ok=True)
    base_name = os.path.splitext(source_file_name)[0]
    docx_path = os.path.join(output_dir, f"{base_name}.docx")
    txt_path = os.path.join(output_dir, f"{base_name}.txt")

    # 1. Ghi tệp .txt UTF-8
    with open(txt_path, 'w', encoding='utf-8') as f:
        f.write(transcript_text.strip() + "\n")

    # 2. Tạo tệp Word .docx chuẩn A4, Times New Roman 13pt, dãn dòng 1.2, thụt đầu dòng 1.27cm
    doc = Document()
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(30)
        section.right_margin = Mm(15)

    # Tiêu đề tài liệu
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r = p_title.add_run("BẢN BÓC TÁCH NỘI DUNG GHI ÂM")
    set_font(r, size_pt=14, bold=True)

    # Thông tin nguồn
    now_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_before = Pt(0)
    p_meta.paragraph_format.space_after = Pt(12)
    r_meta = p_meta.add_run(f"(Nguồn: {source_file_name} — Thời gian xử lý: {now_str})")
    set_font(r_meta, size_pt=11.5, italic=True)

    # Thân bài: từng đoạn phát biểu
    paragraphs = [p.strip() for p in transcript_text.split('\n') if p.strip()]
    for p_text in paragraphs:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0.5) # 1.27 cm
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)

        # Kiểm tra nếu đoạn bắt đầu bằng nhãn người nói (ví dụ "Chủ trì:", "Thầy Nam:", "[01:15]")
        if any(p_text.startswith(prefix) for prefix in ["Chủ trì:", "Thầy", "Cô", "Người nói", "["]):
            parts = p_text.split(':', 1)
            if len(parts) == 2:
                r_speaker = p.add_run(parts[0] + ": ")
                set_font(r_speaker, size_pt=13, bold=True)
                r_content = p.add_run(parts[1].strip())
                set_font(r_content, size_pt=13, bold=False)
            else:
                r_body = p.add_run(p_text)
                set_font(r_body, size_pt=13, bold=False)
        else:
            r_body = p.add_run(p_text)
            set_font(r_body, size_pt=13, bold=False)

    doc.save(docx_path)
    return docx_path, txt_path

if __name__ == '__main__':
    if len(sys.argv) >= 3:
        src = sys.argv[1]
        text_content = sys.argv[2]
        out_d = sys.argv[3] if len(sys.argv) >= 4 else "TROLYTHIEN/13_CHUYEN_GHI_AM/Ket_qua"
        d_path, t_path = export_transcript_to_docx_and_txt(src, text_content, out_d)
        print(f"Exported docx: {d_path}")
        print(f"Exported txt: {t_path}")
    else:
        print("Usage: python chuyen_ghi_am.py <source_name> <text> [output_dir]")
