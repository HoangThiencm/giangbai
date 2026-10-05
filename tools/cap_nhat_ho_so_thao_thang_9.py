import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="1F2937", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def safe_save_doc(doc, target_path):
    try:
        doc.save(target_path)
        print("Saved successfully:", target_path)
        return target_path
    except PermissionError:
        base, ext = os.path.splitext(target_path)
        fallback_path = f"{base}_CapNhat{ext}"
        try:
            doc.save(fallback_path)
            print(f"File locked by Word. Saved to fallback: {fallback_path}")
            return fallback_path
        except Exception as e:
            print(f"Error saving docx: {e}")
            return None

def apply_base_page_setup(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.59)
        s.bottom_margin = Inches(0.59)
        s.left_margin = Inches(0.79)
        s.right_margin = Inches(0.59)
        s.different_first_page_header_footer = False
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(13)
    normal_style.paragraph_format.space_before = Pt(0)
    normal_style.paragraph_format.space_after = Pt(3)
    normal_style.paragraph_format.line_spacing = 1.0

def build_header_national(doc, so_bb="... /BB-TT"):
    tbl_hdr = doc.add_table(rows=1, cols=2)
    tbl_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left, c_right = tbl_hdr.rows[0].cells
    c_left.width = Inches(3.2)
    c_right.width = Inches(3.8)

    p_l = c_left.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_l1 = p_l.add_run("TRƯỜNG THCS TRẦN PHÚ\n")
    r_l1.font.bold = True
    r_l1.font.size = Pt(12)
    r_l2 = p_l.add_run("TỔ TOÁN – TIN\n")
    r_l2.font.bold = True
    r_l2.font.size = Pt(12)
    r_l3 = p_l.add_run(f"Số: {so_bb}")
    r_l3.font.italic = True
    r_l3.font.size = Pt(11)

    p_r = c_right.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_r1 = p_r.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
    r_r1.font.bold = True
    r_r1.font.size = Pt(12)
    r_r2 = p_r.add_run("Độc lập - Tự do - Hạnh phúc\n")
    r_r2.font.bold = True
    r_r2.font.size = Pt(12)
    r_r3 = p_r.add_run("Xuân Đông, ngày 30 tháng 9 năm 2026")
    r_r3.font.italic = True
    r_r3.font.size = Pt(11)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(6)

def add_signatures(doc):
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(14)
    tbl_sign = doc.add_table(rows=1, cols=2)
    tbl_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_left, s_right = tbl_sign.rows[0].cells
    s_left.width = Inches(3.5)
    s_right.width = Inches(3.5)

    p_sl = s_left.paragraphs[0]
    p_sl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sl.add_run("TỔ TRƯỞNG CHUYÊN MÔN\n").bold = True
    p_sl.add_run("(Ký và ghi rõ họ tên)\n\n\n\n\n")
    p_sl.add_run("....................................................")

    p_sr = s_right.paragraphs[0]
    p_sr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sr.add_run("PHÊ DUYỆT CỦA BAN GIÁM HIỆU\n").bold = True
    p_sr.add_run("HIỆU TRƯỞNG / PHÓ HIỆU TRƯỞNG\n\n\n\n\n")
    p_sr.add_run("....................................................")

# ==============================================================================
# 1. CẬP NHẬT BIÊN BẢN HÀNH CHÍNH (THÊM CÔ THẢO VÀO STT 3)
# ==============================================================================
def update_formal_minutes_with_thao():
    doc = docx.Document()
    apply_base_page_setup(doc)
    build_header_national(doc, so_bb="01 /BB-TT")

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("BIÊN BẢN KIỂM TRA HỒ SƠ GIÁO ÁN TỔ CHUYÊN MÔN\n")
    r_t.font.bold = True
    r_t.font.size = Pt(14)
    r_sub = p_title.add_run("ĐỢT THÁNG 9/2026 — NĂM HỌC 2026 - 2027")
    r_sub.font.bold = True
    r_sub.font.size = Pt(13)

    # Section I
    p_sec1 = doc.add_paragraph()
    p_sec1.add_run("I. THÀNH PHẦN KIỂM TRA\n").bold = True
    p_sec1.add_run("1. Tổ trưởng chuyên môn: ").bold = True
    p_sec1.add_run("Trưởng ban kiểm tra.\n")
    p_sec1.add_run("2. Tổ phó chuyên môn: ").bold = True
    p_sec1.add_run("Phó ban kiểm tra.\n")
    p_sec1.add_run("3. Các thành viên: ").bold = True
    p_sec1.add_run("Toàn thể giáo viên tổ Toán – Tin.")

    # Section II
    p_sec2 = doc.add_paragraph()
    p_sec2.add_run("II. CĂN CỨ VÀ TIÊU CHÍ ĐÁNH GIÁ CHUYÊN MÔN\n").bold = True
    p_sec2.add_run("1. Căn cứ Phụ lục 4 Kế hoạch bài dạy (Công văn 5512/BGDĐT-GDTrH) và Quy định chuyên môn trường THCS Trần Phú:\n")
    p_sec2.add_run("   • Hình thức thể thức: Font Times New Roman 13pt (in đứng); Căn lề Trên 1.5cm, Dưới 1.5cm, Trái 2.0cm, Phải 1.5cm; Giãn dòng: Before 0pt, After 3pt, Single.\n")
    p_sec2.add_run("   • Header bắt buộc: Trường THCS Trần Phú | Giáo viên: [Họ và tên]; Footer bắt buộc: Môn/Phân môn | Trang | Năm học.\n")
    p_sec2.add_run("2. Tiến độ thực hiện: Bắt buộc khớp từng phân môn và tuần dạy đã duyệt tại Phụ lục 1 và Phụ lục 3 (Kế hoạch giáo dục của giáo viên).\n")
    p_sec2.add_run("3. Tích hợp NLS / AI: Khớp 1-1 cả Mã và Câu mô tả với Phụ lục 3; trong bài dạy bắt buộc phải in đậm, nghiêng.\n")
    p_sec2.add_run("4. Kiến thức toán học & Ký hiệu: Chuẩn xác tuyệt đối về dấu mũ góc, dấu toán học (+/-), biến đổi đại số.\n")

    # Section III: Bảng tổng hợp toàn tổ (3 GV)
    p_sec3 = doc.add_paragraph()
    p_sec3.add_run("III. BẢNG TỔNG HỢP TIẾN ĐỘ VÀ XẾP LOẠI TOÀN TỔ").bold = True

    t_sum = doc.add_table(rows=4, cols=6)
    t_sum.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_sum, color="1F2937", sz="4")

    headers = [
        "STT", "Họ và tên GV", "Môn / Khối lớp", "Tổng số tiết đã nộp",
        "Tiến độ PPCT Tháng 9", "Kết luận / Xếp loại"
    ]
    hdr_cells = t_sum.rows[0].cells
    col_widths = [Inches(0.5), Inches(1.7), Inches(1.3), Inches(1.1), Inches(1.4), Inches(1.0)]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(10)
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "E2E8F0")
        set_cell_margins(hdr_cells[i], top=70, bottom=70, left=70, right=70)

    # Row 1: Thầy Danh
    r1_cells = t_sum.rows[1].cells
    r1_data = [
        "1", "HỒ ĐĂNG DANH", "Toán 6, Toán 9", "33 tiết (4 phân môn)",
        "Đạt (Có giải trình lý do chính đáng)", "DUYỆT\n(Xếp loại: Tốt)"
    ]
    for i, d in enumerate(r1_data):
        r1_cells[i].text = d
        p = r1_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in [0, 4, 5] else WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.size = Pt(9.5)
            if i == 5:
                run.font.bold = True
                run.font.color.rgb = RGBColor(22, 101, 52)
            elif i == 1:
                run.font.bold = True
            elif i == 4:
                run.font.color.rgb = RGBColor(22, 101, 52)
        r1_cells[i].width = col_widths[i]
        set_cell_margins(r1_cells[i], top=60, bottom=60, left=60, right=60)

    # Row 2: Thầy Hải
    r2_cells = t_sum.rows[2].cells
    r2_data = [
        "2", "TRẦN LONG HẢI", "Toán 8", "20 tiết (2 phân môn)",
        "Đủ 100% Tháng 9 & vượt Tuần 5", "TRẢ HỒ SƠ\n(Sửa sai sót toán)"
    ]
    for i, d in enumerate(r2_data):
        r2_cells[i].text = d
        p = r2_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in [0, 4, 5] else WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.size = Pt(9.5)
            if i == 5:
                run.font.bold = True
                run.font.color.rgb = RGBColor(185, 28, 28)
            elif i == 1:
                run.font.bold = True
        r2_cells[i].width = col_widths[i]
        set_cell_margins(r2_cells[i], top=60, bottom=60, left=60, right=60)

    # Row 3: Cô Thảo
    r3_cells = t_sum.rows[3].cells
    r3_data = [
        "3", "NGUYỄN THỊ THẢO", "Toán 6, Toán 7", "34 tiết (4 phân môn)",
        "Đủ 100% Tháng 9 (12t Số 6) & vượt Tuần 5", "DUYỆT\n(Xếp loại: Khá)"
    ]
    for i, d in enumerate(r3_data):
        r3_cells[i].text = d
        p = r3_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in [0, 4, 5] else WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.size = Pt(9.5)
            if i == 5:
                run.font.bold = True
                run.font.color.rgb = RGBColor(22, 101, 52)
            elif i == 1:
                run.font.bold = True
            elif i == 4:
                run.font.color.rgb = RGBColor(22, 101, 52)
        r3_cells[i].width = col_widths[i]
        set_cell_margins(r3_cells[i], top=60, bottom=60, left=60, right=60)

    # Section IV: Đánh giá chi tiết từng giáo viên
    p_sec4 = doc.add_paragraph()
    p_sec4.paragraph_format.space_before = Pt(12)
    p_sec4.add_run("IV. ĐÁNH GIÁ VÀ NHẬN XÉT CỤ THỂ TỪNG GIÁO VIÊN\n").bold = True

    # 1. Thầy Danh
    p_d = doc.add_paragraph()
    p_d.add_run("1. Giáo viên: HỒ ĐĂNG DANH (Môn Toán 6, Toán 9)\n").bold = True
    p_d.add_run("• Ưu điểm: ").bold = True
    p_d.add_run("Hồ sơ giáo án chuẩn bị chu đáo, nộp đủ 4 phân môn (Đại số 9, Hình học 9, Số học 6, Hình học 6) với tổng số 33 tiết. Thể thức văn bản và Header/Footer đúng quy chuẩn trường THCS Trần Phú. Cấu trúc bài dạy đảm bảo đủ 4 hoạt động theo CV 5512.\n")
    p_d.add_run("• Về tiến độ: ").bold = True
    p_d.add_run("Đối với 02 tiết còn thiếu của Tuần 4 (Tiết 4 Bài 12 Hình 9 và Tiết 12 Số 6), tổ chuyên môn chấp thuận giải trình có lý do chính đáng và kế hoạch dạy bù trong tuần tiếp theo của thầy, thống nhất nghiệm thu tiến độ đạt yêu cầu.\n")
    p_d.add_run("• Tồn tại cần lưu ý: ").bold = True
    p_d.add_run("Các câu mô tả mục tiêu chỉ báo NLS và AI tại Bài 1 (Đại số 9), Bài 6 và Bài 7 (Số học 6) chưa được định dạng in đậm, nghiêng theo quy định chuyên môn.\n")
    p_d.add_run("• Kết luận và xếp loại: ").bold = True
    r_d_c = p_d.add_run("DUYỆT HỒ SƠ (Xếp loại: TỐT). ")
    r_d_c.bold = True
    r_d_c.font.color.rgb = RGBColor(22, 101, 52)
    p_d.add_run("Đề nghị thầy hoàn thiện định dạng in đậm nghiêng phần tích hợp NLS/AI trước khi giảng dạy trên lớp.\n")

    # 2. Thầy Hải
    p_h = doc.add_paragraph()
    p_h.add_run("2. Giáo viên: TRẦN LONG HẢI (Môn Toán 8)\n").bold = True
    p_h.add_run("• Ưu điểm: ").bold = True
    p_h.add_run("Hồ sơ giáo án chuẩn bị với khối lượng lớn (20 tiết — 89 trang), tiến độ đảm bảo 100% tháng 9 và vượt tuần 5. Thể thức văn bản, Header/Footer đúng chuẩn quy định trường THCS Trần Phú. Các bài dạy sau (Bài 11, 12, 13 Hình học 8 và các bài Đại số 8) soạn rất công phu, chuẩn mực.\n")
    p_h.add_run("• Tồn tại sai sót chuyên môn toán học nghiêm trọng: ").bold = True
    p_h.add_run("Tại Trang 5 Bài 10 (Hình học 8) phát hiện các sai sót toán học nghiêm trọng: (1) Mất dấu mũ góc ở các công thức (A, B, C, D và H, E, F, G); (2) Sai bản chất định lý tổng các góc trong tứ giác: viết H + E + F - G = 360°; (3) Sai quy tắc chuyển vế và thứ tự phép tính: D = 360° - A + B + C = 50° và F = 360° - H - E + G = 125°; (4) Chưa in đậm, nghiêng mô tả NLS tại Bài 10 (Hình 8) và Bài 3 (Đại 8).\n")
    p_h.add_run("• Kết luận và xếp loại: ").bold = True
    r_h_c = p_h.add_run("ĐỀ NGHỊ TRẢ HỒ SƠ. ")
    r_h_c.bold = True
    r_h_c.font.color.rgb = RGBColor(185, 28, 28)
    p_h.add_run("Yêu cầu thầy Trần Long Hải đính chính chuẩn xác toàn bộ lỗi toán học và ký hiệu tại Trang 5 Bài 10 và in đậm nghiêng các phần NLS trước khi phê duyệt chính thức.\n")

    # 3. Cô Thảo
    p_t = doc.add_paragraph()
    p_t.add_run("3. Giáo viên: NGUYỄN THỊ THẢO (Môn Toán 6, Toán 7)\n").bold = True
    p_t.add_run("• Ưu điểm: ").bold = True
    p_t.add_run("Hồ sơ giáo án chuẩn bị rất đầy đủ, công phu với tổng cộng 34 tiết (123 trang) thuộc 4 phân môn (Số học 6: 12 tiết; Hình học 6: 5 tiết; Đại số 7: 9 tiết; Hình học 7: 8 tiết). Đảm bảo 100% tiến độ tháng 9 và vượt tuần 5 (đặc biệt phân môn Số học 6 nộp đủ 12 tiết đến hết bài Luyện tập chung tuần 4). Thể thức văn bản, Header/Footer đúng mẫu quy định. Có phần hướng dẫn, giao việc riêng cho học sinh hòa nhập (HSHN) rất chu đáo. Nội dung tích hợp Năng lực số tại Bài 6 và Bài 7 Số học 6 được định dạng IN ĐẬM, NGHIÊNG rất chuẩn xác theo đúng quy định chuyên môn.\n")
    p_t.add_run("• Tồn tại cần hoàn thiện: ").bold = True
    p_t.add_run("Tại phân môn Hình học 7 (30 trang), xuất hiện lỗi kỹ thuật trình bày ký hiệu góc có tính hệ thống: thay vì sử dụng dấu mũ góc chuẩn (^), phần mềm hiển thị thành dấu chấm trên đầu chữ cái đầu tiên (ví dụ: ṁBy, ṁBx, ṅBx, ẋBy, ẋAB, ẏBA, ṀNQ, ṄQP...) hoặc viết trần (mBx, ANM, ACD). Tại Trang 8 (Bài 3.3c) bị lỗi font tạo thành các khối đen đè chữ; Trang 9 (Bài 3.5) có dòng viết 'ṁBn = ẋBy = 180°' chưa chuẩn xác về diễn đạt toán học.\n")
    p_t.add_run("• Kết luận và xếp loại: ").bold = True
    r_t_c = p_t.add_run("DUYỆT HỒ SƠ (Xếp loại: KHÁ). ")
    r_t_c.bold = True
    r_t_c.font.color.rgb = RGBColor(22, 101, 52)
    p_t.add_run("Yêu cầu cô Nguyễn Thị Thảo đính chính, chuẩn hóa lại toàn bộ ký hiệu dấu mũ góc trong tệp Hình học 7 để nâng xếp loại Tốt trong đợt kiểm tra tiếp theo.\n")

    # 4. Giáo viên tiếp theo
    p_next = doc.add_paragraph()
    p_next.add_run("4. Giáo viên tiếp theo: ").bold = True
    p_next.add_run("[Đang chờ nạp hồ sơ kiểm tra cuốn chiếu tiếp theo...]\n").italic = True

    add_signatures(doc)

    out_docx = os.path.abspath("TROLYTHIEN/3_DUYET_GIAO_AN/Ket_qua/Bien_Ban_Kiem_Tra_Ho_So_To_Toan_Thang_9.docx")
    safe_save_doc(doc, out_docx)

# ==============================================================================
# 2. CẬP NHẬT TỆP HỆ THỐNG DUYỆT GIÁO ÁN (THÊM CÔ THẢO)
# ==============================================================================
def update_system_comment_doc_with_thao():
    doc = docx.Document()
    apply_base_page_setup(doc)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("TỔNG HỢP Ý KIẾN VÀ KẾT LUẬN NHẬP HỆ THỐNG DUYỆT GIÁO ÁN\n")
    r_t.font.bold = True
    r_t.font.size = Pt(14)
    r_sub = p_title.add_run("TỔ TOÁN – TIN — ĐỢT THÁNG 9/2026 (NĂM HỌC 2026 - 2027)\n")
    r_sub.font.bold = True
    r_sub.font.size = Pt(12)
    p_guide = doc.add_paragraph()
    p_guide.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_g = p_guide.add_run("(Tài liệu chuyên dụng để Tổ trưởng sao chép trực tiếp vào phần mềm vnEdu, SMAS, K12Online)\n")
    r_g.italic = True
    r_g.font.size = Pt(10.5)

    def add_comment_entry(doc, title, conclusion, comment_text, is_reject=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(f"• {title}: ")
        r1.bold = True
        r2 = p.add_run(f"[{conclusion}]\n")
        r2.bold = True
        r2.font.color.rgb = RGBColor(185, 28, 28) if is_reject else RGBColor(22, 101, 52)

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        cell.width = Inches(7.0)
        border_col = "DC2626" if is_reject else "2563EB"
        bg_col = "FEF2F2" if is_reject else "F0FDF4"
        text_col = RGBColor(185, 28, 28) if is_reject else RGBColor(22, 101, 52)
        set_cell_background(cell, bg_col)
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="none"/>\n'
            f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_col}"/>\n'
            f'  <w:bottom w:val="none"/>\n'
            f'  <w:right w:val="none"/>\n'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(1)
        cp.paragraph_format.space_after = Pt(1)
        r_l = cp.add_run("👉 [MẪU NHẬN XÉT HỆ THỐNG]: ")
        r_l.bold = True
        r_l.font.size = Pt(10.5)
        r_l.font.color.rgb = text_col
        r_c = cp.add_run(f'"{comment_text}"')
        r_c.italic = True
        r_c.font.size = Pt(10.5)
        r_c.font.color.rgb = RGBColor(30, 41, 59)

        p_end = doc.add_paragraph()
        p_end.paragraph_format.space_before = Pt(0)
        p_end.paragraph_format.space_after = Pt(3)

    # 1. THẦY DANH
    p_d_title = doc.add_paragraph()
    p_d_title.paragraph_format.space_before = Pt(10)
    p_d_title.add_run("1. THẦY HỒ ĐĂNG DANH (MÔN TOÁN 6, TOÁN 9)\n").bold = True
    add_comment_entry(
        doc, "Giáo án Phân môn Đại số 9 (14 tiết)", "ĐẠT / DUYỆT",
        "Giáo án soạn đủ 14 tiết theo đúng PPCT, đảm bảo tiến độ tháng 9 và vượt tuần 5; cấu trúc đủ 4 hoạt động theo CV 5512. Tích hợp NLS (5.3) và AI (9.B2) tại Bài 1 đúng địa chỉ. Lưu ý: Cần in đậm, nghiêng câu mô tả năng lực số và AI theo đúng quy định chuyên môn trước khi đưa vào giảng dạy.",
        is_reject=False
    )
    add_comment_entry(
        doc, "Giáo án Phân môn Hình học 9 (03 tiết)", "ĐẠT / DUYỆT (Ghi nhận giải trình lý do chính đáng)",
        "Kế hoạch bài dạy Bài 11 soạn kỹ, đúng phân phối chương trình. Chấp thuận giải trình của giáo viên về lý do chính đáng đối với Tiết 4 (Bài 12) và ghi nhận kế hoạch dạy bù trong tuần tiếp theo. Duyệt.",
        is_reject=False
    )
    add_comment_entry(
        doc, "Giáo án Phân môn Số học 6 (11 tiết)", "ĐẠT / DUYỆT (Ghi nhận giải trình lý do chính đáng)",
        "Soạn 11 tiết bám sát chuẩn kiến thức kỹ năng. Chấp thuận giải trình lý do chính đáng về Tiết 12 và ghi nhận kế hoạch dạy bù. Lưu ý giáo viên in đậm, nghiêng câu mô tả năng lực số tại Bài 6, Bài 7 theo quy định. Duyệt.",
        is_reject=False
    )
    add_comment_entry(
        doc, "Giáo án Phân môn Hình học 6 (05 tiết)", "ĐẠT / DUYỆT",
        "Giáo án soạn đạt yêu cầu, đủ 05 tiết (đảm bảo tiến độ tháng 9 và vượt tuần 5). Cấu trúc 4 hoạt động rõ ràng, phân chia hoạt động GV - HS mạch lạc, hình vẽ và hệ thống câu hỏi phù hợp đối tượng học sinh lớp 6. Duyệt.",
        is_reject=False
    )
    add_comment_entry(
        doc, "Nhận xét chung gói hồ sơ Thầy Danh (Nếu duyệt theo gói)", "DUYỆT (XẾP LOẠI TỐT)",
        "DUYỆT. Hồ sơ giáo án nộp đầy đủ các phân môn Toán 6, Toán 9. Tổ chuyên môn chấp thuận giải trình lý do chính đáng về 02 tiết tuần 4 và ghi nhận kế hoạch dạy bù. Nhắc nhở giáo viên in đậm, nghiêng nội dung tích hợp NLS/AI theo quy định.",
        is_reject=False
    )

    # 2. THẦY HẢI
    p_h_title = doc.add_paragraph()
    p_h_title.paragraph_format.space_before = Pt(12)
    p_h_title.add_run("2. THẦY TRẦN LONG HẢI (MÔN TOÁN 8)\n").bold = True
    add_comment_entry(
        doc, "Bài 10: Tứ giác (Hình 8 — Tiết 1, 2)", "KHÔNG DUYỆT / TRẢ HỒ SƠ",
        "TRẢ HỒ SƠ do có sai sót nghiêm trọng về kiến thức và ký hiệu toán học tại Trang 5: (1) Mất dấu mũ góc ở toàn bộ các công thức (viết trần A, B, C, D); (2) Sai bản chất định lý tổng các góc: viết H + E + F - G = 360° (nhầm dấu trừ); (3) Sai quy tắc chuyển vế và thứ tự phép tính: viết D = 360° - A + B + C = 50° và F = 360° - H - E + G = 125°. Đề nghị đính chính chuẩn xác ký hiệu góc và biểu thức toán học, in đậm nghiêng NLS 1.1, 5.3 trước khi duyệt.",
        is_reject=True
    )
    add_comment_entry(
        doc, "Gói Phân môn Hình học 8 (49 trang — 6 bài)", "KHÔNG DUYỆT / TRẢ LẠI HỒ SƠ",
        "TRẢ LẠI HỒ SƠ. Yêu cầu thầy Trần Long Hải đính chính dứt điểm toàn bộ lỗi kiến thức định lý và ký hiệu góc tại Trang 5 Bài 10, in đậm nghiêng NLS trước khi phê duyệt toàn tệp.",
        is_reject=True
    )
    add_comment_entry(
        doc, "Gói Phân môn Đại số 8 (40 trang — 6 bài)", "ĐẠT YÊU CẦU",
        "ĐẠT YÊU CẦU. Kế hoạch bài dạy soạn đầy đủ 10 tiết, đúng PPCT và vượt tiến độ. Nhắc nhở giáo viên in đậm, nghiêng câu mô tả năng lực số tại Bài 3 theo đúng quy định.",
        is_reject=False
    )

    # 3. CÔ THẢO
    p_t_title = doc.add_paragraph()
    p_t_title.paragraph_format.space_before = Pt(12)
    p_t_title.add_run("3. CÔ NGUYỄN THỊ THẢO (MÔN TOÁN 6, TOÁN 7 — 34 TIẾT, 123 TRANG)\n").bold = True

    add_comment_entry(
        doc, "Phân môn Số học 6 (45 trang — 12 tiết: Bài 1 đến Bài 7 & LTC Tiết 12)", "ĐẠT / DUYỆT (XUẤT SẮC)",
        "Giáo án soạn rất công phu, chuẩn bị đầy đủ 12 tiết đảm bảo 100% tiến độ Tháng 9. Cấu trúc chuẩn 4 hoạt động CV 5512, có nội dung giao việc chu đáo cho học sinh hòa nhập (HSHN). Tích hợp Năng lực số (5.3 và 3.1) tại Bài 6 và Bài 7 đúng địa chỉ và đã IN ĐẬM, NGHIÊNG rất chuẩn xác theo quy chế chuyên môn. Duyệt.",
        is_reject=False
    )
    add_comment_entry(
        doc, "Phân môn Hình học 6 (18 trang — 05 tiết: Bài 18, Bài 19)", "ĐẠT / DUYỆT",
        "Kế hoạch bài dạy soạn tốt, đủ 05 tiết đảm bảo tiến độ tháng 9 và vượt tuần 5. Ký hiệu góc chuẩn xác, hình vẽ trực quan, phân chia hoạt động GV - HS mạch lạc, có bài tập phù hợp cho đối tượng học sinh hòa nhập. Duyệt.",
        is_reject=False
    )
    add_comment_entry(
        doc, "Phân môn Đại số 7 (30 trang — 09 tiết: Bài 1, 2, LTC, Bài 3)", "ĐẠT / DUYỆT",
        "Soạn đủ 09 tiết đảm bảo 100% tiến độ tháng 9 và vượt tuần 5. Cấu trúc 4 hoạt động rõ ràng, các bước thực hiện phép tính số hữu tỉ và lũy thừa chuẩn xác, bài tập củng cố bám sát SGK Kết nối tri thức. Duyệt.",
        is_reject=False
    )
    add_comment_entry(
        doc, "Phân môn Hình học 7 (30 trang — 08 tiết: Bài 8, 9, LTC, Bài 10)", "DUYỆT CÓ LƯU Ý (ĐÍNH CHÍNH KÝ HIỆU GÓC)",
        "Kế hoạch bài dạy đảm bảo đủ 08 tiết đúng tiến độ tháng 9. Nội dung kiến thức hai góc kề bù, đối đỉnh, hai đường thẳng song song và tiên đề Euclid đầy đủ. Lưu ý: Cần chỉnh sửa lại lỗi kỹ thuật hiển thị ký hiệu góc (các góc bị hiển thị thành dấu chấm trên đầu chữ cái như ṁBy, ẋAB... và lỗi font ở trang 8, 9) thành dấu mũ góc chuẩn xác trước khi giảng dạy.",
        is_reject=False
    )
    add_comment_entry(
        doc, "Nhận xét chung toàn bộ hồ sơ Cô Nguyễn Thị Thảo (Duyệt theo gói)", "DUYỆT (XẾP LOẠI KHÁ)",
        "DUYỆT HỒ SƠ (Xếp loại Khá). Hồ sơ nộp đầy đủ 34 tiết (123 trang), tiến độ đạt 100% tháng 9 và vượt tuần 5. Tích hợp NLS Số học 6 in đậm nghiêng rất chuẩn mực. Đề nghị cô Thảo chuẩn hóa lại ký hiệu dấu mũ góc trong phân môn Hình học 7 để nâng xếp loại Tốt trong đợt sau.",
        is_reject=False
    )

    out_sys = os.path.abspath("TROLYTHIEN/3_DUYET_GIAO_AN/Ket_qua/Cap_Nhat_He_Thong_Duyet_Giao_An_Thang_9.docx")
    safe_save_doc(doc, out_sys)

# ==============================================================================
# 3. TẠO PHIẾU NHẬN XÉT CÁ NHÂN CHO CÔ NGUYỄN THỊ THẢO
# ==============================================================================
def create_phieu_thao():
    doc = docx.Document()
    apply_base_page_setup(doc)
    build_header_national(doc, so_bb="... /PNX-TT")

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("PHIẾU NHẬN XÉT HỒ SƠ GIÁO ÁN CÁ NHÂN\n")
    r_t.font.bold = True
    r_t.font.size = Pt(14)
    r_sub = p_title.add_run("ĐỢT THÁNG 9/2026 — NĂM HỌC 2026 - 2027\n")
    r_sub.font.bold = True
    r_sub.font.size = Pt(13)

    p_gv = doc.add_paragraph()
    p_gv.paragraph_format.space_before = Pt(6)
    p_gv.add_run("Họ và tên giáo viên: ").bold = True
    p_gv.add_run("NGUYỄN THỊ THẢO\n").bold = True
    p_gv.add_run("Tổ chuyên môn: ").bold = True
    p_gv.add_run("Toán – Tin\n")
    p_gv.add_run("Nhiệm vụ giảng dạy: ").bold = True
    p_gv.add_run("Toán 6, Toán 7\n")

    p_sec1 = doc.add_paragraph()
    p_sec1.add_run("I. KẾT QUẢ KIỂM TRA SỐ LƯỢNG VÀ TIẾN ĐỘ\n").bold = True

    t_detail = doc.add_table(rows=5, cols=4)
    t_detail.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_detail, color="1F2937", sz="4")
    hdrs = ["Phân môn", "Số tiết đã nộp", "Tiến độ tháng 9", "Ghi chú chuyên môn"]
    for i, h in enumerate(hdrs):
        t_detail.rows[0].cells[i].text = h
        t_detail.rows[0].cells[i].paragraphs[0].runs[0].font.bold = True
        set_cell_background(t_detail.rows[0].cells[i], "E2E8F0")
        set_cell_margins(t_detail.rows[0].cells[i], top=60, bottom=60, left=60, right=60)

    rows_data = [
        ("Số học 6 (45 trang)", "12 tiết: Bài 1 đến 7 & LTC Tiết 12", "Đạt 100% Tháng 9 (Tuần 1-4)", "Tích hợp NLS 5.3 & 3.1 in đậm, nghiêng chuẩn"),
        ("Hình học 6 (18 trang)", "05 tiết: Bài 18, Bài 19", "Đạt và vượt (tuần 5)", "Ký hiệu góc chuẩn, có giao việc cho HSHN"),
        ("Đại số 7 (30 trang)", "09 tiết: Bài 1, 2, LTC, Bài 3", "Đạt và vượt (tuần 5)", "Nội dung chuẩn xác, bám sát SGK"),
        ("Hình học 7 (30 trang)", "08 tiết: Bài 8, 9, LTC, Bài 10", "Đạt 100% Tháng 9 (Tuần 1-4)", "Cần sửa lỗi hiển thị ký hiệu góc")
    ]
    for r_idx, r_val in enumerate(rows_data, start=1):
        for c_idx, val in enumerate(r_val):
            cell = t_detail.rows[r_idx].cells[c_idx]
            cell.text = val
            p = cell.paragraphs[0]
            for run in p.runs:
                run.font.size = Pt(9.5)
            set_cell_margins(cell, top=50, bottom=50, left=50, right=50)

    p_sec2 = doc.add_paragraph()
    p_sec2.paragraph_format.space_before = Pt(8)
    p_sec2.add_run("II. NHẬN XÉT CỦA TỔ CHUYÊN MÔN\n").bold = True
    p_sec2.add_run("1. Ưu điểm: ").bold = True
    p_sec2.add_run("Hồ sơ giáo án chuẩn bị số lượng lớn (34 tiết — 123 trang), tiến độ đảm bảo 100% Tháng 9 và vượt tuần 5 ở các phân môn Hình 6 và Đại 7. Cấu trúc 4 hoạt động theo CV 5512 rõ ràng, chuẩn mực. Đặc biệt, phân môn Số học 6 nộp đủ 12 tiết và tích hợp Năng lực số (NLS 5.3, 3.1) được định dạng IN ĐẬM, NGHIÊNG rất chuẩn xác theo đúng quy chế chuyên môn trường THCS Trần Phú. Kế hoạch bài dạy có phân công hỗ trợ học sinh hòa nhập (HSHN) rất chu đáo.\n")
    p_sec2.add_run("2. Điểm cần đính chính, hoàn thiện: ").bold = True
    p_sec2.add_run("Tại phân môn Hình học 7, toàn bộ ký hiệu góc đang bị lỗi hiển thị thành dấu chấm trên đầu chữ cái đầu tiên (ví dụ: ṁBy, ṁBx, ẋBy, ẋAB, ṀNQ...) hoặc viết trần thay vì dấu mũ góc chuẩn (^). Đề nghị cô Thảo kiểm tra lại phần mềm gõ công thức (MathType/Equation) và đính chính chuẩn xác dấu mũ góc.\n")
    p_sec2.add_run("3. Xếp loại hồ sơ: ").bold = True
    r_xl = p_sec2.add_run("KHÁ (Có điều kiện nâng Tốt sau khi đính chính ký hiệu Hình học 7)\n")
    r_xl.bold = True
    r_xl.font.color.rgb = RGBColor(22, 101, 52)

    add_signatures(doc)

    out_phieu = os.path.abspath("TROLYTHIEN/3_DUYET_GIAO_AN/Ket_qua/Phieu_Nhan_Xet_Ho_So_Nguyen_Thi_Thao_Thang_9.docx")
    safe_save_doc(doc, out_phieu)

if __name__ == "__main__":
    os.makedirs("TROLYTHIEN/3_DUYET_GIAO_AN/Ket_qua", exist_ok=True)
    update_formal_minutes_with_thao()
    update_system_comment_doc_with_thao()
    create_phieu_thao()
    print("HOAN TAT CAP NHAT HO SO CO NGUYEN THI THAO!")
