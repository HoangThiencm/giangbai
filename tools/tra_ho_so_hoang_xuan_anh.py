# -*- coding: utf-8 -*-
"""
ĐIỀU CHỈNH KẾT QUẢ THẨM ĐỊNH HỒ SƠ THẦY HOÀNG XUÂN ÁNH:
TRẢ HỒ SƠ PHÂN MÔN HĐTN 6 DO VI PHẠM QUY CHẾ CHUYÊN MÔN:
TÍCH HỢP NLS VÀ AI KHÔNG IN ĐẬM, KHÔNG IN NGHIÊNG.
Chuẩn hóa 100% thể thức văn bản hành chính theo Nghị định 30/2020/NĐ-CP
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_table_borders_nd30(table, color="000000", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_callout_border(cell, border_color="DC2626", bg_color="FEF2F2"):
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=140)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def apply_page_setup_nd30(doc):
    for s in doc.sections:
        s.page_width = Inches(8.27)
        s.page_height = Inches(11.69)
        s.top_margin = Inches(0.79)      # 20 mm
        s.bottom_margin = Inches(0.79)   # 20 mm
        s.left_margin = Inches(1.18)     # 30 mm
        s.right_margin = Inches(0.59)    # 15 mm
        s.different_first_page_header_footer = False

def safe_save_doc(doc, target_path):
    try:
        doc.save(target_path)
        print("Đã lưu thành công:", target_path)
        return target_path
    except PermissionError:
        base, ext = os.path.splitext(target_path)
        fallback_path = f"{base}_CapNhat{ext}"
        doc.save(fallback_path)
        print(f"Tệp đang mở. Đã lưu sang: {fallback_path}")
        return fallback_path

# ==============================================================================
# 1. TẠO LẠI PHIẾU NHẬN XÉT: HOÀNG XUÂN ÁNH (07/PĐG-TT) — TRẢ HỒ SƠ
# ==============================================================================
def recreate_phieu_anh_reject():
    print("\n--- BẮT ĐẦU TẠO LẠI PHIẾU NHẬN XÉT THẦY HOÀNG XUÂN ÁNH (TRẢ HỒ SƠ) ---")
    doc = docx.Document()
    apply_page_setup_nd30(doc)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(13)

    # 1. Quốc hiệu tiêu ngữ (Table 0)
    tbl_h = doc.add_table(rows=1, cols=2)
    tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
    c0 = tbl_h.cell(0, 0)
    c1 = tbl_h.cell(0, 1)
    set_cell_margins(c0, 0, 0, 0, 0)
    set_cell_margins(c1, 0, 0, 0, 0)

    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(0); p0.paragraph_format.space_after = Pt(0)
    p0.paragraph_format.line_spacing = 1.15
    r = p0.add_run("TRƯỜNG THCS TRẦN PHÚ\n")
    r.font.name = "Times New Roman"; r.font.size = Pt(12)
    r = p0.add_run("TỔ TOÁN – TIN\n")
    r.font.name = "Times New Roman"; r.font.size = Pt(12); r.bold = True
    r = p0.add_run("———————\n")
    r.font.name = "Times New Roman"; r.font.size = Pt(10)
    r = p0.add_run("Số: 07/PĐG-TT")
    r.font.name = "Times New Roman"; r.font.size = Pt(12)

    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0); p1.paragraph_format.space_after = Pt(0)
    p1.paragraph_format.line_spacing = 1.15
    r = p1.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
    r.font.name = "Times New Roman"; r.font.size = Pt(12); r.bold = True
    r = p1.add_run("Độc lập - Tự do - Hạnh phúc\n")
    r.font.name = "Times New Roman"; r.font.size = Pt(13); r.bold = True
    r = p1.add_run("—————————————\n")
    r.font.name = "Times New Roman"; r.font.size = Pt(10)
    r = p1.add_run("Xuân Đông, ngày 30 tháng 9 năm 2026")
    r.font.name = "Times New Roman"; r.font.size = Pt(13); r.italic = True

    # 2. Tiêu đề phiếu
    p_t0 = doc.add_paragraph()
    p_t0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t0.paragraph_format.space_before = Pt(14); p_t0.paragraph_format.space_after = Pt(2)
    r_t0 = p_t0.add_run("PHIẾU NHẬN XÉT, ĐÁNH GIÁ HỒ SƠ BÀI DẠY")
    r_t0.font.name = "Times New Roman"; r_t0.font.size = Pt(14); r_t0.bold = True

    p_t1 = doc.add_paragraph()
    p_t1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t1.paragraph_format.space_before = Pt(0); p_t1.paragraph_format.space_after = Pt(12)
    r_t1 = p_t1.add_run("(ĐỢT THÁNG 9/2026 — NĂM HỌC 2026 - 2027)")
    r_t1.font.name = "Times New Roman"; r_t1.font.size = Pt(13); r_t1.italic = True

    # 3. Thông tin giáo viên
    info_items = [
        ("Họ và tên giáo viên: ", "HOÀNG XUÂN ÁNH", True),
        ("Tổ chuyên môn: ", "Toán – Tin, Trường THCS Trần Phú", False),
        ("Nhiệm vụ giảng dạy: ", "Toán 7 (Đại số 7, Hình học 7) và HĐTN 6", False),
        ("Tổng số hồ sơ nộp: ", "03 phân môn (80 trang PDF, 29 tiết)", False),
    ]

    for label, val, is_val_bold in info_items:
        p_i = doc.add_paragraph()
        p_i.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_i.paragraph_format.first_line_indent = Inches(0.5)
        p_i.paragraph_format.space_before = Pt(2)
        p_i.paragraph_format.space_after = Pt(2 if label != "Tổng số hồ sơ nộp: " else 6)
        p_i.paragraph_format.line_spacing = 1.2
        r_lbl = p_i.add_run(label)
        r_lbl.font.name = "Times New Roman"; r_lbl.font.size = Pt(13); r_lbl.bold = True
        r_val = p_i.add_run(val)
        r_val.font.name = "Times New Roman"; r_val.font.size = Pt(13); r_val.bold = is_val_bold

    # 4. Mục I. KẾT QUẢ THẨM ĐỊNH CHI TIẾT THEO TỪNG PHÂN MÔN
    p_sec1 = doc.add_paragraph()
    p_sec1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_sec1.paragraph_format.space_before = Pt(8); p_sec1.paragraph_format.space_after = Pt(4)
    p_sec1.paragraph_format.line_spacing = 1.2
    r_s1 = p_sec1.add_run("I. KẾT QUẢ THẨM ĐỊNH CHI TIẾT THEO TỪNG PHÂN MÔN")
    r_s1.font.name = "Times New Roman"; r_s1.font.size = Pt(13); r_s1.bold = True

    t_data = doc.add_table(rows=4, cols=5)
    t_data.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_nd30(t_data, color="000000", sz="4")

    headers = ["Phân môn", "Số trang", "Số tiết", "Tiến độ & Tích hợp NLS / AI", "Đánh giá"]
    col_widths = [Inches(1.2), Inches(0.9), Inches(0.8), Inches(2.7), Inches(1.1)]

    for c_idx, h in enumerate(headers):
        cell = t_data.cell(0, c_idx)
        cell.width = col_widths[c_idx]
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.name = "Times New Roman"; r.font.size = Pt(10.5); r.bold = True

    rows_anh = [
        ("Đại số 7", "31 trang", "09 tiết",
         "Đạt 100% Tháng 9 & vượt tuần 5 (Bài 1, 2, LTC, Bài 3). Phép tính số hữu tỉ và lũy thừa chuẩn xác, 4 hoạt động CV 5512 rõ ràng.",
         "DUYỆT (Tốt)"),
        ("Hình học 7", "31 trang", "08 tiết",
         "Đạt 100% Tháng 9 (Bài 8, 9, LTC, Bài 10). Ký hiệu góc và song song hiển thị dấu mũ góc rõ nét, chuẩn mực, không lỗi font.",
         "DUYỆT (Tốt)"),
        ("HĐTN 6", "18 trang", "12 tiết",
         "VI PHẠM QUY CHẾ CHUYÊN MÔN: Các câu mô tả mục tiêu chỉ báo NLS (2.2.TC1a) và AI (6.C2.2) KHÔNG IN ĐẬM, KHÔNG IN NGHIÊNG theo quy định.",
         "TRẢ HỒ SƠ"),
    ]

    for r_idx, row_val in enumerate(rows_anh, start=1):
        for c_idx, val in enumerate(row_val):
            cell = t_data.cell(r_idx, c_idx)
            cell.width = col_widths[c_idx]
            set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [1, 2, 4] else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.name = "Times New Roman"; r.font.size = Pt(10.5)
            if c_idx in [0, 4]:
                r.bold = True
            if c_idx == 4:
                r.font.color.rgb = RGBColor(22, 163, 74) if "DUYỆT" in val else RGBColor(220, 38, 38)

    # 5. Mục II. NHẬN XÉT CHUYÊN MÔN CHI TIẾT
    p_sec2 = doc.add_paragraph()
    p_sec2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_sec2.paragraph_format.space_before = Pt(12); p_sec2.paragraph_format.space_after = Pt(4)
    p_sec2.paragraph_format.line_spacing = 1.2
    r_s2 = p_sec2.add_run("II. NHẬN XÉT CHUYÊN MÔN CHI TIẾT")
    r_s2.font.name = "Times New Roman"; r_s2.font.size = Pt(13); r_s2.bold = True

    # Ưu điểm
    p_u = doc.add_paragraph()
    p_u.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_u.paragraph_format.first_line_indent = Inches(0.5)
    p_u.paragraph_format.space_before = Pt(2); p_u.paragraph_format.space_after = Pt(3)
    p_u.paragraph_format.line_spacing = 1.2
    r_ulbl = p_u.add_run("1. Ưu điểm: ")
    r_ulbl.font.name = "Times New Roman"; r_ulbl.font.size = Pt(13); r_ulbl.bold = True
    r_utxt = p_u.add_run(
        "Hồ sơ bài dạy môn Toán 7 (Đại số 7 và Hình học 7) soạn rất chu đáo, đảm bảo 100% tiến độ Tháng 9 và vượt tuần 5 ở Đại 7. Cấu trúc 4 hoạt động CV 5512 rõ ràng. Thể thức văn bản, Header/Footer đúng chuẩn quy định của trường THCS Trần Phú. Ký hiệu góc và công thức hình học trong Hình học 7 hiển thị chuẩn mực, rõ nét, không lỗi font."
    )
    r_utxt.font.name = "Times New Roman"; r_utxt.font.size = Pt(13)

    # Tồn tại vi phạm quy chế
    p_l = doc.add_paragraph()
    p_l.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_l.paragraph_format.first_line_indent = Inches(0.5)
    p_l.paragraph_format.space_before = Pt(2); p_l.paragraph_format.space_after = Pt(3)
    p_l.paragraph_format.line_spacing = 1.2
    r_llbl = p_l.add_run("2. Tồn tại vi phạm quy chế chuyên môn (Phân môn HĐTN 6): ")
    r_llbl.font.name = "Times New Roman"; r_llbl.font.size = Pt(13); r_llbl.bold = True
    r_llbl.font.color.rgb = RGBColor(185, 28, 28)
    r_ltxt = p_l.add_run(
        "Tại phân môn HĐTN 6 (Chủ đề 1), giáo viên vi phạm quy định bắt buộc của tổ chuyên môn về hình thức tích hợp: Toàn bộ các câu mô tả chỉ báo Năng lực số (2.2.TC1a) và Trí tuệ nhân tạo (6.C2.2) tại Trang 1, Trang 10 đều KHÔNG ĐƯỢC IN ĐẬM, KHÔNG ĐƯỢC IN NGHIÊNG mà để chữ đứng, bôi màu đỏ bình thường. Quy chế chuyên môn bắt buộc 100% mục tiêu tích hợp phải định dạng in đậm, nghiêng để thuận tiện theo dõi, giám sát."
    )
    r_ltxt.font.name = "Times New Roman"; r_ltxt.font.size = Pt(13)

    # Khoảng đệm trước Table 2
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0); p_sp.paragraph_format.space_after = Pt(4)

    # Table 2: Callout kết luận và yêu cầu TRẢ HỒ SƠ
    tbl_callout = doc.add_table(rows=1, cols=1)
    tbl_callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_call = tbl_callout.cell(0, 0)
    set_callout_border(c_call, border_color="DC2626", bg_color="FEF2F2")

    p_call = c_call.paragraphs[0]
    p_call.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_call.paragraph_format.space_before = Pt(2); p_call.paragraph_format.space_after = Pt(2)
    p_call.paragraph_format.line_spacing = 1.2
    r_clbl = p_call.add_run("👉 KẾT LUẬN & YÊU CẦU: ")
    r_clbl.font.name = "Times New Roman"; r_clbl.font.size = Pt(13); r_clbl.bold = True
    r_clbl.font.color.rgb = RGBColor(185, 28, 28)
    r_ctxt = p_call.add_run(
        "Tổ chuyên môn TRẢ HỒ SƠ phân môn Hoạt động trải nghiệm, hướng nghiệp 6 (HĐTN 6) của thầy Hoàng Xuân Ánh. Yêu cầu thầy Ánh định dạng IN ĐẬM, NGHIÊNG chuẩn xác toàn bộ câu mô tả mục tiêu chỉ báo NLS (2.2.TC1a) và AI (6.C2.2) theo đúng quy chế chuyên môn trường THCS Trần Phú và nộp lại trước ngày 10/10/2026. (Phân môn Toán 7 nghiệm thu đạt yêu cầu)."
    )
    r_ctxt.font.name = "Times New Roman"; r_ctxt.font.size = Pt(13)
    r_ctxt.font.color.rgb = RGBColor(30, 41, 59)

    # Khoảng đệm trước Table 3
    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(4); p_sp2.paragraph_format.space_after = Pt(6)

    # Table 3: Chữ ký 2 bên
    tbl_sign = doc.add_table(rows=1, cols=2)
    tbl_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    cs0 = tbl_sign.cell(0, 0)
    cs1 = tbl_sign.cell(0, 1)
    set_cell_margins(cs0, 60, 0, 0, 0)
    set_cell_margins(cs1, 60, 0, 0, 0)

    ps0 = cs0.paragraphs[0]
    ps0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ps0.paragraph_format.line_spacing = 1.15
    r_s0 = ps0.add_run("GIÁO VIÊN BỘ MÔN\n\n\n\n\n")
    r_s0.font.name = "Times New Roman"; r_s0.font.size = Pt(13); r_s0.bold = True
    r_n0 = ps0.add_run("Hoàng Xuân Ánh")
    r_n0.font.name = "Times New Roman"; r_n0.font.size = Pt(13); r_n0.bold = True

    ps1 = cs1.paragraphs[0]
    ps1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ps1.paragraph_format.line_spacing = 1.15
    r_s1 = ps1.add_run("TỔ TRƯỞNG CHUYÊN MÔN\n\n\n\n\n")
    r_s1.font.name = "Times New Roman"; r_s1.font.size = Pt(13); r_s1.bold = True
    r_n1 = ps1.add_run("Hoàng Tấn Thiên")
    r_n1.font.name = "Times New Roman"; r_n1.font.size = Pt(13); r_n1.bold = True

    out_phieu = r"TROLYTHIEN/3_DUYET_GIAO_AN/Ket_qua/Phieu_Nhan_Xet_Ho_So_Hoang_Xuan_Anh_Thang_9.docx"
    safe_save_doc(doc, out_phieu)
    print("HOÀN THÀNH TẠO LẠI PHIẾU NHẬN XÉT THẦY ÁNH (TRẢ HỒ SƠ)!")

# ==============================================================================
# 2. CẬP NHẬT BIÊN BẢN TỔ (Bien_Ban_Kiem_Tra_Ho_So_To_Toan_Thang_9)
# ==============================================================================
def update_bien_ban_to_reject():
    print("\n--- BẮT ĐẦU CẬP NHẬT BIÊN BẢN TỔ (THẦY ÁNH TRẢ HỒ SƠ) ---")
    docx_path = r"TROLYTHIEN/3_DUYET_GIAO_AN/Ket_qua/Bien_Ban_Kiem_Tra_Ho_So_To_Toan_Thang_9.docx"
    doc = docx.Document(docx_path)

    # 1. Bảng III (Table 1) - Row 7
    t = doc.tables[1]
    for r in t.rows:
        if "hoàng xuân ánh" in r.cells[1].text.lower():
            cells = r.cells
            cells[4].text = "Toán 7 đạt tốt; HĐTN 6 vi phạm quy định NLS/AI (không in đậm, nghiêng)"
            cells[5].text = "TRẢ HỒ SƠ\n(Sửa HĐTN 6)"
            for ci in [4, 5]:
                p = cells[ci].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 4 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10.5)
                    if ci == 5:
                        run.bold = True
                        run.font.color.rgb = RGBColor(220, 38, 38) # ĐỎ TRẢ HỒ SƠ

    # 2. Mục IV.7: Đánh giá chi tiết Thầy Hoàng Xuân Ánh
    idx_p7 = -1
    for i, p in enumerate(doc.paragraphs):
        if "7. Thầy Hoàng Xuân Ánh" in p.text:
            idx_p7 = i
            break

    if idx_p7 != -1:
        # Cập nhật kết luận tại idx_p7 + 4
        for k in range(idx_p7, min(idx_p7 + 6, len(doc.paragraphs))):
            p = doc.paragraphs[k]
            if "- Kết luận và xếp loại:" in p.text:
                p.text = ""
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.first_line_indent = Inches(0.5)
                p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.2
                r1 = p.add_run("- Kết luận và xếp loại: ")
                r1.font.name = "Times New Roman"; r1.font.size = Pt(13); r1.bold = True
                r2 = p.add_run("Đề nghị trả hồ sơ phân môn HĐTN 6. ")
                r2.font.name = "Times New Roman"; r2.font.size = Pt(13); r2.bold = True
                r2.font.color.rgb = RGBColor(220, 38, 38)
                r3 = p.add_run("Yêu cầu thầy Hoàng Xuân Ánh định dạng in đậm, nghiêng toàn bộ câu mô tả mục tiêu chỉ báo NLS (2.2.TC1a) và AI (6.C2.2) trong HĐTN 6 theo đúng quy chế chuyên môn trước khi trình duyệt lại.")
                r3.font.name = "Times New Roman"; r3.font.size = Pt(13)

    # 3. Mục V.1 và V.3
    for p in doc.paragraphs:
        if "1. Đánh giá chung: Đợt kiểm tra hồ sơ tháng 9/2026 đã tiến hành thẩm định" in p.text:
            p.text = ""
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Inches(0.5)
            p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.2
            r1 = p.add_run("1. Đánh giá chung: ")
            r1.font.name = "Times New Roman"; r1.font.size = Pt(13); r1.bold = True
            r2 = p.add_run("Đợt kiểm tra hồ sơ tháng 9/2026 đã tiến hành thẩm định 07 giáo viên trong tổ. Kết quả: 06 giáo viên đạt chuẩn được phê duyệt chính thức (trong đó 05 giáo viên xếp loại Tốt, 01 giáo viên xếp loại Khá); 01 giáo viên tạm thời trả hồ sơ để hoàn thiện định dạng chỉ báo NLS/AI môn HĐTN 6 (thầy Hoàng Xuân Ánh).")
            r2.font.name = "Times New Roman"; r2.font.size = Pt(13)

        if "3. Ghi nhận khắc phục:" in p.text or "3. Yêu cầu đối với giáo viên trả hồ sơ:" in p.text:
            p.text = ""
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Inches(0.5)
            p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.2
            r1 = p.add_run("3. Yêu cầu đối với giáo viên trả hồ sơ: ")
            r1.font.name = "Times New Roman"; r1.font.size = Pt(13); r1.bold = True
            r2 = p.add_run("Thầy Hoàng Xuân Ánh khẩn trương hoàn thiện việc in đậm, nghiêng câu mô tả chỉ báo NLS và AI trong tệp HĐTN 6 và nộp lại hồ sơ cho tổ chuyên môn trước ngày 10 tháng 10 năm 2026.")
            r2.font.name = "Times New Roman"; r2.font.size = Pt(13)

    safe_save_doc(doc, docx_path)
    print("HOÀN THÀNH CẬP NHẬT BIÊN BẢN TỔ DOCX!")

    # Cập nhật Bien_Ban_Kiem_Tra_Ho_So_To_Toan_Thang_9.md
    md_bb_path = r"TROLYTHIEN/3_DUYET_GIAO_AN/Ket_qua/Bien_Ban_Kiem_Tra_Ho_So_To_Toan_Thang_9.md"
    with open(md_bb_path, "r", encoding="utf-8") as f:
        bb_md = f.read()

    old_r7_md = "| **7** | **HOÀNG XUÂN ÁNH** | Toán 7, HĐTN 6 | 29 tiết (80 trang) | Đủ 100% Tháng 9 & vượt Tuần 5; NLS HĐTN 6 lưu ý in đậm nghiêng | **DUYỆT**<br>*(Xếp loại: Tốt)* |"
    new_r7_md = "| **7** | **HOÀNG XUÂN ÁNH** | Toán 7, HĐTN 6 | 29 tiết (80 trang) | Toán 7 đạt tốt; HĐTN 6 vi phạm NLS/AI không in đậm nghiêng | **TRẢ HỒ SƠ**<br>*(Sửa HĐTN 6)* |"
    bb_md = bb_md.replace(old_r7_md, new_r7_md)

    old_s7_md = "- **Kết luận và xếp loại**: **Duyệt hồ sơ (Xếp loại: Tốt)**. Đề nghị thầy Ánh hoàn thiện in đậm nghiêng mục tiêu NLS/AI HĐTN 6."
    new_s7_md = "- **Kết luận và xếp loại**: **Đề nghị trả hồ sơ**. Yêu cầu thầy Hoàng Xuân Ánh định dạng in đậm, nghiêng toàn bộ câu mô tả mục tiêu chỉ báo NLS (2.2.TC1a) và AI (6.C2.2) trong HĐTN 6 theo đúng quy chế chuyên môn trước khi trình duyệt lại."
    bb_md = bb_md.replace(old_s7_md, new_s7_md)

    old_v_md = "Đợt kiểm tra hồ sơ tháng 9/2026 đã tiến hành thẩm định 07 giáo viên trong tổ. Kết quả: 100% giáo viên (07/07 đồng chí) đạt chuẩn được phê duyệt chính thức (trong đó 06 giáo viên xếp loại Tốt, 01 giáo viên xếp loại Khá). Không còn giáo viên nào bị trả hồ sơ."
    new_v_md = "Đợt kiểm tra hồ sơ tháng 9/2026 đã tiến hành thẩm định 07 giáo viên trong tổ. Kết quả: 06 giáo viên đạt chuẩn được phê duyệt chính thức (trong đó 05 giáo viên xếp loại Tốt, 01 giáo viên xếp loại Khá); 01 giáo viên tạm thời trả hồ sơ để hoàn thiện định dạng chỉ báo NLS/AI môn HĐTN 6 (thầy Hoàng Xuân Ánh)."
    bb_md = bb_md.replace(old_v_md, new_v_md)

    with open(md_bb_path, "w", encoding="utf-8") as f:
        f.write(bb_md)
    print("HOÀN THÀNH CẬP NHẬT BIÊN BẢN TỔ MARKDOWN!")

# ==============================================================================
# 3. CẬP NHẬT TÀI LIỆU NỘP HỆ THỐNG (Cap_Nhat_He_Thong_Duyet_Giao_An_Thang_9)
# ==============================================================================
def update_cap_nhat_he_thong_reject():
    print("\n--- BẮT ĐẦU CẬP NHẬT TÀI LIỆU NỘP HỆ THỐNG (THẦY ÁNH TRẢ HỒ SƠ) ---")
    docx_sys_path = r"TROLYTHIEN/3_DUYET_GIAO_AN/Ket_qua/Cap_Nhat_He_Thong_Duyet_Giao_An_Thang_9.docx"
    doc = docx.Document(docx_sys_path)

    for i, p in enumerate(doc.paragraphs):
        if "Phân môn Hoạt động trải nghiệm, hướng nghiệp 6" in p.text and "THẦY HOÀNG XUÂN ÁNH" not in p.text:
            # Kiểm tra xem có phải của thầy Ánh không (sau Mục 7)
            pass
        if "7. THẦY HOÀNG XUÂN ÁNH" in p.text:
            idx_sec7 = i
            # Tìm các đoạn phân môn sau idx_sec7
            for j in range(idx_sec7, len(doc.paragraphs)):
                pj = doc.paragraphs[j]
                if "Phân môn Hoạt động trải nghiệm, hướng nghiệp 6" in pj.text:
                    pj.text = "• Phân môn Hoạt động trải nghiệm, hướng nghiệp 6 (18 trang — 12 tiết: Chủ đề 1): [KHÔNG DUYỆT / TRẢ HỒ SƠ]"
                    for r in pj.runs:
                        r.font.name = "Times New Roman"; r.font.size = Pt(13); r.bold = True
                        if "[KHÔNG DUYỆT" in r.text:
                            r.font.color.rgb = RGBColor(220, 38, 38)
                if "Nhận xét chung toàn bộ hồ sơ Thầy Hoàng Xuân Ánh" in pj.text:
                    pj.text = "• Nhận xét chung toàn bộ hồ sơ Thầy Hoàng Xuân Ánh (Duyệt theo gói): [TRẢ HỒ SƠ - CHỜ ĐÍNH CHÍNH HĐTN 6]"
                    for r in pj.runs:
                        r.font.name = "Times New Roman"; r.font.size = Pt(13); r.bold = True
                        if "[TRẢ HỒ SƠ" in r.text:
                            r.font.color.rgb = RGBColor(220, 38, 38)

    # Cập nhật nội dung trong callout của Thầy Ánh
    for t in doc.tables:
        c = t.rows[0].cells[0]
        txt = c.text
        if "Soạn đủ 12 tiết Chủ đề 1 đúng PPCT tháng 9. Cấu trúc phân chia hoạt động chung" in txt:
            set_callout_border(c, border_color="DC2626", bg_color="FEF2F2")
            p = c.paragraphs[0]
            p.text = ""
            r1 = p.add_run("👉 Nhận xét hệ thống (Copy & Paste): ")
            r1.font.name = "Times New Roman"; r1.font.size = Pt(11); r1.bold = True
            r1.font.color.rgb = RGBColor(185, 28, 28)
            r2 = p.add_run("TRẢ HỒ SƠ PHÂN MÔN HĐTN 6 do vi phạm quy chế chuyên môn: Toàn bộ câu mô tả mục tiêu chỉ báo Năng lực số (2.2.TC1a) và AI (6.C2.2) tại Trang 1, Trang 10 KHÔNG IN ĐẬM, KHÔNG IN NGHIÊNG theo quy định bắt buộc. Yêu cầu giáo viên định dạng in đậm, nghiêng chuẩn mực và nộp lại để phê duyệt.")
            r2.font.name = "Times New Roman"; r2.font.size = Pt(11)
            r2.font.color.rgb = RGBColor(30, 41, 59)

        elif "DUYỆT TOÀN BỘ HỒ SƠ (Xếp loại: Tốt). Hồ sơ giáo án nộp đầy đủ 29 tiết (80 trang) thuộc 3 phân môn Toán 7, HĐTN 6" in txt:
            set_callout_border(c, border_color="DC2626", bg_color="FEF2F2")
            p = c.paragraphs[0]
            p.text = ""
            r1 = p.add_run("👉 Nhận xét hệ thống (Copy & Paste): ")
            r1.font.name = "Times New Roman"; r1.font.size = Pt(11); r1.bold = True
            r1.font.color.rgb = RGBColor(185, 28, 28)
            r2 = p.add_run("TRẢ LẠI HỒ SƠ. Nghiệm thu đạt yêu cầu 2 phân môn Toán 7 (Đại số 7 và Hình học 7). Tạm thời trả hồ sơ phân môn HĐTN 6 do vi phạm quy định định dạng NLS/AI (không in đậm, nghiêng). Đề nghị thầy Hoàng Xuân Ánh hoàn thiện và nộp lại trước ngày 10/10/2026.")
            r2.font.name = "Times New Roman"; r2.font.size = Pt(11)
            r2.font.color.rgb = RGBColor(30, 41, 59)

    safe_save_doc(doc, docx_sys_path)
    print("HOÀN THÀNH CẬP NHẬT TÀI LIỆU NỘP HỆ THỐNG DOCX!")

    # Cập nhật Cap_Nhat_He_Thong_Duyet_Giao_An_Thang_9.md
    md_sys_path = r"TROLYTHIEN/3_DUYET_GIAO_AN/Ket_qua/Cap_Nhat_He_Thong_Duyet_Giao_An_Thang_9.md"
    with open(md_sys_path, "r", encoding="utf-8") as f:
        sys_md = f.read()

    old_anh_md = """### • Phân môn Hoạt động trải nghiệm, hướng nghiệp 6 (18 trang — 12 tiết: Chủ đề 1: Em với nhà trường): [ĐẠT / DUYỆT (KHÁ/TỐT)]
> **👉 Nhận xét hệ thống (Copy & Paste):**  
> `Soạn đủ 12 tiết Chủ đề 1 đúng PPCT tháng 9. Cấu trúc phân chia hoạt động chung và hoạt động lớp rõ ràng, có phiếu học tập số và bảng rubric tự đánh giá. Tích hợp NLS (2.2.TC1a) và AI (6.C2.2) đúng địa chỉ. Lưu ý: Cần bổ sung in đậm, nghiêng câu mô tả NLS và AI theo đúng quy định chuyên môn trước khi giảng dạy. Duyệt.`

### • Nhận xét chung toàn bộ hồ sơ Thầy Hoàng Xuân Ánh (Duyệt theo gói): [DUYỆT - XẾP LOẠI TỐT]
> **👉 Nhận xét hệ thống (Copy & Paste):**  
> `DUYỆT TOÀN BỘ HỒ SƠ (Xếp loại: Tốt). Hồ sơ giáo án nộp đầy đủ 29 tiết (80 trang) thuộc 3 phân môn Toán 7, HĐTN 6. Đảm bảo 100% tiến độ Tháng 9 và vượt tuần 5 ở Đại 7. Ký hiệu toán học và thể thức chuẩn mực. Nhắc nhở giáo viên in đậm, nghiêng nội dung tích hợp NLS và AI trong HĐTN 6.`"""

    new_anh_md = """### • Phân môn Hoạt động trải nghiệm, hướng nghiệp 6 (18 trang — 12 tiết: Chủ đề 1: Em với nhà trường): [KHÔNG DUYỆT / TRẢ HỒ SƠ]
> **👉 Nhận xét hệ thống (Copy & Paste):**  
> `TRẢ HỒ SƠ PHÂN MÔN HĐTN 6 do vi phạm quy chế chuyên môn: Toàn bộ câu mô tả mục tiêu chỉ báo Năng lực số (2.2.TC1a) và AI (6.C2.2) tại Trang 1, Trang 10 KHÔNG IN ĐẬM, KHÔNG IN NGHIÊNG theo quy định bắt buộc. Yêu cầu giáo viên định dạng in đậm, nghiêng chuẩn mực và nộp lại để phê duyệt.`

### • Nhận xét chung toàn bộ hồ sơ Thầy Hoàng Xuân Ánh (Duyệt theo gói): [TRẢ HỒ SƠ - CHỜ ĐÍNH CHÍNH HĐTN 6]
> **👉 Nhận xét hệ thống (Copy & Paste):**  
> `TRẢ LẠI HỒ SƠ. Nghiệm thu đạt yêu cầu 2 phân môn Toán 7 (Đại số 7 và Hình học 7). Tạm thời trả hồ sơ phân môn HĐTN 6 do vi phạm quy định định dạng NLS/AI (không in đậm, nghiêng). Đề nghị thầy Hoàng Xuân Ánh hoàn thiện và nộp lại trước ngày 10/10/2026.`"""

    if old_anh_md in sys_md:
        sys_md = sys_md.replace(old_anh_md, new_anh_md)

    with open(md_sys_path, "w", encoding="utf-8") as f:
        f.write(sys_md)
    print("HOÀN THÀNH CẬP NHẬT TÀI LIỆU NỘP HỆ THỐNG MARKDOWN!")

if __name__ == "__main__":
    recreate_phieu_anh_reject()
    update_bien_ban_to_reject()
    update_cap_nhat_he_thong_reject()
    print("\n>>> ĐÃ HOÀN TẤT ĐIỀU CHỈNH TRẢ HỒ SƠ THẦY HOÀNG XUÂN ÁNH! <<<")
